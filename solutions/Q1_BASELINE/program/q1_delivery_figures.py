"""Archive and plot the existing Q1 Time A ledger without importing a solver."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
import shutil
import sys
from pathlib import Path

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
BASE = HERE.parent
REPO = BASE.parents[1]
sys.path.insert(0, str(HERE / "vendor" / "figure_skill"))

import matplotlib

matplotlib.use("Agg")
import matplotlib.font_manager as fm
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap
import numpy as np
import openpyxl
import pandas as pd
from PIL import Image
from pypdf import PdfReader

from check_figure import check_figure
from export_figure import export_figure
from figure_audit import audit_figure_directory
from profile_data import profile_data, render_report
from setup_style import setup_style
from visual_qa import audit_layout

RESULTS = BASE / "results"
OTHERS = BASE / "others"
FIGURES = BASE / "figures"
SOURCE = BASE / "result"
DT = 1 / 6
TOL = 1e-5
COLORS = dict(load="#0072B2", pv="#E69F00", grid="#CC79A7",
              charge="#56B4E9", discharge="#D55E00", energy="#009E73",
              reference="#666666", balance="#0072B2", state="#D55E00")
SPECS = [
    ("raw_q1_01_inputs", "raw", "step", "负荷、光伏与电价的时序关系",
     "上下两面板分别展示功率和电价，避免双纵轴造成视觉假相关。", "输入时序及电价"),
    ("raw_q1_02_net_load_distribution", "raw", "histogram", "净负荷的符号和取值分布",
     "144个相邻时段的净负荷直方图，频数之和为144；不是独立重复试验。", "净负荷分布"),
    ("raw_q1_03_load_pv_relation", "raw", "scatter", "光伏对同期负荷的覆盖关系",
     "每点对应一个时段，虚线为光伏等于负荷；不拟合回归或作因果解释。", "负荷与光伏匹配"),
    ("process_q1_01_residuals", "process", "line", "现有账本的能量守恒与状态递推一致性",
     "重算残差按10^-12 kWh缩放展示，验收容差为10^-5 kWh，未画入坐标范围；不是收敛轨迹。", "守恒与状态残差"),
    ("process_q1_02_constraint_activity", "process", "heatmap", "容量和功率边界的约束活跃时段",
     "使用绝对容差10^-5 kWh判断触界，容量行表示时段起点状态；最终状态另在SOC图展示。", "约束触界位置"),
    ("process_q1_03_block_flows", "process", "grouped_bar", "连续四小时段内的充放电转移",
     "每组为连续24段的总电量，不是样本均值；横轴使用旧Time A实际序列边界。", "四小时充放电汇总"),
    ("result_q1_01_dispatch", "result", "step", "购电与储能动作共同满足净负荷",
     "上图为购电和净负荷，下图充电为正、放电为负；电量除以1/6 h得到区间平均功率。", "购电与储能调度"),
    ("result_q1_02_soc", "result", "line", "储能状态满足容量边界和序列闭环",
     "145个边界状态；SOC按12000 kWh额定容量计算；首末均为50%。折线只连接边界值。", "储能状态轨迹"),
    ("result_q1_03_cost", "result", "bar", "现有储能方案降低该实例的购电费用",
     "两个确定性单序列总费用；无储能参照为sum(p*max(load-PV,0)/6)，不调用优化器。", "购电费用对比"),
]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, np.generic):
        return clean(value.item())
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(clean(value), ensure_ascii=False, indent=2,
                               allow_nan=False) + "\n", encoding="utf-8")


def frozen_paths():
    return sorted([p for p in SOURCE.rglob("*") if p.is_file()] +
                  list(BASE.glob("*.md")) + [HERE / "q1_reproduce_v2.py"])


def snapshot(paths):
    return {str(p.relative_to(BASE)): digest(p) for p in paths}


def archive():
    mapping = []
    pairs = [(p, RESULTS / "baseline" / p.relative_to(SOURCE))
             for p in SOURCE.rglob("*") if p.is_file()]
    pairs.extend((BASE / name, OTHERS / "baseline_reports" / name) for name in
                 ["Q1_baseline论文报告.md", "交付清单.md", "REPOSITORY_STATUS.md"])
    for src, dst in sorted(pairs):
        dst.parent.mkdir(parents=True, exist_ok=True)
        if dst.exists() and digest(src) != digest(dst):
            raise AssertionError(f"Refusing to overwrite divergent archive: {dst}")
        if not dst.exists():
            shutil.copy2(src, dst)
        assert digest(src) == digest(dst)
        mapping.append(dict(source=str(src.relative_to(BASE)),
                            canonical_copy=str(dst.relative_to(BASE)), sha256=digest(src)))
    write_json(OTHERS / "archive_mapping.json", mapping)
    return mapping


def clock(minutes):
    return f"{minutes % 1440 // 60:02d}:{minutes % 60:02d}" + ("+1" if minutes >= 1440 else "")


def read_and_check():
    d = pd.read_csv(RESULTS / "baseline/q1_dispatch.csv", float_precision="round_trip")
    summary = json.loads((RESULTS / "baseline/q1_summary.json").read_text(encoding="utf-8-sig"))
    mapping = pd.read_csv(RESULTS / "baseline/q1_time_mapping.csv")
    blocks = pd.read_csv(RESULTS / "baseline/q1_four_hour_blocks.csv", float_precision="round_trip")
    specified = pd.read_csv(RESULTS / "baseline/q1_specified_intervals.csv", float_precision="round_trip")
    assert len(d) == len(mapping) == 144 and len(blocks) == len(specified) == 6
    assert not d.isna().any().any() and np.isfinite(d.select_dtypes("number")).all().all()
    np.testing.assert_array_equal(d.step, np.arange(144))
    for col, shift in [("internal_start", 10), ("internal_end", 20)]:
        assert d[col].tolist() == [clock(shift + 10 * i) for i in range(144)]
    for col in mapping.columns:
        assert d[col].tolist() == mapping[col].tolist(), col
    q, c, out, w = (d[name].to_numpy() for name in
                    ["purchase_kwh", "charge_kwh", "discharge_kwh", "unused_supply_kwh"])
    e = np.r_[d.energy_before_kwh.to_numpy(), d.energy_after_kwh.iloc[-1]]
    np.testing.assert_allclose(e[1:], d.energy_after_kwh, atol=TOL, rtol=0)
    np.testing.assert_allclose(d.load_kwh, d.load_kw * DT, atol=1e-10, rtol=0)
    np.testing.assert_allclose(d.pv_kwh, d.pv_kw * DT, atol=1e-10, rtol=0)
    balance = q + d.pv_kwh.to_numpy() + out - d.load_kwh.to_numpy() - c - w
    state = e[1:] - e[:-1] - 0.9 * c + out / 0.9
    no_storage = np.maximum(d.load_kwh.to_numpy() - d.pv_kwh.to_numpy(), 0)
    cash = float(d.price_yuan_per_kwh @ q)
    reference = float(d.price_yuan_per_kwh @ no_storage)
    assert max(abs(balance).max(), abs(state).max()) <= TOL
    assert min(q.min(), c.min(), out.min(), w.min()) >= -TOL
    assert max(c.max(), out.max()) <= 5000 * DT + TOL
    assert e.min() >= 1200 - TOL and e.max() <= 10800 + TOL
    np.testing.assert_allclose(e[[0, -1]], [6000, 6000], atol=TOL, rtol=0)
    assert not ((c > 1e-7) & (out > 1e-7)).any()
    np.testing.assert_allclose(d.purchase_cost_yuan, d.price_yuan_per_kwh * q, atol=1e-8, rtol=0)
    for actual, stored in [(cash, summary["result"]["total_purchase_cost_yuan"]),
                           (reference, summary["result"]["no_storage_cost_yuan"]),
                           (q.sum(), summary["result"]["total_purchase_kwh"])]:
        np.testing.assert_allclose(actual, stored, atol=1e-7, rtol=0)
    for i, row in blocks.iterrows():
        sl = slice(i * 24, (i + 1) * 24)
        assert row.actual_sequence_interval == f"{clock(10 + i * 240)}-{clock(250 + i * 240)}"
        np.testing.assert_allclose([c[sl].sum(), out[sl].sum()],
                                   [row.charge_kwh, row.discharge_kwh], atol=TOL, rtol=0)
    for _, row in specified.iterrows():
        slots = d.internal_start + "-" + d.internal_end
        values = d.loc[slots == row.internal_interval, "purchase_kwh"]
        assert len(values) == 1
        np.testing.assert_allclose(values.iloc[0], row.purchase_kwh, atol=TOL, rtol=0)

    input_path = REPO / "data/附件1.xlsx"
    template_path = REPO / "data/附件5/result1.xlsx"
    inputs = dict(status="NOT_PRESENT", note="Figure regeneration uses the archived ledger only.")
    if input_path.exists():
        wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
        sheet = wb.worksheets[0]
        rows = list(sheet.values)
        assert (sheet.max_row, sheet.max_column) == (145, 4)
        assert rows[0] == ("时间", "电价", "小区负载", "光伏发电预测功率")
        assert str(rows[1][0]) == "00:10:00" and rows[-1][0] == "0:00+1"
        np.testing.assert_allclose(np.array([row[1:] for row in rows[1:]], dtype=float),
                                   d[["price_yuan_per_kwh", "load_kw", "pv_kw"]], atol=1e-10, rtol=0)
        assert [str(row[0]) for row in rows[1:]] == d.input_time_label.tolist()
        inputs = dict(status="PASS", rows=144, columns=4, explicit_header=list(rows[0]),
                      first_label=str(rows[1][0]), last_label=str(rows[-1][0]), sha256=digest(input_path))
        wb.close()

    wb_path = RESULTS / "baseline/result1_v2.xlsx"
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=False)
    assert len(wb.worksheets) == 2
    plan, battery = wb.worksheets
    assert (plan.max_row, plan.max_column) == (145, 2)
    assert (battery.max_row, battery.max_column) == (7, 5)
    for sheet in wb.worksheets:
        for row in sheet:
            assert all(cell.data_type not in {"e", "f"} for cell in row), "Unexpected error or formula"
    assert [plan.cell(i + 2, 1).value for i in range(144)] == d.template_label.tolist()
    wb_q = np.array([plan.cell(i + 2, 2).value for i in range(144)], float)
    np.testing.assert_allclose(wb_q, q, atol=TOL, rtol=0)
    for i, row in blocks.iterrows():
        np.testing.assert_allclose([battery.cell(i + 2, j).value for j in (2, 3)],
                                   [row.charge_kwh, row.discharge_kwh], atol=TOL, rtol=0)
    np.testing.assert_allclose([battery.cell(i, 5).value for i in (2, 3)], e[[0, -1]], atol=TOL, rtol=0)
    if template_path.exists():
        template = openpyxl.load_workbook(template_path, read_only=True, data_only=False)
        assert template.sheetnames == wb.sheetnames
        writable = {(0, r, 2) for r in range(2, 146)}
        writable |= {(1, r, c) for r in range(2, 8) for c in (2, 3)} | {(1, 2, 5), (1, 3, 5)}
        for k, sheet in enumerate(template.worksheets):
            assert (sheet.max_row, sheet.max_column) == (wb.worksheets[k].max_row, wb.worksheets[k].max_column)
            for row_index, row in enumerate(sheet, 1):
                for col_index, cell in enumerate(row, 1):
                    if (k, row_index, col_index) not in writable:
                        assert cell.value == wb.worksheets[k].cell(row_index, col_index).value
        inputs["template_sha256"] = digest(template_path)
        template.close()
    workbook_audit = dict(status="PASS", sheets=wb.sheetnames, numeric_plan_rows=144,
                          max_plan_difference_kwh=float(abs(wb_q-q).max()),
                          no_formulas_or_error_cells=True, sha256=digest(wb_path))
    wb.close()
    metrics = dict(purchase_cost_yuan=cash, no_storage_cost_yuan=reference,
                   saving_yuan=reference-cash, saving_fraction=1-cash/reference,
                   purchase_kwh=float(q.sum()), total_charge_kwh=float(c.sum()),
                   total_discharge_kwh=float(out.sum()), battery_loss_kwh=float(c.sum()-out.sum()),
                   max_balance_residual_kwh=float(abs(balance).max()),
                   max_state_residual_kwh=float(abs(state).max()),
                   soc_min_percent=float(e.min()/120), soc_max_percent=float(e.max()/120),
                   initial_energy_kwh=float(e[0]), final_energy_kwh=float(e[-1]),
                   simultaneous_charge_discharge_steps=0, first_interval="00:10-00:20",
                   last_interval="00:00+1-00:10+1", interval_count=144, state_count=145)
    audit = dict(status="PASS", scope="Delivery consistency and physical feasibility, no optimization",
                 time_convention="LEGACY_TIME_A_00:10_TO_NEXT_00:10", metrics=metrics,
                 numerical_tolerance_kwh=TOL, inputs=inputs, workbook=workbook_audit,
                 optimality_note="Historical solver claims preserved; no new dual certificate or optimality run.")
    write_json(OTHERS / "delivery_data_audit.json", audit)
    return d, e, blocks, balance, state, metrics


def tables(d, e, blocks, balance, state, metrics):
    common = d[["step", "internal_start", "internal_end"]].copy()
    inputs = pd.concat([common, d[["price_yuan_per_kwh", "load_kw", "pv_kw"]]], axis=1)
    distribution = common.assign(net_load_kw=d.load_kw-d.pv_kw)
    relation = common.assign(load_kw=d.load_kw, pv_kw=d.pv_kw)
    residuals = common.assign(balance_residual_kwh=balance, state_residual_kwh=state)
    activity = common.assign(energy_at_lower_bound=(abs(e[:-1]-1200) <= TOL).astype(int),
                             energy_at_upper_bound=(abs(e[:-1]-10800) <= TOL).astype(int),
                             charge_at_power_limit=(abs(d.charge_kwh-5000*DT) <= TOL).astype(int),
                             discharge_at_power_limit=(abs(d.discharge_kwh-5000*DT) <= TOL).astype(int))
    dispatch = common.assign(grid_purchase_kw=d.purchase_kwh/DT, net_load_kw=d.load_kw-d.pv_kw,
                             charge_kw=d.charge_kwh/DT, discharge_kw=d.discharge_kwh/DT)
    soc = pd.DataFrame(dict(boundary=np.arange(145), legacy_time=[clock(10+10*i) for i in range(145)],
                            energy_kwh=e, soc_percent=e/120))
    cost = pd.DataFrame(dict(scenario=["no_storage", "archived_storage_solution"],
                             cost_yuan=[metrics["no_storage_cost_yuan"], metrics["purchase_cost_yuan"]]))
    dfs = [inputs, distribution, relation, residuals, activity, blocks, dispatch, soc, cost]
    dest = RESULTS / "figure_data"
    dest.mkdir(parents=True, exist_ok=True)
    records = []
    for spec, frame in zip(SPECS, dfs):
        path = dest / f"{spec[0]}.csv"
        frame.to_csv(path, index=False, encoding="utf-8-sig")
        readback = pd.read_csv(path, float_precision="round_trip")
        pd.testing.assert_frame_equal(frame.reset_index(drop=True), readback, check_dtype=False,
                                      rtol=0, atol=1e-10)
        records.append(dict(figure=spec[0], rows=len(frame), columns=frame.columns.tolist(),
                            path=str(path.relative_to(BASE)), sha256=digest(path), roundtrip="PASS"))
    write_json(OTHERS / "figure_source_audit.json", records)
    profile = profile_data(d)
    profile["scope_note"] = "Single ordered 144-slot series; no independence, confidence intervals, or outlier removal assumed."
    profile["label_columns"] = ["internal_start", "internal_end", "input_time_label", "template_label"]
    write_json(OTHERS / "data_profile.json", profile)
    (OTHERS / "data_profile.txt").write_text(render_report(profile), encoding="utf-8")
    return dfs


def style():
    font = HERE / "vendor/fonts/NotoSansSC-Regular.ttf"
    fm.fontManager.addfont(str(font))
    config = setup_style(journal="general", lang="zh", use_sciplots=False)
    plt.rcParams.update({"figure.dpi": 100, "font.size": 9, "axes.labelsize": 9, "axes.titlesize": 10,
                         "xtick.labelsize": 8, "ytick.labelsize": 8, "legend.fontsize": 8,
                         "font.sans-serif": [config["cjk_font"], "DejaVu Sans"],
                         "axes.grid": False, "svg.hashsalt": "q1-baseline-delivery"})
    return dict(config, font_sha256=digest(font))


def time_axis(ax):
    ax.set_xlim(0, 144)
    ax.set_xticks([0, 24, 48, 72, 96, 120, 144],
                  ["00:10", "04:10", "08:10", "12:10", "16:10", "20:10", "00:10+1"])
    ax.set_xlabel("旧Time A序列时刻 (+1为次日)")


def step(ax, values, **kwargs):
    ax.stairs(np.asarray(values), np.arange(145), baseline=None, **kwargs)


def legend(ax, columns=2):
    ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0, 1.01),
              ncol=columns, borderaxespad=0)


def draw(index, frame, metrics):
    multi = index in (0, 6)
    size = (7.2, 4.8 if multi else 3.6)
    fig, axes = plt.subplots(2 if multi else 1, 1, figsize=size, sharex=multi)
    ax = axes[0] if multi else axes
    if index == 0:
        step(ax, frame.load_kw, color=COLORS["load"], label="负荷")
        step(ax, frame.pv_kw, color=COLORS["pv"], linestyle="--", label="光伏")
        ax.set_ylabel("区间平均功率 (kW)")
        ax.set_ylim(bottom=0)
        legend(ax)
        step(axes[1], frame.price_yuan_per_kwh, color=COLORS["reference"])
        axes[1].set_ylabel("电价 (元/kWh)")
        axes[1].set_ylim(bottom=0)
        time_axis(axes[1])
    elif index == 1:
        ax.hist(frame.net_load_kw, bins=16, color=COLORS["load"], edgecolor="white")
        ax.axvline(0, color=COLORS["reference"], linestyle="--", linewidth=1)
        ax.set(xlabel="净负荷 = 负荷 - 光伏 (kW)", ylabel="时段数", title="144段净负荷分布")
    elif index == 2:
        daylight = frame.pv_kw > 0
        for mask, color, marker, label in [(daylight, COLORS["pv"], "o", "有光伏时段"),
                                           (~daylight, COLORS["load"], "x", "零光伏时段")]:
            ax.scatter(frame.load_kw[mask], frame.pv_kw[mask], c=color, marker=marker,
                       s=20, alpha=0.7, label=label)
        limit = float(max(frame.load_kw.max(), frame.pv_kw.max()) * 1.05)
        ax.plot([0, limit], [0, limit], color=COLORS["reference"], linestyle="--", label="光伏 = 负荷")
        ax.set(xlabel="负荷 (kW)", ylabel="光伏 (kW)", xlim=(-100, limit), ylim=(-180, limit))
        legend(ax, 3)
    elif index == 3:
        ax.plot(frame.step, frame.balance_residual_kwh/1e-12,
                color=COLORS["balance"], label="能量守恒残差")
        ax.plot(frame.step, frame.state_residual_kwh/1e-12,
                color=COLORS["state"], linestyle="--", label="储能状态残差")
        ax.axhline(0, color="#999999", linewidth=0.6)
        ax.set_ylabel("重算残差 (10^-12 kWh)")
        time_axis(ax)
        legend(ax)
    elif index == 4:
        matrix = frame.iloc[:, 3:].to_numpy().T
        grid = ax.pcolormesh(np.arange(145), np.arange(5), matrix,
                            cmap=ListedColormap(["#F0F0F0", "#0072B2"]), vmin=0, vmax=1,
                            linewidth=0, rasterized=False)
        ax.set_yticks(np.arange(4)+0.5, ["储量下限", "储量上限", "充电功率上限", "放电功率上限"])
        ax.set_ylim(4, 0)
        cb = fig.colorbar(grid, ax=ax, ticks=[0, 1], shrink=0.65, pad=0.02)
        cb.ax.set_yticklabels(["未触界", "触界"])
        cb.set_label("绝对容差 10^-5 kWh")
        time_axis(ax)
    elif index == 5:
        x = np.arange(6)
        ax.bar(x-0.19, frame.charge_kwh, width=0.38, color=COLORS["charge"],
               edgecolor="#333333", linewidth=0.3, label="充电")
        ax.bar(x+0.19, frame.discharge_kwh, width=0.38, color=COLORS["discharge"],
               edgecolor="#333333", linewidth=0.3, hatch="//", label="放电")
        labels = [s.replace("-", "\n") for s in frame.actual_sequence_interval]
        ax.set_xticks(x, labels)
        ax.set(xlabel="旧Time A连续四小时段 (起点 / 终点)", ylabel="组内总电量 (kWh)")
        legend(ax)
    elif index == 6:
        step(ax, frame.grid_purchase_kw, color=COLORS["grid"], label="购电")
        step(ax, frame.net_load_kw, color=COLORS["reference"], linestyle="--", label="净负荷")
        ax.set_ylabel("区间平均功率 (kW)")
        ax.axhline(0, color="#999999", linewidth=0.5)
        legend(ax)
        step(axes[1], frame.charge_kw, color=COLORS["charge"], label="充电 (+)")
        step(axes[1], -frame.discharge_kw, color=COLORS["discharge"], linestyle="--", label="放电 (-)")
        axes[1].set_ylabel("储能功率 (kW)")
        axes[1].axhline(0, color="#999999", linewidth=0.5)
        axes[1].set_ylim(-5500, 5500)
        legend(axes[1])
        time_axis(axes[1])
    elif index == 7:
        ax.plot(frame.boundary, frame.soc_percent, color=COLORS["energy"], linewidth=1.4, label="SOC边界状态")
        ax.axhline(10, color=COLORS["reference"], linestyle=":", label="容量边界 (10%, 90%)")
        ax.axhline(90, color=COLORS["reference"], linestyle=":")
        ax.scatter([0, 144], [50, 50], marker="o", facecolor="white", edgecolor=COLORS["energy"],
                   zorder=3, s=24, clip_on=False, label="首末状态 50%")
        ax.set(ylabel="SOC (%)", ylim=(0, 100))
        time_axis(ax)
        legend(ax, 3)
    else:
        bars = ax.bar([0, 1], frame.cost_yuan, width=0.5,
                      color=[COLORS["reference"], COLORS["grid"]], edgecolor="#333333", linewidth=0.3)
        bars[0].set_hatch("//")
        ax.bar_label(bars, labels=[f"{v:,.2f}" for v in frame.cost_yuan], padding=4, fontsize=9)
        ax.set_xticks([0, 1], ["无储能参照", "现有储能方案"])
        ax.set(ylabel="序列总购电费 (元)", xlabel="同一输入下的确定性方案",
               ylim=(0, metrics["no_storage_cost_yuan"]*1.19),
               title=f"节约 {metrics['saving_yuan']:,.2f} 元 ({metrics['saving_fraction']:.2%})")
    if multi:
        for panel, label in zip(axes, "ab"):
            panel.set_title(label, loc="left", fontweight="normal", pad=25)
    return fig, size


def render(frames, metrics):
    config = style()
    FIGURES.mkdir(parents=True, exist_ok=True)
    previews = OTHERS / "visual_qa"
    previews.mkdir(parents=True, exist_ok=True)
    records = []
    contracts = []
    for i, (spec, frame) in enumerate(zip(SPECS, frames)):
        name, category, chart, claim, evidence, title = spec
        fig, size = draw(i, frame, metrics)
        issues = audit_layout(fig)
        assert not any(level == "FAIL" for level, _ in issues), (name, issues)
        exported = export_figure(fig, str(FIGURES / name), formats=["png", "pdf", "svg"],
                                  size_inches=size, dpi=300, tight=False, grayscale_preview=False)
        # The skill grayscale helper rewrites the main PNG with tight cropping;
        # derive QA copies separately to preserve the specified final dimensions.
        with Image.open(FIGURES / f"{name}.png") as picture:
            picture.convert("L").save(previews / f"{name}_grayscale.png", dpi=(300, 300))
        records.append(dict(figure=name, layout_issues=issues, size_inches=size,
                            formats=[Path(p).suffix for p in exported]))
        contracts.append(dict(id=name, category=category, question="q1", title=title,
                              chart_type=chart, central_claim=claim, evidence_chain=evidence,
                              layout="quantitative grid" if i in (0, 6) else "single quantitative panel",
                              backend="Python/matplotlib", size_inches=size, dpi=300,
                              export_formats=["png", "pdf", "svg"], editable_svg_text=True,
                              source_csv=f"results/figure_data/{name}.csv",
                              statistical_scope="单个确定性序列，不作统计推断；无误差棒、置信区间或显著性检验。",
                              time_scope="旧Time A：当天00:10至次日00:10；不代表新自然日口径。"))
        plt.close(fig)
    write_json(OTHERS / "figure_contracts.json", dict(style=config, figures=contracts))
    write_json(OTHERS / "figure_layout_audit.json", records)
    captions = ["# Q1图件索引与可直接引用的图注", "",
                "以下为已有Q1_BASELINE的后处理图件。时间均为旧Time A：当天00:10至次日00:10。",
                "每张图对应一个源CSV，完整格式为PNG(300 DPI)、PDF、SVG；PDF嵌入字体，SVG保留可编辑文本。",
                "三类各3张，共9张逻辑图。过程图是账本核查与调度摘要，不是求解器迭代记录。", ""]
    for record in contracts:
        captions.extend([f"## {record['id']}", "", f"**{record['title']}。** {record['evidence_chain']}",
                         record["statistical_scope"], "", f"源数据：`{record['source_csv']}`。", ""])
    (OTHERS / "FIGURE_CAPTIONS.md").write_text("\n".join(captions), encoding="utf-8")


def check_pdf_embedding(path):
    fonts = []
    for page in PdfReader(path).pages:
        resources = page["/Resources"].get_object()
        for name, ref in resources["/Font"].get_object().items():
            root = ref.get_object()
            assert root["/Subtype"] != "/Type3"
            descendants = root.get("/DescendantFonts", [ref])
            for descendant in descendants:
                font = descendant.get_object()
                assert font["/Subtype"] != "/Type3"
                descriptor = font["/FontDescriptor"].get_object()
                streams = [key for key in ("/FontFile", "/FontFile2", "/FontFile3") if key in descriptor]
                assert streams, f"Unembedded font in {path}: {name}"
                sizes = {key: len(descriptor[key].get_object().get_data()) for key in streams}
                assert min(sizes.values()) > 0
                fonts.append(dict(resource=str(name), root_type=str(root["/Subtype"]),
                                  descendant_type=str(font["/Subtype"]),
                                  base_font=str(root["/BaseFont"]), embedded_stream_bytes=sizes))
    assert fonts
    return dict(path=str(path.relative_to(BASE)), status="PASS", fonts=fonts)


def check_artifacts():
    records = []
    font_records = []
    for name, *_ in SPECS:
        csv = RESULTS / "figure_data" / f"{name}.csv"
        assert csv.exists()
        for ext in ("png", "pdf", "svg"):
            p = FIGURES / f"{name}.{ext}"
            issues, info = check_figure(str(p), min_dpi=300,
                                        target_inches=(7.2, 4.8 if name in (SPECS[0][0], SPECS[6][0]) else 3.6))
            assert not any(level == "FAIL" for level, _ in issues), (name, issues)
            if ext == "pdf":
                font_records.append(check_pdf_embedding(p))
            if ext == "png":
                with Image.open(p) as im:
                    pixels = np.asarray(im.convert("RGB"))
                    assert pixels.std() > 5 and np.mean(pixels.min(axis=2) < 245) > 0.01
            records.append(dict(path=str(p.relative_to(BASE)), issues=issues, metadata=info, sha256=digest(p)))
    category = audit_figure_directory(FIGURES, questions=("q1",))
    assert category["ok"], category["issues"]
    write_json(OTHERS / "pdf_font_embedding_audit.json", dict(status="PASS", files=font_records,
               method="Inspect each Type0 DescendantFonts FontDescriptor and require nonempty embedded font streams."))
    write_json(OTHERS / "figure_file_audit.json", dict(status="PASS", files=records,
               warning_resolution="The unchanged skill checker inspects only the root FontDescriptor, producing Type0 embedding warnings. All descendants are separately verified in pdf_font_embedding_audit.json."))
    write_json(OTHERS / "figure_category_audit.json", category)
    audit = json.loads((OTHERS / "figure_source_audit.json").read_text(encoding="utf-8"))
    for entry in audit:
        assert digest(BASE / entry["path"]) == entry["sha256"]


def manifest():
    excluded = {"others/delivery_manifest.json"}
    files = sorted(p for p in BASE.rglob("*") if p.is_file() and "__pycache__" not in p.parts
                   and str(p.relative_to(BASE)) not in excluded)
    write_json(OTHERS / "delivery_manifest.json", dict(
        status="PASS", kind="Postprocess-only archive supplement; not a model version",
        time_convention="LEGACY_TIME_A", solver_called=False, command=sys.argv, python=sys.version,
        platform=platform.platform(), packages=dict(numpy=np.__version__, pandas=pd.__version__,
                                                    matplotlib=matplotlib.__version__, openpyxl=openpyxl.__version__),
        files=snapshot(files), excluded=list(excluded),
        excluded_patterns=["**/__pycache__/**"], file_count=len(files)))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit-only", action="store_true", help="Check existing deliverables without drawing")
    args = parser.parse_args()
    before = snapshot(frozen_paths())
    archive()
    d, e, blocks, balance, state, metrics = read_and_check()
    if not args.audit_only:
        frames = tables(d, e, blocks, balance, state, metrics)
        render(frames, metrics)
    check_artifacts()
    assert before == snapshot(frozen_paths()), "A historical file changed during postprocessing"
    write_json(OTHERS / "historical_preservation_audit.json", dict(status="PASS", files=before))
    manifest()
    print(json.dumps(dict(status="PASS", solver_called=False, logical_figures=9,
                          exports=27, metrics=metrics), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
