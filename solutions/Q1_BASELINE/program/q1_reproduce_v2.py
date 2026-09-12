from __future__ import annotations

import hashlib
import argparse
import json
import platform
import shutil
import sys
import time
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import scipy
from scipy.optimize import linprog


ROOT = Path(__file__).resolve().parents[3]
INPUT_XLSX = ROOT / "data" / "附件1.xlsx"
TEMPLATE_XLSX = ROOT / "data" / "附件5" / "result1.xlsx"
ROUTE_SPEC = ROOT / "plan" / "C题融合路线与验证规约_v2.md"
OUTPUT_DIR = ROOT / "solutions" / "Q1_BASELINE" / "result"
OUTPUT_XLSX = OUTPUT_DIR / "result1_v2.xlsx"

INTERVALS = 144
DELTA_HOURS = 1.0 / 6.0
CAPACITY_KWH = 12000.0
ENERGY_MIN_KWH = 1200.0
ENERGY_MAX_KWH = 10800.0
INITIAL_ENERGY_KWH = 6000.0
FINAL_ENERGY_KWH = 6000.0
POWER_LIMIT_KW = 5000.0
STEP_LIMIT_KWH = POWER_LIMIT_KW * DELTA_HOURS
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
ABS_AUDIT_TOL_KWH = 1e-5
NORM_AUDIT_TOL = 1e-8
MUTEX_TOL_KWH = 1e-7


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clock(minutes: int) -> str:
    day = minutes // 1440
    value = minutes % 1440
    text = f"{value // 60:02d}:{value % 60:02d}"
    return text if day == 0 else f"{text}+{day}"


def load_q1_data():
    raw = pd.read_excel(INPUT_XLSX, sheet_name=0, header=0)
    if raw.shape != (INTERVALS, 4):
        raise AssertionError(f"附件1应为144行4列，实际为{raw.shape}")
    if raw.isna().any().any():
        raise AssertionError("附件1存在缺失值")
    numeric = raw.iloc[:, 1:4].apply(pd.to_numeric, errors="raise").to_numpy(float)
    price = numeric[:, 0]
    load_kw = numeric[:, 1]
    pv_kw = numeric[:, 2]
    if np.any(price < 0) or np.any(load_kw < 0) or np.any(pv_kw < 0):
        raise AssertionError("附件1出现负价格、负负荷或负光伏")
    return raw, price, load_kw, pv_kw


def solve_q1_lp(price, load_kw, pv_kw):
    load_kwh = np.asarray(load_kw) * DELTA_HOURS
    pv_kwh = np.asarray(pv_kw) * DELTA_HOURS
    count = INTERVALS
    offset_q = 0
    offset_c = count
    offset_d = 2 * count
    offset_w = 3 * count
    offset_e = 4 * count
    variables = 5 * count + 1

    objective = np.zeros(variables)
    objective[offset_q:offset_q + count] = price
    equalities = np.zeros((2 * count, variables))
    targets = np.zeros(2 * count)

    for step in range(count):
        balance_row = 2 * step
        equalities[balance_row, offset_q + step] = 1.0
        equalities[balance_row, offset_c + step] = -1.0
        equalities[balance_row, offset_d + step] = 1.0
        equalities[balance_row, offset_w + step] = -1.0
        targets[balance_row] = load_kwh[step] - pv_kwh[step]

        state_row = balance_row + 1
        equalities[state_row, offset_e + step] = -1.0
        equalities[state_row, offset_e + step + 1] = 1.0
        equalities[state_row, offset_c + step] = -ETA_CHARGE
        equalities[state_row, offset_d + step] = 1.0 / ETA_DISCHARGE

    bounds = [(0.0, None)] * count
    bounds += [(0.0, STEP_LIMIT_KWH)] * count
    bounds += [(0.0, STEP_LIMIT_KWH)] * count
    bounds += [(0.0, None)] * count
    bounds += [(ENERGY_MIN_KWH, ENERGY_MAX_KWH)] * (count + 1)
    bounds[offset_e] = (INITIAL_ENERGY_KWH, INITIAL_ENERGY_KWH)
    bounds[offset_e + count] = (FINAL_ENERGY_KWH, FINAL_ENERGY_KWH)

    started = time.perf_counter()
    result = linprog(
        objective,
        A_eq=equalities,
        b_eq=targets,
        bounds=bounds,
        method="highs",
    )
    elapsed = time.perf_counter() - started
    if not result.success:
        raise RuntimeError(f"Q1 LP求解失败: {result.message}")

    values = result.x
    solution = {
        "purchase_kwh": values[offset_q:offset_q + count],
        "charge_kwh": values[offset_c:offset_c + count],
        "discharge_kwh": values[offset_d:offset_d + count],
        "unused_supply_kwh": values[offset_w:offset_w + count],
        "energy_kwh": values[offset_e:offset_e + count + 1],
        "solver_objective_yuan": float(result.fun),
        "solve_seconds": float(elapsed),
        "solver_status": int(result.status),
        "solver_message": result.message,
    }
    return solution


def audit_solution(price, load_kw, pv_kw, solution):
    purchase = solution["purchase_kwh"]
    charge = solution["charge_kwh"]
    discharge = solution["discharge_kwh"]
    unused = solution["unused_supply_kwh"]
    energy = solution["energy_kwh"]
    load_kwh = np.asarray(load_kw) * DELTA_HOURS
    pv_kwh = np.asarray(pv_kw) * DELTA_HOURS

    balance = purchase + pv_kwh + discharge - load_kwh - charge - unused
    state = energy[1:] - energy[:-1] - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
    scale = max(1.0, float(np.max(load_kwh)), float(np.max(pv_kwh)))
    simultaneous = np.flatnonzero((charge > MUTEX_TOL_KWH) & (discharge > MUTEX_TOL_KWH))
    cash = float(np.dot(price, purchase))
    no_storage_purchase = np.maximum(load_kwh - pv_kwh, 0.0)
    no_storage_cash = float(np.dot(price, no_storage_purchase))

    audit = {
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_state_residual_kwh": float(np.max(np.abs(state))),
        "normalized_max_residual": float(max(np.max(np.abs(balance)), np.max(np.abs(state))) / scale),
        "minimum_flow_kwh": float(min(np.min(purchase), np.min(charge), np.min(discharge), np.min(unused))),
        "minimum_energy_kwh": float(np.min(energy)),
        "maximum_energy_kwh": float(np.max(energy)),
        "maximum_charge_kwh": float(np.max(charge)),
        "maximum_discharge_kwh": float(np.max(discharge)),
        "simultaneous_charge_discharge_steps": simultaneous.astype(int).tolist(),
        "initial_energy_kwh": float(energy[0]),
        "final_energy_kwh": float(energy[-1]),
        "cash_recalculated_yuan": cash,
        "objective_cash_difference_yuan": float(abs(cash - solution["solver_objective_yuan"])),
        "no_storage_cash_yuan": no_storage_cash,
        "saving_vs_no_storage_yuan": float(no_storage_cash - cash),
        "saving_vs_no_storage_fraction": float(1.0 - cash / no_storage_cash),
    }

    checks = {name: bool(value) for name, value in {
        "time_count_144": len(purchase) == INTERVALS,
        "state_count_145": len(energy) == INTERVALS + 1,
        "absolute_residual_pass": max(audit["max_balance_residual_kwh"], audit["max_state_residual_kwh"]) <= ABS_AUDIT_TOL_KWH,
        "normalized_residual_pass": audit["normalized_max_residual"] <= NORM_AUDIT_TOL,
        "nonnegative_pass": audit["minimum_flow_kwh"] >= -ABS_AUDIT_TOL_KWH,
        "energy_bounds_pass": audit["minimum_energy_kwh"] >= ENERGY_MIN_KWH - ABS_AUDIT_TOL_KWH and audit["maximum_energy_kwh"] <= ENERGY_MAX_KWH + ABS_AUDIT_TOL_KWH,
        "power_bounds_pass": audit["maximum_charge_kwh"] <= STEP_LIMIT_KWH + ABS_AUDIT_TOL_KWH and audit["maximum_discharge_kwh"] <= STEP_LIMIT_KWH + ABS_AUDIT_TOL_KWH,
        "mutual_exclusion_pass": len(simultaneous) == 0,
        "endpoint_pass": abs(energy[0] - INITIAL_ENERGY_KWH) <= ABS_AUDIT_TOL_KWH and abs(energy[-1] - FINAL_ENERGY_KWH) <= ABS_AUDIT_TOL_KWH,
        "cash_recalculation_pass": audit["objective_cash_difference_yuan"] <= 1e-7,
    }.items()}
    if not all(checks.values()):
        failed = [name for name, passed in checks.items() if not passed]
        raise AssertionError(f"Q1硬审计失败: {failed}")

    milp_certificate = {
        "method": "LP-relaxation lower bound plus an LP optimum that already satisfies every binary mutual-exclusion constraint",
        "lp_lower_bound_yuan": cash,
        "milp_feasible_solution_yuan": cash,
        "certified_milp_optimum_yuan": cash,
        "proof": "MILP feasible set is a subset of the LP relaxation; the LP optimum is mutually exclusive and therefore MILP-feasible, so both bounds coincide.",
        "instance_only": True,
    }
    return audit, checks, milp_certificate


def build_time_mapping(input_frame, template_labels, time_mapping):
    mapping = []
    for step in range(INTERVALS):
        source_step = step if time_mapping == "A" else (step + 1) % INTERVALS
        mapping.append({
            "step": step,
            "source_step": source_step,
            "input_time_label": str(input_frame.iloc[source_step, 0]),
            "internal_start": clock(10 + step * 10),
            "internal_end": clock(20 + step * 10),
            "template_label": str(template_labels[step]),
        })
    if mapping[0]["internal_start"] != "00:10" or mapping[0]["internal_end"] != "00:20":
        raise AssertionError("首时段内部映射错误")
    if mapping[-1]["internal_start"] != "00:00+1" or mapping[-1]["internal_end"] != "00:10+1":
        raise AssertionError("末时段内部映射错误")
    if time_mapping == "A" and not mapping[0]["input_time_label"].startswith("00:10"):
        raise AssertionError("附件1首项不是00:10，无法对齐首段00:10-00:20")
    if time_mapping == "A" and mapping[-1]["input_time_label"] not in {"0:00+1", "00:00+1"}:
        raise AssertionError("附件1末项不是0:00+1，无法解释为次日00:00-00:10")
    if time_mapping == "B" and (mapping[0]["source_step"] != 1 or mapping[-1]["source_step"] != 0):
        raise AssertionError("Q1 B映射必须按典型日周期左移一槽")
    if str(template_labels[0]) != "0:10-0:20":
        raise AssertionError("result1首行标签不是0:10-0:20")
    if str(template_labels[-1]) != "0:00+1-0:10+1":
        raise AssertionError("result1末行未明确标记为次日0:00-0:10")
    return pd.DataFrame(mapping)


def write_outputs(input_frame, price, load_kw, pv_kw, solution, audit, checks, milp_certificate,
                  time_mapping):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(TEMPLATE_XLSX)
    plan_sheet = workbook.worksheets[0]
    storage_sheet = workbook.worksheets[1]
    if plan_sheet.max_row != 145 or plan_sheet.max_column != 2:
        raise AssertionError("result1计划工作表结构变化")
    if storage_sheet.max_row != 7 or storage_sheet.max_column != 5:
        raise AssertionError("result1储能工作表结构变化")

    template_labels = [plan_sheet.cell(row=row, column=1).value for row in range(2, 146)]
    time_mapping_frame = build_time_mapping(input_frame, template_labels, time_mapping)
    purchase = solution["purchase_kwh"]
    charge = solution["charge_kwh"]
    discharge = solution["discharge_kwh"]
    energy = solution["energy_kwh"]
    unused = solution["unused_supply_kwh"]

    for step in range(INTERVALS):
        plan_sheet.cell(row=step + 2, column=2, value=float(purchase[step]))
    block_rows = []
    for block in range(6):
        start = block * 24
        end = start + 24
        charge_total = float(np.sum(charge[start:end]))
        discharge_total = float(np.sum(discharge[start:end]))
        storage_sheet.cell(row=block + 2, column=2, value=charge_total)
        storage_sheet.cell(row=block + 2, column=3, value=discharge_total)
        block_rows.append({
            "template_block": f"{4 * block:02d}:00-{4 * (block + 1):02d}:00",
            "actual_sequence_interval": f"{clock(10 + start * 10)}-{clock(10 + end * 10)}",
            "charge_kwh": charge_total,
            "discharge_kwh": discharge_total,
        })
    storage_sheet.cell(row=2, column=5, value=float(energy[0]))
    storage_sheet.cell(row=3, column=5, value=float(energy[-1]))
    workbook.save(OUTPUT_XLSX)

    dispatch = pd.DataFrame({
        "step": np.arange(INTERVALS),
        "internal_start": time_mapping_frame["internal_start"],
        "internal_end": time_mapping_frame["internal_end"],
        "input_time_label": time_mapping_frame["input_time_label"],
        "template_label": time_mapping_frame["template_label"],
        "price_yuan_per_kwh": price,
        "load_kw": load_kw,
        "pv_kw": pv_kw,
        "load_kwh": np.asarray(load_kw) * DELTA_HOURS,
        "pv_kwh": np.asarray(pv_kw) * DELTA_HOURS,
        "purchase_kwh": purchase,
        "charge_kwh": charge,
        "discharge_kwh": discharge,
        "unused_supply_kwh": unused,
        "energy_before_kwh": energy[:-1],
        "energy_after_kwh": energy[1:],
        "purchase_cost_yuan": price * purchase,
    })
    dispatch.to_csv(OUTPUT_DIR / "q1_dispatch.csv", index=False, encoding="utf-8-sig")
    time_mapping_frame.to_csv(OUTPUT_DIR / "q1_time_mapping.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(block_rows).to_csv(OUTPUT_DIR / "q1_four_hour_blocks.csv", index=False, encoding="utf-8-sig")

    requested = ["10:00", "12:00", "14:00", "16:00", "18:00", "20:00"]
    specified_rows = []
    for start_text in requested:
        matched = dispatch.loc[dispatch["internal_start"] == start_text]
        if len(matched) != 1:
            raise AssertionError(f"指定时段{start_text}映射数量不是1")
        row = matched.iloc[0]
        specified_rows.append({
            "internal_interval": f"{row['internal_start']}-{row['internal_end']}",
            "purchase_kwh": float(row["purchase_kwh"]),
        })
    specified = pd.DataFrame(specified_rows)
    specified.to_csv(OUTPUT_DIR / "q1_specified_intervals.csv", index=False, encoding="utf-8-sig")

    reopened = openpyxl.load_workbook(OUTPUT_XLSX, data_only=False)
    reopened_plan = reopened.worksheets[0]
    reopened_storage = reopened.worksheets[1]
    plan_readback = np.array([reopened_plan.cell(row=row, column=2).value for row in range(2, 146)], dtype=float)
    charge_readback = np.array([reopened_storage.cell(row=row, column=2).value for row in range(2, 8)], dtype=float)
    discharge_readback = np.array([reopened_storage.cell(row=row, column=3).value for row in range(2, 8)], dtype=float)
    template_readback = {
        "sheet_names_preserved": reopened.sheetnames == workbook.sheetnames,
        "plan_rows": int(len(plan_readback)),
        "plan_max_abs_difference_kwh": float(np.max(np.abs(plan_readback - purchase))),
        "plan_total_difference_kwh": float(abs(np.sum(plan_readback) - np.sum(purchase))),
        "cash_difference_yuan": float(abs(np.dot(price, plan_readback) - audit["cash_recalculated_yuan"])),
        "charge_blocks_max_abs_difference_kwh": float(np.max(np.abs(charge_readback - np.array([row["charge_kwh"] for row in block_rows])))),
        "discharge_blocks_max_abs_difference_kwh": float(np.max(np.abs(discharge_readback - np.array([row["discharge_kwh"] for row in block_rows])))),
        "initial_energy_difference_kwh": float(abs(reopened_storage.cell(row=2, column=5).value - energy[0])),
        "final_energy_difference_kwh": float(abs(reopened_storage.cell(row=3, column=5).value - energy[-1])),
    }
    template_readback["pass"] = bool(
        template_readback["sheet_names_preserved"]
        and template_readback["plan_rows"] == INTERVALS
        and max(value for key, value in template_readback.items() if key.endswith(("_kwh", "_yuan"))) <= ABS_AUDIT_TOL_KWH
    )
    if not template_readback["pass"]:
        raise AssertionError(f"result1模板回读失败: {template_readback}")

    manifest = {
        "route": "题目分析报告.md v1.1 governs current status; the v2 route supplies the archived mathematical contract.",
        "inputs": {
            str(INPUT_XLSX.relative_to(ROOT)): sha256(INPUT_XLSX),
            str(TEMPLATE_XLSX.relative_to(ROOT)): sha256(TEMPLATE_XLSX),
            str(ROUTE_SPEC.relative_to(ROOT)): sha256(ROUTE_SPEC),
        },
        "outputs": {
            str(OUTPUT_XLSX.relative_to(ROOT)): sha256(OUTPUT_XLSX),
        },
        "runtime": {
            "python": sys.version,
            "platform": platform.platform(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
        },
        "parameters": {
            "delta_hours": DELTA_HOURS,
            "capacity_kwh": CAPACITY_KWH,
            "energy_min_kwh": ENERGY_MIN_KWH,
            "energy_max_kwh": ENERGY_MAX_KWH,
            "initial_energy_kwh": INITIAL_ENERGY_KWH,
            "final_energy_kwh": FINAL_ENERGY_KWH,
            "power_limit_kw": POWER_LIMIT_KW,
            "step_limit_kwh": STEP_LIMIT_KWH,
            "eta_charge": ETA_CHARGE,
            "eta_discharge": ETA_DISCHARGE,
            "round_trip_efficiency": ETA_CHARGE * ETA_DISCHARGE,
            "time_mapping": time_mapping,
            "mapping_b_terminal_rule": (
                "Q1 fixed typical day is assumed periodic, so the final slot uses source row 0."
                if time_mapping == "B" else None
            ),
        },
        "contracts": {
            "q1_state_rule": "single-day closed cycle: E0=E144=6000 kWh",
            "q2_q4_state_rule": "continuous cross-day rolling state; never reset to 6000 each day",
            "time_rule": "first interval 00:10-00:20; last interval next-day 00:00-00:10",
            "time_mapping": time_mapping,
        },
        "command": (
            "python solutions/Q1_BASELINE/program/q1_reproduce_v2.py "
            f"--time-mapping {time_mapping} --output-dir {OUTPUT_DIR}"
        ),
    }
    (OUTPUT_DIR / "q1_input_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    summary = {
        "status": "PASS",
        "evidence_scope": "Q1 only; no Q2-Q4 annual result is claimed",
        "solver": {
            "lp_objective_yuan": solution["solver_objective_yuan"],
            "solve_seconds": solution["solve_seconds"],
            "status": solution["solver_status"],
            "message": solution["solver_message"],
        },
        "result": {
            "total_purchase_kwh": float(np.sum(purchase)),
            "total_purchase_cost_yuan": audit["cash_recalculated_yuan"],
            "no_storage_cost_yuan": audit["no_storage_cash_yuan"],
            "saving_fraction": audit["saving_vs_no_storage_fraction"],
            "specified_intervals": specified_rows,
            "four_hour_blocks": block_rows,
            "initial_energy_kwh": float(energy[0]),
            "final_energy_kwh": float(energy[-1]),
            "charge_efficiency": ETA_CHARGE,
            "discharge_efficiency": ETA_DISCHARGE,
            "round_trip_efficiency": ETA_CHARGE * ETA_DISCHARGE,
        },
        "audit": audit,
        "checks": checks,
        "milp_optimality_certificate": milp_certificate,
        "template_readback": template_readback,
        "manifest": manifest,
    }
    (OUTPUT_DIR / "q1_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main():
    global OUTPUT_DIR, OUTPUT_XLSX
    parser = argparse.ArgumentParser(description="Reproduce the archived Q1 numerical baseline")
    parser.add_argument(
        "--output-dir",
        default=str(OUTPUT_DIR),
        help="New or archived output directory; use a new runs/ path for independent checks.",
    )
    parser.add_argument("--time-mapping", choices=["A", "B"], default="A")
    args = parser.parse_args()
    OUTPUT_DIR = Path(args.output_dir).resolve()
    OUTPUT_XLSX = OUTPUT_DIR / "result1_v2.xlsx"
    input_frame, price, load_kw, pv_kw = load_q1_data()
    if args.time_mapping == "B":
        # Q1 is a fixed typical day, so the right-endpoint sensitivity is periodic.
        price, load_kw, pv_kw = (np.roll(values, -1) for values in (price, load_kw, pv_kw))
    solution = solve_q1_lp(price, load_kw, pv_kw)
    audit, checks, milp_certificate = audit_solution(price, load_kw, pv_kw, solution)
    summary = write_outputs(
        input_frame, price, load_kw, pv_kw, solution, audit, checks,
        milp_certificate, args.time_mapping,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
