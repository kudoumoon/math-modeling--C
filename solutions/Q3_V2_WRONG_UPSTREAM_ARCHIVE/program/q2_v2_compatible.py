"""Full Q2 causal rolling backtest under the frozen V2 contract.

The script first selects a simple forecast/alpha using only completed January
plan days, then runs B0, M0_audited and R1 over the official 334-plan-day
window.  R2's next-day piecewise-linear terminal value is tested separately
and adopted only when its predeclared gate passes.  The official result2
template is populated only with the accepted executable strategy.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import platform
import shutil
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
import pandas as pd
import scipy
from scipy import sparse
from scipy.optimize import linprog

import q2_pipeline_v1 as core


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "Q2" / "result"
TEMPLATE_PATH = core.TEMPLATE_PATH
N_STEPS = core.N_STEPS
DT_HOURS = core.DT_HOURS
SOC_MIN = core.SOC_MIN_KWH
SOC_MAX = core.SOC_MAX_KWH
INITIAL_SOC = core.INITIAL_SOC_KWH
STEP_LIMIT = core.STEP_LIMIT_KWH
ETA_C = core.ETA_CHARGE
ETA_D = core.ETA_DISCHARGE
ABS_TOL = core.ABS_TOL_KWH
NORM_TOL = core.NORM_TOL
EPS_THROUGHPUT = core.EPS_THROUGHPUT
OUTPUT_START_INDEX = 31  # 2025-02-01
OUTPUT_END_INDEX = 364   # 2025-12-31, inclusive
JAN_SELECTION_LAST_INDEX = 29  # 2025-01-30; plan ends Jan 31 00:10
ALPHAS = (0.65, 0.70, 0.75, 0.80, 0.85)
LOAD_MODELS = core.LOAD_CANDIDATE_ORDER
PV_MODELS = core.PV_CANDIDATE_ORDER
R2_GRID = np.array([1200, 2800, 4400, 6000, 7600, 9200, 10800], dtype=float)
BOOTSTRAP_SEED = 20260911
BOOTSTRAP_REPS = 2000
R2_MIN_RELATIVE_SAVING = 0.002
R2_ANCHOR_SENSITIVITY = (4400.0, 7600.0)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


@dataclass(frozen=True)
class StrategyConfig:
    name: str
    load_model: str
    pv_model: str
    alpha: Optional[float]
    terminal_mode: str = "soft"

    @property
    def combo(self) -> Tuple[str, str]:
        return self.load_model, self.pv_model


@dataclass
class Plan:
    day_index: int
    decision_time: pd.Timestamp
    predicted_start_soc: float
    forecast: core.Forecast
    margin: np.ndarray
    residual_days: Tuple[int, ...]
    residual_max_timestamp: Optional[pd.Timestamp]
    q: np.ndarray
    planned_charge: np.ndarray
    planned_discharge: np.ndarray
    planned_soc: np.ndarray
    planning_cash: float
    planning_objective: float
    max_lp_residual: float
    solve_seconds: float
    config: StrategyConfig
    terminal_curve: Optional[Dict[str, object]]


def _observed_value(
    data: np.ndarray,
    dataset: core.Dataset,
    source_day: int,
    step: int,
    issue_time: pd.Timestamp,
) -> Optional[float]:
    if source_day < 0 or source_day >= len(dataset.dates):
        return None
    end = dataset.interval_starts(source_day)[step] + pd.Timedelta(minutes=10)
    if end <= issue_time:
        return float(data[source_day, step])
    return None


def _series_forecast(
    dataset: core.Dataset,
    target_day: int,
    issue_time: pd.Timestamp,
    kind: str,
    model: str,
) -> Tuple[np.ndarray, Optional[pd.Timestamp], str]:
    data = dataset.load_kwh if kind == "load" else dataset.pv_kwh
    cold = dataset.cold_load_kwh if kind == "load" else dataset.cold_pv_kwh
    values = np.zeros(N_STEPS)
    used_ends: List[pd.Timestamp] = []

    for step in range(N_STEPS):
        chosen: List[Tuple[int, float]] = []
        if model == "lag7":
            offsets: Iterable[int] = (7, 14, 21, 28)
            needed = 1
        elif model == "lag7_lag14_mean":
            offsets = (7, 14)
            needed = 2
        elif model == "same_week_up_to4_mean":
            offsets = (7, 14, 21, 28)
            needed = 4
        elif model == "same_weekday_2w":
            # Q2 v2 point forecast: mean of the two most recent same-weekday
            # observations.  The causal timestamp filter below remains active.
            offsets = (7, 14)
            needed = 2
        elif model == "lag1":
            offsets = range(1, 32)
            needed = 1
        elif model == "recent3_mean":
            offsets = range(1, 32)
            needed = 3
        elif model == "recent_5d":
            # Q2 v2 point forecast: mean of the five most recent completed
            # days, filtered interval-by-interval by the issue timestamp.
            offsets = range(1, 32)
            needed = 5
        else:
            raise ValueError(f"未知{kind}预测器: {model}")

        for offset in offsets:
            source = target_day - int(offset)
            value = _observed_value(data, dataset, source, step, issue_time)
            if value is not None:
                chosen.append((source, value))
                if len(chosen) >= needed:
                    break
        if chosen:
            values[step] = float(np.mean([value for _, value in chosen]))
            for source, _ in chosen:
                used_ends.append(dataset.interval_starts(source)[step] + pd.Timedelta(minutes=10))
        else:
            values[step] = float(cold[step])

    source_max = max(used_ends) if used_ends else None
    return values, source_max, f"{kind}:{model};issue={issue_time.isoformat()}"


def make_forecast(
    dataset: core.Dataset,
    target_day: int,
    issue_time: pd.Timestamp,
    combo: Tuple[str, str],
) -> core.Forecast:
    load, load_max, load_rule = _series_forecast(dataset, target_day, issue_time, "load", combo[0])
    pv, pv_max, pv_rule = _series_forecast(dataset, target_day, issue_time, "pv", combo[1])
    maxima = [x for x in (load_max, pv_max) if x is not None]
    return core.Forecast(
        load_kwh=load,
        pv_kwh=pv,
        source_max_timestamp=max(maxima) if maxima else None,
        source_rule=f"{load_rule}|{pv_rule}",
    )


def build_forecast_library(
    dataset: core.Dataset,
    combos: Sequence[Tuple[str, str]],
) -> Dict[Tuple[str, str], Dict[int, core.Forecast]]:
    library: Dict[Tuple[str, str], Dict[int, core.Forecast]] = {}
    for combo in combos:
        library[combo] = {
            day: make_forecast(dataset, day, dataset.dates[day], combo)
            for day in range(len(dataset.dates))
        }
    return library


def causal_margin(
    dataset: core.Dataset,
    issue_time: pd.Timestamp,
    forecasts: Dict[int, core.Forecast],
    alpha: Optional[float],
) -> Tuple[np.ndarray, Tuple[int, ...], Optional[pd.Timestamp]]:
    completed = [
        day for day in sorted(forecasts)
        if dataset.plan_end(day) <= issue_time
    ][-core.MAX_RESIDUAL_DAYS:]
    max_time = dataset.plan_end(completed[-1]) if completed else None
    if alpha is None or len(completed) < core.MIN_RESIDUAL_DAYS:
        return np.zeros(N_STEPS), tuple(completed), max_time
    residuals = np.vstack([
        (dataset.load_kwh[day] - dataset.pv_kwh[day]) - forecasts[day].net_kwh
        for day in completed
    ])
    return np.quantile(residuals, alpha, axis=0), tuple(completed), max_time


def solve_plan(
    risk_net: np.ndarray,
    price: np.ndarray,
    start_soc: float,
    terminal_mode: str,
    hard_anchor: Optional[float] = None,
    pwl_curve: Optional[Tuple[np.ndarray, np.ndarray]] = None,
    kappa_multiplier: float = 1.0,
) -> Dict[str, object]:
    n = N_STEPS
    oq, oc, od, ow = 0, n, 2 * n, 3 * n
    oe = 4 * n
    extra = oe + n + 1
    if terminal_mode == "soft":
        oup, odn = extra, extra + 1
        nvar = extra + 2
    elif terminal_mode == "pwl":
        oz = extra
        nvar = extra + 1
    elif terminal_mode in {"none", "hard"}:
        nvar = extra
    else:
        raise ValueError(terminal_mode)

    objective = np.zeros(nvar)
    objective[oq:oq + n] = price
    objective[oc:oc + n] = EPS_THROUGHPUT
    objective[od:od + n] = EPS_THROUGHPUT
    if terminal_mode == "soft":
        kappa = kappa_multiplier * ETA_D * float(np.median(price))
        objective[oup] = kappa
        objective[odn] = kappa
    elif terminal_mode == "pwl":
        objective[oz] = 1.0

    eq_rows = 2 * n + (1 if terminal_mode == "soft" else 0)
    aeq = sparse.lil_matrix((eq_rows, nvar))
    beq = np.zeros(eq_rows)
    for t in range(n):
        aeq[t, oq + t] = 1.0
        aeq[t, oc + t] = -1.0
        aeq[t, od + t] = 1.0
        aeq[t, ow + t] = -1.0
        beq[t] = float(risk_net[t])
        row = n + t
        aeq[row, oe + t] = -1.0
        aeq[row, oe + t + 1] = 1.0
        aeq[row, oc + t] = -ETA_C
        aeq[row, od + t] = 1.0 / ETA_D
    if terminal_mode == "soft":
        aeq[2 * n, oe + n] = 1.0
        aeq[2 * n, oup] = -1.0
        aeq[2 * n, odn] = 1.0
        beq[2 * n] = start_soc

    aub = None
    bub = None
    if terminal_mode == "pwl":
        if pwl_curve is None:
            raise ValueError("PWL模式缺少价值曲线")
        grid, values = pwl_curve
        slopes = np.diff(values) / np.diff(grid)
        if np.max(slopes) > 1e-7 or np.min(np.diff(slopes)) < -1e-7:
            raise AssertionError("次日价值曲线不是单调非增凸函数")
        intercepts = values[:-1] - slopes * grid[:-1]
        aub = sparse.lil_matrix((len(slopes), nvar))
        bub = np.zeros(len(slopes))
        for i, (slope, intercept) in enumerate(zip(slopes, intercepts)):
            aub[i, oe + n] = float(slope)
            aub[i, oz] = -1.0
            bub[i] = -float(intercept)

    bounds: List[Tuple[Optional[float], Optional[float]]] = []
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(0.0, STEP_LIMIT)] * n)
    bounds.extend([(0.0, STEP_LIMIT)] * n)
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(SOC_MIN, SOC_MAX)] * (n + 1))
    bounds[oe] = (start_soc, start_soc)
    if terminal_mode == "soft":
        bounds.extend([(0.0, None), (0.0, None)])
    elif terminal_mode == "pwl":
        bounds.append((0.0, None))
    if terminal_mode == "hard":
        if hard_anchor is None:
            raise ValueError("hard模式缺少anchor")
        bounds[oe + n] = (hard_anchor, hard_anchor)

    started = time.perf_counter()
    result = linprog(
        objective,
        A_ub=None if aub is None else aub.tocsr(),
        b_ub=bub,
        A_eq=aeq.tocsr(),
        b_eq=beq,
        bounds=bounds,
        method="highs",
    )
    elapsed = time.perf_counter() - started
    if not result.success:
        raise RuntimeError(f"LP失败: {result.status} {result.message}")
    x = np.asarray(result.x)
    eq_residual = np.asarray(aeq.tocsr() @ x - beq)
    ub_residual = np.array([0.0]) if aub is None else np.maximum(np.asarray(aub.tocsr() @ x - bub), 0.0)
    return {
        "q": x[oq:oq + n],
        "charge": x[oc:oc + n],
        "discharge": x[od:od + n],
        "soc": x[oe:oe + n + 1],
        "cash": float(np.dot(price, x[oq:oq + n])),
        "objective": float(result.fun),
        "max_residual": float(max(np.max(np.abs(eq_residual)), np.max(ub_residual))),
        "solve_seconds": elapsed,
    }


def next_day_value_curve(
    dataset: core.Dataset,
    day_index: int,
    config: StrategyConfig,
    forecast_library: Dict[Tuple[str, str], Dict[int, core.Forecast]],
    anchor: float = 6000.0,
) -> Dict[str, object]:
    if day_index + 1 >= len(dataset.dates):
        raise ValueError("最后一天没有附件内次日")
    issue = dataset.dates[day_index]
    future_forecast = make_forecast(dataset, day_index + 1, issue, config.combo)
    margin, residual_days, residual_max = causal_margin(
        dataset, issue, forecast_library[config.combo], config.alpha
    )
    future_risk = future_forecast.net_kwh + margin
    values = []
    for start in R2_GRID:
        solved = solve_plan(
            future_risk,
            dataset.price,
            float(start),
            terminal_mode="hard",
            hard_anchor=anchor,
        )
        values.append(float(solved["cash"]))
    values_array = np.asarray(values)
    slopes = np.diff(values_array) / np.diff(R2_GRID)
    monotone = bool(np.max(slopes) <= 1e-7)
    convex = bool(np.min(np.diff(slopes)) >= -1e-7)
    return {
        "grid": R2_GRID.copy(),
        "values": values_array,
        "slopes": slopes,
        "monotone": monotone,
        "convex": convex,
        "forecast_source_max_timestamp": future_forecast.source_max_timestamp,
        "residual_source_max_timestamp": residual_max,
        "residual_days": residual_days,
        "anchor": anchor,
    }


def make_plan(
    dataset: core.Dataset,
    day_index: int,
    soc_midnight: float,
    previous: Optional[Plan],
    config: StrategyConfig,
    forecast_library: Dict[Tuple[str, str], Dict[int, core.Forecast]],
    r2_anchor: float = 6000.0,
) -> Plan:
    issue = dataset.dates[day_index]
    if previous is None:
        predicted_start = INITIAL_SOC
    else:
        bridge = core.feedback_step(
            soc_midnight,
            float(previous.q[-1]),
            float(previous.forecast.load_kwh[-1]),
            float(previous.forecast.pv_kwh[-1]),
        )
        predicted_start = float(bridge["soc_after_kwh"])
    forecast = forecast_library[config.combo][day_index]
    margin, residual_days, residual_max = causal_margin(
        dataset, issue, forecast_library[config.combo], config.alpha
    )
    risk_net = forecast.net_kwh + margin
    terminal_curve = None
    if day_index == OUTPUT_END_INDEX:
        terminal_mode = "hard"
        solved = solve_plan(risk_net, dataset.price, predicted_start, terminal_mode, hard_anchor=6000.0)
    elif config.terminal_mode == "hard":
        solved = solve_plan(risk_net, dataset.price, predicted_start, "hard", hard_anchor=6000.0)
    elif config.terminal_mode == "pwl":
        terminal_curve = next_day_value_curve(
            dataset, day_index, config, forecast_library, anchor=r2_anchor
        )
        if not terminal_curve["monotone"] or not terminal_curve["convex"]:
            raise AssertionError(f"{dataset.dates[day_index].date()} R2价值曲线形状失败")
        solved = solve_plan(
            risk_net,
            dataset.price,
            predicted_start,
            "pwl",
            pwl_curve=(terminal_curve["grid"], terminal_curve["values"]),
        )
    else:
        solved = solve_plan(risk_net, dataset.price, predicted_start, config.terminal_mode)
    return Plan(
        day_index=day_index,
        decision_time=issue,
        predicted_start_soc=predicted_start,
        forecast=forecast,
        margin=margin,
        residual_days=residual_days,
        residual_max_timestamp=residual_max,
        q=np.asarray(solved["q"]),
        planned_charge=np.asarray(solved["charge"]),
        planned_discharge=np.asarray(solved["discharge"]),
        planned_soc=np.asarray(solved["soc"]),
        planning_cash=float(solved["cash"]),
        planning_objective=float(solved["objective"]),
        max_lp_residual=float(solved["max_residual"]),
        solve_seconds=float(solved["solve_seconds"]),
        config=config,
        terminal_curve=terminal_curve,
    )


def execute_step(dataset: core.Dataset, plan: Plan, step: int, soc: float) -> Dict[str, object]:
    start = dataset.interval_starts(plan.day_index)[step]
    q = float(plan.q[step])
    load = float(dataset.load_kwh[plan.day_index, step])
    pv = float(dataset.pv_kwh[plan.day_index, step])
    action = core.feedback_step(soc, q, load, pv)
    price = float(dataset.price[step])
    return {
        "interval_start": start,
        "interval_end": start + pd.Timedelta(minutes=10),
        "decision_time": plan.decision_time,
        "plan_day": dataset.dates[plan.day_index],
        "natural_day": start.normalize(),
        "step": step,
        "template_label": dataset.template_labels[step],
        "input_time_label": dataset.input_time_labels[step],
        "strategy": plan.config.name,
        "load_model": plan.config.load_model,
        "pv_model": plan.config.pv_model,
        "alpha": plan.config.alpha,
        "terminal_mode": plan.config.terminal_mode,
        "load_actual_kw": float(dataset.load_kw[plan.day_index, step]),
        "pv_actual_kw": float(dataset.pv_kw[plan.day_index, step]),
        "load_actual_kwh": load,
        "pv_actual_kwh": pv,
        "load_forecast_kwh": float(plan.forecast.load_kwh[step]),
        "pv_forecast_kwh": float(plan.forecast.pv_kwh[step]),
        "net_forecast_kwh": float(plan.forecast.net_kwh[step]),
        "residual_margin_kwh": float(plan.margin[step]),
        "plan_grid_kwh": q,
        "emergency_grid_kwh": float(action["emergency_kwh"]),
        "charge_kwh": float(action["charge_kwh"]),
        "discharge_kwh": float(action["discharge_kwh"]),
        "soc_before_kwh": float(soc),
        "soc_after_kwh": float(action["soc_after_kwh"]),
        "unused_supply_kwh": float(action["unused_supply_kwh"]),
        "price_yuan_per_kwh": price,
        "plan_cost_yuan": price * q,
        "emergency_cost_yuan": 5.0 * price * float(action["emergency_kwh"]),
        "forecast_source_max_timestamp": plan.forecast.source_max_timestamp,
        "residual_source_max_timestamp": plan.residual_max_timestamp,
        "residual_history_days": len(plan.residual_days),
    }


def simulate(
    dataset: core.Dataset,
    first_day: int,
    last_day: int,
    config: StrategyConfig,
    forecast_library: Dict[Tuple[str, str], Dict[int, core.Forecast]],
    warmup_config: Optional[StrategyConfig] = None,
    warmup_last_day: int = OUTPUT_START_INDEX - 1,
    r2_anchor: float = 6000.0,
) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[int, Plan]]:
    records: List[Dict[str, object]] = []
    decisions: List[Dict[str, object]] = []
    plans: Dict[int, Plan] = {}
    soc_midnight = INITIAL_SOC
    previous: Optional[Plan] = None

    for day in range(first_day, last_day + 1):
        active = warmup_config if warmup_config is not None and day <= warmup_last_day else config
        plan = make_plan(dataset, day, soc_midnight, previous, active, forecast_library, r2_anchor)
        plans[day] = plan
        decisions.append({
            "decision_day": dataset.dates[day],
            "decision_time": dataset.dates[day],
            "strategy": active.name,
            "load_model": active.load_model,
            "pv_model": active.pv_model,
            "alpha": active.alpha,
            "terminal_mode": active.terminal_mode,
            "decision_soc_0000_kwh": soc_midnight,
            "predicted_plan_start_soc_0010_kwh": plan.predicted_start_soc,
            "realized_plan_start_soc_0010_kwh": np.nan,
            "forecast_source_max_timestamp": plan.forecast.source_max_timestamp,
            "residual_source_max_timestamp": plan.residual_max_timestamp,
            "residual_history_days": len(plan.residual_days),
            "planning_cash_yuan": plan.planning_cash,
            "planning_objective_yuan": plan.planning_objective,
            "planning_max_residual_kwh": plan.max_lp_residual,
            "solve_seconds": plan.solve_seconds,
        })
        if previous is None:
            actual_start = INITIAL_SOC
        else:
            bridge = execute_step(dataset, previous, N_STEPS - 1, soc_midnight)
            records.append(bridge)
            actual_start = float(bridge["soc_after_kwh"])
        decisions[-1]["realized_plan_start_soc_0010_kwh"] = actual_start
        soc = actual_start
        for step in range(N_STEPS - 1):
            row = execute_step(dataset, plan, step, soc)
            records.append(row)
            soc = float(row["soc_after_kwh"])
        soc_midnight = soc
        previous = plan

    if previous is not None:
        last = execute_step(dataset, previous, N_STEPS - 1, soc_midnight)
        records.append(last)

    log = pd.DataFrame(records).sort_values("interval_start", kind="stable").reset_index(drop=True)
    return log, pd.DataFrame(decisions), plans


def metric_window(log: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp) -> pd.DataFrame:
    return log.loc[(log["interval_start"] >= start) & (log["interval_start"] < end)].copy()


def summarize_window(frame: pd.DataFrame) -> Dict[str, float]:
    return {
        "intervals": int(len(frame)),
        "plan_purchase_kwh": float(frame["plan_grid_kwh"].sum()),
        "plan_cost_yuan": float(frame["plan_cost_yuan"].sum()),
        "emergency_purchase_kwh": float(frame["emergency_grid_kwh"].sum()),
        "emergency_cost_yuan": float(frame["emergency_cost_yuan"].sum()),
        "total_cash_yuan": float(frame[["plan_cost_yuan", "emergency_cost_yuan"]].sum().sum()),
        "unused_supply_kwh": float(frame["unused_supply_kwh"].sum()),
        "end_soc_kwh": float(frame.iloc[-1]["soc_after_kwh"]),
        "soc_min_kwh": float(min(frame["soc_before_kwh"].min(), frame["soc_after_kwh"].min())),
        "soc_max_kwh": float(max(frame["soc_before_kwh"].max(), frame["soc_after_kwh"].max())),
    }


def january_score(log: pd.DataFrame) -> Dict[str, float]:
    cash = float(log[["plan_cost_yuan", "emergency_cost_yuan"]].sum().sum())
    end_soc = float(log.iloc[-1]["soc_after_kwh"])
    emergency = float(log["emergency_grid_kwh"].sum())
    lambda_asset = ETA_D * float(np.median(core.load_dataset().price))
    return {
        "cash_yuan": cash,
        "emergency_kwh": emergency,
        "end_soc_kwh": end_soc,
        "lambda_asset_yuan_per_kwh": lambda_asset,
        "score_yuan": cash - lambda_asset * (end_soc - INITIAL_SOC),
    }


def select_r1(
    dataset: core.Dataset,
    library: Dict[Tuple[str, str], Dict[int, core.Forecast]],
) -> Tuple[StrategyConfig, pd.DataFrame]:
    rows: List[Dict[str, object]] = []
    combos = [(l, p) for l in LOAD_MODELS for p in PV_MODELS]
    for combo_index, combo in enumerate(combos):
        cfg = StrategyConfig(f"candidate_{combo_index}", combo[0], combo[1], 0.70, "soft")
        log, _, _ = simulate(dataset, 0, JAN_SELECTION_LAST_INDEX, cfg, library)
        rows.append({"stage": "forecast", "combo_order": combo_index, "load_model": combo[0], "pv_model": combo[1], "alpha": 0.70, **january_score(log)})
    forecast_table = pd.DataFrame(rows).sort_values(
        ["score_yuan", "emergency_kwh", "combo_order"], kind="stable"
    )
    retained = [tuple(x) for x in forecast_table[["load_model", "pv_model"]].head(2).to_numpy()]

    for combo in retained:
        combo_index = combos.index(combo)
        for alpha in ALPHAS:
            cfg = StrategyConfig("alpha_candidate", combo[0], combo[1], alpha, "soft")
            log, _, _ = simulate(dataset, 0, JAN_SELECTION_LAST_INDEX, cfg, library)
            rows.append({"stage": "alpha", "combo_order": combo_index, "load_model": combo[0], "pv_model": combo[1], "alpha": alpha, **january_score(log)})
    table = pd.DataFrame(rows)
    alpha_table = table.loc[table["stage"] == "alpha"].sort_values(
        ["score_yuan", "emergency_kwh", "alpha", "combo_order"], kind="stable"
    )
    best = alpha_table.iloc[0]
    m0 = table.loc[
        (table["stage"] == "forecast")
        & (table["load_model"] == "lag7")
        & (table["pv_model"] == "lag7")
    ].iloc[0]
    if float(best["score_yuan"]) >= float(m0["score_yuan"]) - 1e-9:
        selected = StrategyConfig("R1_fallback_M0", "lag7", "lag7", 0.70, "soft")
    else:
        selected = StrategyConfig("R1", str(best["load_model"]), str(best["pv_model"]), float(best["alpha"]), "soft")
    table["selected"] = (
        (table["stage"] == "alpha")
        & (table["load_model"] == selected.load_model)
        & (table["pv_model"] == selected.pv_model)
        & (np.isclose(table["alpha"].astype(float), float(selected.alpha)))
    )
    return selected, table


def daily_costs(frame: pd.DataFrame) -> pd.Series:
    return frame.assign(cash=frame["plan_cost_yuan"] + frame["emergency_cost_yuan"]).groupby("plan_day")["cash"].sum()


def block_bootstrap(diff: np.ndarray, block: int, reps: int = BOOTSTRAP_REPS) -> Dict[str, float]:
    rng = np.random.default_rng(BOOTSTRAP_SEED + block)
    n = len(diff)
    starts = np.arange(max(1, n - block + 1))
    means = np.empty(reps)
    blocks_needed = int(np.ceil(n / block))
    for i in range(reps):
        sample = np.concatenate([diff[s:s + block] for s in rng.choice(starts, size=blocks_needed, replace=True)])[:n]
        means[i] = float(np.mean(sample))
    return {
        "block_days": block,
        "reps": reps,
        "mean_daily_saving_yuan": float(np.mean(diff)),
        "ci95_low_yuan_per_day": float(np.quantile(means, 0.025)),
        "ci95_high_yuan_per_day": float(np.quantile(means, 0.975)),
    }


def emergency_events(frame: pd.DataFrame) -> pd.DataFrame:
    events = core.merge_emergency_events(frame)
    if events.empty:
        return pd.DataFrame(columns=["date", "start", "end", "time_period", "emergency_kwh", "duration_minutes"])
    events["date"] = pd.to_datetime(events["start"]).dt.normalize()
    events["time_period"] = [core.format_event_time(pd.Timestamp(s), pd.Timestamp(e)) for s, e in zip(events["start"], events["end"])]
    events["duration_minutes"] = (pd.to_datetime(events["end"]) - pd.to_datetime(events["start"])).dt.total_seconds() / 60.0
    return events[["date", "start", "end", "time_period", "emergency_kwh", "duration_minutes"]]


def natural_blocks(frame: pd.DataFrame, days: pd.DatetimeIndex) -> pd.DataFrame:
    parts = [core.block_summary(metric_window(frame, day, day + pd.Timedelta(days=1)), day) for day in days]
    return pd.concat(parts, ignore_index=True)


def _copy_cell_style(source: openpyxl.cell.Cell, target: openpyxl.cell.Cell) -> None:
    if source.has_style:
        target._style = copy.copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)


def write_result2(
    output_path: Path,
    dataset: core.Dataset,
    log: pd.DataFrame,
    plans: Dict[int, Plan],
    blocks: pd.DataFrame,
    events_wc: pd.DataFrame,
) -> Dict[str, object]:
    shutil.copy2(TEMPLATE_PATH, output_path)
    book = openpyxl.load_workbook(output_path)
    plan_sheet = book["计划购电量"]
    storage_sheet = book["充放电量"]
    emergency_sheet = book["紧急购电量"]
    output_days = dataset.dates[OUTPUT_START_INDEX:OUTPUT_END_INDEX + 1]

    for row_offset, day_index in enumerate(range(OUTPUT_START_INDEX, OUTPUT_END_INDEX + 1), start=2):
        plan = plans[day_index]
        plan_sheet.cell(row_offset, 1, dataset.dates[day_index].to_pydatetime())
        for step, value in enumerate(plan.q, start=2):
            plan_sheet.cell(row_offset, step, float(value))
        plan_sheet.cell(row_offset, 146, float(plan.q.sum()))
        plan_sheet.cell(row_offset, 147, float(np.dot(dataset.price, plan.q)))

    storage_style = [[copy.copy(storage_sheet.cell(r, c)) for c in range(1, 7)] for r in range(2, 8)]
    storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
    for day_offset, day in enumerate(output_days):
        day_blocks = blocks.loc[blocks["natural_day"] == day].sort_values("block")
        natural = metric_window(log, day, day + pd.Timedelta(days=1)).sort_values("interval_start")
        if len(day_blocks) != 6 or len(natural) != N_STEPS:
            raise AssertionError(f"自然日块不完整: {day.date()}")
        for block_index, (_, item) in enumerate(day_blocks.iterrows()):
            row = 2 + day_offset * 6 + block_index
            for c in range(1, 7):
                _copy_cell_style(storage_style[block_index][c - 1], storage_sheet.cell(row, c))
            storage_sheet.cell(row, 1, day.to_pydatetime() if block_index == 0 else None)
            storage_sheet.cell(row, 2, item["block"])
            storage_sheet.cell(row, 3, float(item["charge_kwh"]))
            storage_sheet.cell(row, 4, float(item["discharge_kwh"]))
            storage_sheet.cell(row, 5, "0:00" if block_index == 0 else ("24:00" if block_index == 1 else None))
            storage_sheet.cell(row, 6, float(natural.iloc[0]["soc_before_kwh"]) if block_index == 0 else (float(natural.iloc[-1]["soc_after_kwh"]) if block_index == 1 else None))

    emergency_style = [[copy.copy(emergency_sheet.cell(r, c)) for c in range(1, 4)] for r in range(2, 5)]
    emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
    output_event_rows = events_wc.copy().reset_index(drop=True)
    if output_event_rows.empty:
        output_event_rows = pd.DataFrame([{"date": output_days[0], "time_period": None, "emergency_kwh": 0.0}])
    last_date = None
    for i, item in output_event_rows.iterrows():
        row = 2 + i
        for c in range(1, 4):
            _copy_cell_style(emergency_style[i % 3][c - 1], emergency_sheet.cell(row, c))
        event_date = pd.Timestamp(item["date"])
        emergency_sheet.cell(row, 1, event_date.to_pydatetime() if last_date != event_date else None)
        emergency_sheet.cell(row, 2, item["time_period"])
        emergency_sheet.cell(row, 3, float(item["emergency_kwh"]))
        last_date = event_date

    book.save(output_path)

    check = openpyxl.load_workbook(output_path, data_only=False)
    p = check["计划购电量"]
    read_q = np.asarray([[p.cell(r, c).value for c in range(2, 146)] for r in range(2, 336)], dtype=float)
    expected_q = np.vstack([plans[i].q for i in range(OUTPUT_START_INDEX, OUTPUT_END_INDEX + 1)])
    read_plan_totals = np.asarray([p.cell(r, 146).value for r in range(2, 336)], dtype=float)
    read_plan_costs = np.asarray([p.cell(r, 147).value for r in range(2, 336)], dtype=float)
    expected_plan_totals = expected_q.sum(axis=1)
    expected_plan_costs = expected_q @ dataset.price
    read_plan_dates = pd.DatetimeIndex(pd.to_datetime([p.cell(r, 1).value for r in range(2, 336)]))
    s = check["充放电量"]
    e = check["紧急购电量"]
    read_charge = np.asarray([s.cell(r, 3).value for r in range(2, 2 + 334 * 6)], dtype=float)
    read_discharge = np.asarray([s.cell(r, 4).value for r in range(2, 2 + 334 * 6)], dtype=float)
    expected_charge = blocks["charge_kwh"].to_numpy(float)
    expected_discharge = blocks["discharge_kwh"].to_numpy(float)
    expected_soc_0000 = []
    expected_soc_2400 = []
    read_soc_0000 = []
    read_soc_2400 = []
    for day_offset, day in enumerate(output_days):
        natural = metric_window(log, day, day + pd.Timedelta(days=1)).sort_values("interval_start")
        expected_soc_0000.append(float(natural.iloc[0]["soc_before_kwh"]))
        expected_soc_2400.append(float(natural.iloc[-1]["soc_after_kwh"]))
        read_soc_0000.append(float(s.cell(2 + day_offset * 6, 6).value))
        read_soc_2400.append(float(s.cell(3 + day_offset * 6, 6).value))
    event_dates: List[pd.Timestamp] = []
    event_periods: List[object] = []
    event_values: List[float] = []
    carried_date: Optional[pd.Timestamp] = None
    for row in range(2, e.max_row + 1):
        if e.cell(row, 1).value is not None:
            carried_date = pd.Timestamp(e.cell(row, 1).value)
        event_dates.append(carried_date)
        event_periods.append(e.cell(row, 2).value)
        event_values.append(float(e.cell(row, 3).value))
    expected_event_count = max(1, len(events_wc))
    event_match = bool(len(event_values) == expected_event_count)
    if len(events_wc):
        event_match &= bool(
            pd.DatetimeIndex(event_dates).equals(pd.DatetimeIndex(pd.to_datetime(events_wc["date"])))
            and event_periods == events_wc["time_period"].tolist()
            and np.max(np.abs(np.asarray(event_values) - events_wc["emergency_kwh"].to_numpy(float))) <= ABS_TOL
        )
    return {
        "sheet_names_preserved": check.sheetnames == ["计划购电量", "充放电量", "紧急购电量"],
        "plan_rows": p.max_row - 1,
        "plan_columns": p.max_column,
        "plan_numeric_cells": int(np.isfinite(read_q).sum()),
        "plan_max_abs_diff_kwh": float(np.max(np.abs(read_q - expected_q))),
        "plan_dates_match": bool(read_plan_dates.equals(output_days)),
        "plan_totals_max_abs_diff_kwh": float(np.max(np.abs(read_plan_totals - expected_plan_totals))),
        "plan_costs_max_abs_diff_yuan": float(np.max(np.abs(read_plan_costs - expected_plan_costs))),
        "storage_rows": s.max_row - 1,
        "storage_charge_max_abs_diff_kwh": float(np.max(np.abs(read_charge - expected_charge))),
        "storage_discharge_max_abs_diff_kwh": float(np.max(np.abs(read_discharge - expected_discharge))),
        "storage_soc_0000_max_abs_diff_kwh": float(np.max(np.abs(np.asarray(read_soc_0000) - np.asarray(expected_soc_0000)))),
        "storage_soc_2400_max_abs_diff_kwh": float(np.max(np.abs(np.asarray(read_soc_2400) - np.asarray(expected_soc_2400)))),
        "emergency_rows": e.max_row - 1,
        "emergency_events_match": event_match,
    }


def audit_full(
    dataset: core.Dataset,
    log: pd.DataFrame,
    decisions: pd.DataFrame,
    plans: Dict[int, Plan],
    workbook_readback: Dict[str, object],
    blocks: pd.DataFrame,
    events_wp: pd.DataFrame,
    events_wc: pd.DataFrame,
) -> Dict[str, object]:
    starts = pd.DatetimeIndex(log["interval_start"])
    ends = pd.DatetimeIndex(log["interval_end"])
    durations = (ends - starts).total_seconds() / 60.0
    gaps = starts[1:].asi8 - ends[:-1].asi8
    balance = (
        log["plan_grid_kwh"] + log["pv_actual_kwh"] + log["discharge_kwh"]
        + log["emergency_grid_kwh"] - log["load_actual_kwh"]
        - log["charge_kwh"] - log["unused_supply_kwh"]
    ).to_numpy(float)
    state = (
        log["soc_after_kwh"] - log["soc_before_kwh"]
        - ETA_C * log["charge_kwh"] + log["discharge_kwh"] / ETA_D
    ).to_numpy(float)
    scale = max(1.0, float(log[["load_actual_kwh", "pv_actual_kwh", "plan_grid_kwh"]].to_numpy().max()))
    wp_start = dataset.dates[OUTPUT_START_INDEX] + pd.Timedelta(minutes=10)
    wp_end = dataset.dates[OUTPUT_END_INDEX] + pd.Timedelta(days=1, minutes=10)
    wc_start = dataset.dates[OUTPUT_START_INDEX]
    wc_end = dataset.dates[OUTPUT_END_INDEX] + pd.Timedelta(days=1)
    wp = metric_window(log, wp_start, wp_end)
    wc = metric_window(log, wc_start, wc_end)
    forecast_time_ok = all(
        pd.isna(row.forecast_source_max_timestamp)
        or pd.Timestamp(row.forecast_source_max_timestamp) <= pd.Timestamp(row.decision_time)
        for row in decisions.itertuples(index=False)
    )
    residual_time_ok = all(
        pd.isna(row.residual_source_max_timestamp)
        or pd.Timestamp(row.residual_source_max_timestamp) <= pd.Timestamp(row.decision_time)
        for row in decisions.itertuples(index=False)
    )
    terminal_value_time_ok = all(
        plan.terminal_curve is None
        or (
            (plan.terminal_curve["forecast_source_max_timestamp"] is None
             or pd.Timestamp(plan.terminal_curve["forecast_source_max_timestamp"]) <= plan.decision_time)
            and (plan.terminal_curve["residual_source_max_timestamp"] is None
                 or pd.Timestamp(plan.terminal_curve["residual_source_max_timestamp"]) <= plan.decision_time)
        )
        for plan in plans.values()
    )
    # Causal fault injection on a representative required date.  Both the
    # entire target-day truth and the not-yet-realized prior bridge truth are
    # changed, then the forecast library and plan are rebuilt from scratch.
    leakage_day = int(np.flatnonzero(dataset.dates == pd.Timestamp("2025-03-20"))[0])
    perturbed_load = dataset.load_kw.copy()
    perturbed_pv = dataset.pv_kw.copy()
    perturbed_load[leakage_day] += 1234.5
    perturbed_pv[leakage_day] += 678.9
    perturbed_load[leakage_day - 1, -1] += 2345.6
    perturbed_pv[leakage_day - 1, -1] += 987.6
    injection_effective = bool(
        np.max(np.abs(perturbed_load - dataset.load_kw)) > 0
        and np.max(np.abs(perturbed_pv - dataset.pv_kw)) > 0
    )
    perturbed_dataset = core.Dataset(
        dates=dataset.dates,
        input_time_labels=dataset.input_time_labels,
        template_labels=dataset.template_labels,
        price=dataset.price,
        load_kw=perturbed_load,
        pv_kw=perturbed_pv,
        load_kwh=perturbed_load * DT_HOURS,
        pv_kwh=perturbed_pv * DT_HOURS,
        cold_load_kwh=dataset.cold_load_kwh,
        cold_pv_kwh=dataset.cold_pv_kwh,
    )
    leakage_config = plans[leakage_day].config
    perturbed_library = build_forecast_library(perturbed_dataset, [leakage_config.combo])
    decision_soc = float(
        decisions.loc[decisions["decision_day"] == dataset.dates[leakage_day], "decision_soc_0000_kwh"].iloc[0]
    )
    rebuilt_plan = make_plan(
        perturbed_dataset,
        leakage_day,
        decision_soc,
        plans[leakage_day - 1],
        leakage_config,
        perturbed_library,
    )
    future_truth_plan_invariant = bool(
        np.max(np.abs(rebuilt_plan.q - plans[leakage_day].q)) <= 1e-9
        and abs(rebuilt_plan.predicted_start_soc - plans[leakage_day].predicted_start_soc) <= 1e-9
    )
    plan_cost_ok = np.max(np.abs(log["plan_cost_yuan"] - log["price_yuan_per_kwh"] * log["plan_grid_kwh"])) <= 1e-10
    emergency_cost_ok = np.max(np.abs(log["emergency_cost_yuan"] - 5.0 * log["price_yuan_per_kwh"] * log["emergency_grid_kwh"])) <= 1e-10
    expected_blocks = 334 * 6
    wp_event_sum = float(events_wp["emergency_kwh"].sum()) if not events_wp.empty else 0.0
    wc_event_sum = float(events_wc["emergency_kwh"].sum()) if not events_wc.empty else 0.0
    checks = {
        "01_all_intervals_10min": bool(np.allclose(durations, 10.0)),
        "02_no_duplicate_timestamps": bool(not starts.duplicated().any()),
        "03_no_time_gaps_after_initial_boundary": bool(np.all(gaps == 0)),
        "04_kw_to_kwh_conversion": bool(np.max(np.abs(log["load_actual_kw"] * DT_HOURS - log["load_actual_kwh"])) <= 1e-12 and np.max(np.abs(log["pv_actual_kw"] * DT_HOURS - log["pv_actual_kwh"])) <= 1e-12),
        "05_soc_bounds": bool(log["soc_before_kwh"].min() >= SOC_MIN - ABS_TOL and log["soc_after_kwh"].max() <= SOC_MAX + ABS_TOL),
        "06_charge_discharge_power_bounds": bool(log[["charge_kwh", "discharge_kwh"]].to_numpy().max() <= STEP_LIMIT + ABS_TOL),
        "07_mutual_exclusion": bool(not ((log["charge_kwh"] > 1e-10) & (log["discharge_kwh"] > 1e-10)).any()),
        "08_nonnegative_grid_purchase": bool(log[["plan_grid_kwh", "emergency_grid_kwh"]].to_numpy().min() >= -ABS_TOL),
        "09_energy_balance_and_normalized_residual": bool(np.max(np.abs(balance)) <= ABS_TOL and max(np.max(np.abs(balance)), np.max(np.abs(state))) / scale <= NORM_TOL),
        "10_no_future_leakage": bool(
            forecast_time_ok
            and residual_time_ok
            and terminal_value_time_ok
            and injection_effective
            and future_truth_plan_invariant
        ),
        "11_continuous_soc": bool(np.max(np.abs(log["soc_before_kwh"].to_numpy()[1:] - log["soc_after_kwh"].to_numpy()[:-1])) <= ABS_TOL),
        "12_plan_cost_is_p_times_q": bool(plan_cost_ok),
        "13_emergency_cost_is_5p_times_h": bool(emergency_cost_ok),
        "14_emergency_never_charges_battery": bool(not ((log["emergency_grid_kwh"] > 1e-10) & (log["charge_kwh"] > 1e-10)).any()),
        "15_result2_plan_window_334x144": bool(len(wp) == 334 * N_STEPS and workbook_readback["plan_rows"] == 334 and workbook_readback["plan_numeric_cells"] == 334 * N_STEPS),
        "16_natural_day_blocks_and_soc": bool(len(wc) == 334 * N_STEPS and len(blocks) == expected_blocks and (blocks["interval_count"] == 24).all() and workbook_readback["storage_rows"] == expected_blocks),
        "17_events_and_template_readback": bool(
            abs(wp_event_sum - wp["emergency_grid_kwh"].sum()) <= ABS_TOL
            and abs(wc_event_sum - wc["emergency_grid_kwh"].sum()) <= ABS_TOL
            and workbook_readback["sheet_names_preserved"]
            and workbook_readback["plan_dates_match"]
            and workbook_readback["plan_max_abs_diff_kwh"] <= ABS_TOL
            and workbook_readback["plan_totals_max_abs_diff_kwh"] <= ABS_TOL
            and workbook_readback["plan_costs_max_abs_diff_yuan"] <= ABS_TOL
            and workbook_readback["storage_charge_max_abs_diff_kwh"] <= ABS_TOL
            and workbook_readback["storage_discharge_max_abs_diff_kwh"] <= ABS_TOL
            and workbook_readback["storage_soc_0000_max_abs_diff_kwh"] <= ABS_TOL
            and workbook_readback["storage_soc_2400_max_abs_diff_kwh"] <= ABS_TOL
            and workbook_readback["emergency_events_match"]
        ),
    }
    return {
        "checks": checks,
        "metrics": {
            "max_balance_residual_kwh": float(np.max(np.abs(balance))),
            "max_state_residual_kwh": float(np.max(np.abs(state))),
            "normalized_max_residual": float(max(np.max(np.abs(balance)), np.max(np.abs(state))) / scale),
            "simultaneous_steps": int(((log["charge_kwh"] > 1e-10) & (log["discharge_kwh"] > 1e-10)).sum()),
            "forecast_time_ok": forecast_time_ok,
            "residual_time_ok": residual_time_ok,
            "terminal_value_time_ok": terminal_value_time_ok,
            "leakage_fault_injection_effective": injection_effective,
            "future_truth_plan_invariant": future_truth_plan_invariant,
            "planning_lp_max_residual_kwh": float(max(plan.max_lp_residual for plan in plans.values())),
            "planning_solve_p95_seconds": float(np.quantile([plan.solve_seconds for plan in plans.values()], 0.95)),
        },
        "wp": summarize_window(wp),
        "wc": summarize_window(wc),
        "boundary_reconciliation": {
            "wc_only_cash_yuan": float(wc.iloc[0]["plan_cost_yuan"] + wc.iloc[0]["emergency_cost_yuan"]),
            "wp_only_cash_yuan": float(wp.iloc[-1]["plan_cost_yuan"] + wp.iloc[-1]["emergency_cost_yuan"]),
            "identity_residual_yuan": float(
                summarize_window(wc)["total_cash_yuan"] - summarize_window(wp)["total_cash_yuan"]
                - (wc.iloc[0]["plan_cost_yuan"] + wc.iloc[0]["emergency_cost_yuan"])
                + (wp.iloc[-1]["plan_cost_yuan"] + wp.iloc[-1]["emergency_cost_yuan"])
            ),
        },
    }


def solve_oracle(
    dataset: core.Dataset,
    start_soc: float,
    end_soc: float,
) -> Dict[str, float]:
    actual_net = (dataset.load_kwh[OUTPUT_START_INDEX:OUTPUT_END_INDEX + 1] - dataset.pv_kwh[OUTPUT_START_INDEX:OUTPUT_END_INDEX + 1]).reshape(-1)
    price = np.tile(dataset.price, OUTPUT_END_INDEX - OUTPUT_START_INDEX + 1)
    n = len(actual_net)
    oq, oc, od, ow, oe = 0, n, 2 * n, 3 * n, 4 * n
    nvar = 5 * n + 1
    objective = np.zeros(nvar)
    objective[oq:oq + n] = price
    objective[oc:oc + n] = EPS_THROUGHPUT
    objective[od:od + n] = EPS_THROUGHPUT
    aeq = sparse.lil_matrix((2 * n, nvar))
    beq = np.zeros(2 * n)
    for t in range(n):
        aeq[t, oq + t] = 1.0
        aeq[t, oc + t] = -1.0
        aeq[t, od + t] = 1.0
        aeq[t, ow + t] = -1.0
        beq[t] = actual_net[t]
        row = n + t
        aeq[row, oe + t] = -1.0
        aeq[row, oe + t + 1] = 1.0
        aeq[row, oc + t] = -ETA_C
        aeq[row, od + t] = 1.0 / ETA_D
    bounds: List[Tuple[Optional[float], Optional[float]]] = []
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(0.0, STEP_LIMIT)] * n)
    bounds.extend([(0.0, STEP_LIMIT)] * n)
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(SOC_MIN, SOC_MAX)] * (n + 1))
    bounds[oe] = (start_soc, start_soc)
    bounds[oe + n] = (end_soc, end_soc)
    started = time.perf_counter()
    result = linprog(objective, A_eq=aeq.tocsr(), b_eq=beq, bounds=bounds, method="highs")
    elapsed = time.perf_counter() - started
    if not result.success:
        raise RuntimeError(f"Oracle LP失败: {result.status} {result.message}")
    residual = np.max(np.abs(aeq.tocsr() @ result.x - beq))
    oracle_cash = float(np.dot(price, np.asarray(result.x)[oq:oq + n]))
    return {
        "cash_lower_bound_yuan": oracle_cash,
        "solve_seconds": elapsed,
        "max_residual_kwh": float(residual),
        "start_soc_kwh": start_soc,
        "end_soc_kwh": end_soc,
    }


def run_full(skip_r2: bool = False, skip_oracle: bool = False) -> Dict[str, object]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    dataset = core.load_dataset()
    combos = [(l, p) for l in LOAD_MODELS for p in PV_MODELS]
    library = build_forecast_library(dataset, combos)
    selected_r1, selection_table = select_r1(dataset, library)
    selection_table.to_csv(OUTPUT_DIR / "q2_r1_selection.csv", index=False)

    m0 = StrategyConfig("M0_audited", "lag7", "lag7", 0.70, "soft")
    b0 = StrategyConfig("B0", "lag7", "lag7", None, "soft")
    logs: Dict[str, pd.DataFrame] = {}
    decisions: Dict[str, pd.DataFrame] = {}
    plan_sets: Dict[str, Dict[int, Plan]] = {}
    for config in (b0, m0, selected_r1):
        log, decision, plans = simulate(
            dataset, 0, OUTPUT_END_INDEX, config, library,
            warmup_config=m0, warmup_last_day=OUTPUT_START_INDEX - 1,
        )
        logs[config.name] = log
        decisions[config.name] = decision
        plan_sets[config.name] = plans

    # Required terminal-effect ablations.  They share the selected forecast,
    # alpha, controller, January warm-up and Dec-31 planning anchor.
    terminal_ablations = (
        StrategyConfig("R1_no_terminal_value", selected_r1.load_model, selected_r1.pv_model, selected_r1.alpha, "none"),
        StrategyConfig("R1_daily_fixed_6000", selected_r1.load_model, selected_r1.pv_model, selected_r1.alpha, "hard"),
    )
    for config in terminal_ablations:
        log, decision, plans = simulate(
            dataset, 0, OUTPUT_END_INDEX, config, library,
            warmup_config=m0, warmup_last_day=OUTPUT_START_INDEX - 1,
        )
        logs[config.name] = log
        decisions[config.name] = decision
        plan_sets[config.name] = plans

    wp_start = dataset.dates[OUTPUT_START_INDEX] + pd.Timedelta(minutes=10)
    wp_end = dataset.dates[OUTPUT_END_INDEX] + pd.Timedelta(days=1, minutes=10)
    strategy_rows = []
    for name, log in logs.items():
        row = {"strategy": name, **summarize_window(metric_window(log, wp_start, wp_end))}
        strategy_rows.append(row)
    strategy_table = pd.DataFrame(strategy_rows)

    r2_result: Dict[str, object] = {"attempted": False, "accepted": False, "reason": "--skip-r2" if skip_r2 else None}
    accepted = selected_r1
    if not skip_r2:
        r2 = StrategyConfig("R2", selected_r1.load_model, selected_r1.pv_model, selected_r1.alpha, "pwl")
        r2_result["attempted"] = True
        try:
            r2_log, r2_decisions, r2_plans = simulate(
                dataset, 0, OUTPUT_END_INDEX, r2, library,
                warmup_config=m0, warmup_last_day=OUTPUT_START_INDEX - 1,
            )
            logs[r2.name] = r2_log
            decisions[r2.name] = r2_decisions
            plan_sets[r2.name] = r2_plans
            r1_wp = metric_window(logs[selected_r1.name], wp_start, wp_end)
            r2_wp = metric_window(r2_log, wp_start, wp_end)
            r1_summary = summarize_window(r1_wp)
            r2_summary = summarize_window(r2_wp)
            daily_diff = (daily_costs(r1_wp) - daily_costs(r2_wp)).to_numpy(float)
            bootstrap = {str(block): block_bootstrap(daily_diff, block) for block in (3, 7, 14)}
            lambda_asset = ETA_D * float(np.median(dataset.price))
            adjusted_saving = (
                r1_summary["total_cash_yuan"] - r2_summary["total_cash_yuan"]
                + lambda_asset * (r2_summary["end_soc_kwh"] - r1_summary["end_soc_kwh"])
            )
            relative = adjusted_saving / r1_summary["total_cash_yuan"]
            shape_ok = all(
                plan.terminal_curve is None
                or (plan.terminal_curve["monotone"] and plan.terminal_curve["convex"])
                for plan in r2_plans.values()
            )
            preliminary_gate = bool(
                relative >= R2_MIN_RELATIVE_SAVING
                and bootstrap["7"]["ci95_low_yuan_per_day"] > 0
                and shape_ok
            )
            anchor_sensitivity: List[Dict[str, object]] = []
            anchor_stable = False
            if preliminary_gate:
                for anchor in R2_ANCHOR_SENSITIVITY:
                    sensitivity_config = StrategyConfig(
                        f"R2_anchor_{int(anchor)}",
                        selected_r1.load_model,
                        selected_r1.pv_model,
                        selected_r1.alpha,
                        "pwl",
                    )
                    sensitivity_log, _, _ = simulate(
                        dataset, 0, OUTPUT_END_INDEX, sensitivity_config, library,
                        warmup_config=m0,
                        warmup_last_day=OUTPUT_START_INDEX - 1,
                        r2_anchor=anchor,
                    )
                    sensitivity_wp = metric_window(sensitivity_log, wp_start, wp_end)
                    sensitivity_summary = summarize_window(sensitivity_wp)
                    sensitivity_diff = (daily_costs(r1_wp) - daily_costs(sensitivity_wp)).to_numpy(float)
                    sensitivity_bootstrap = block_bootstrap(sensitivity_diff, 7)
                    sensitivity_adjusted_saving = (
                        r1_summary["total_cash_yuan"] - sensitivity_summary["total_cash_yuan"]
                        + lambda_asset * (sensitivity_summary["end_soc_kwh"] - r1_summary["end_soc_kwh"])
                    )
                    anchor_sensitivity.append({
                        "next_day_anchor_kwh": anchor,
                        "summary": sensitivity_summary,
                        "asset_adjusted_saving_yuan": float(sensitivity_adjusted_saving),
                        "relative_asset_adjusted_saving": float(sensitivity_adjusted_saving / r1_summary["total_cash_yuan"]),
                        "bootstrap_7day": sensitivity_bootstrap,
                    })
                anchor_stable = all(
                    item["asset_adjusted_saving_yuan"] > 0
                    and item["bootstrap_7day"]["ci95_low_yuan_per_day"] > 0
                    for item in anchor_sensitivity
                )
            gate = bool(preliminary_gate and anchor_stable)
            r2_result.update({
                "summary": r2_summary,
                "asset_adjusted_saving_yuan": float(adjusted_saving),
                "relative_asset_adjusted_saving": float(relative),
                "bootstrap": bootstrap,
                "value_curve_shape_ok": shape_ok,
                "preliminary_gate": preliminary_gate,
                "anchor_sensitivity": anchor_sensitivity,
                "anchor_stable": anchor_stable,
                "accepted": gate,
            })
            strategy_table = pd.concat([strategy_table, pd.DataFrame([{"strategy": "R2", **r2_summary}])], ignore_index=True)
            if gate:
                accepted = r2
                r2_result["reason"] = "通过0.2%资产调整收益、7日块区间、价值曲线形状及上下锚点稳定性门槛"
            else:
                r2_result["reason"] = "未同时通过0.2%资产调整收益、7日块区间、价值曲线形状及上下锚点稳定性门槛，回退R1"
        except Exception as exc:
            r2_result.update({"accepted": False, "reason": f"R2执行失败，回退R1: {type(exc).__name__}: {exc}"})

    strategy_table.to_csv(OUTPUT_DIR / "q2_strategy_comparison.csv", index=False)
    accepted_log = logs[accepted.name]
    accepted_decisions = decisions[accepted.name]
    accepted_plans = plan_sets[accepted.name]
    wc_start = dataset.dates[OUTPUT_START_INDEX]
    wc_end = dataset.dates[OUTPUT_END_INDEX] + pd.Timedelta(days=1)
    wp = metric_window(accepted_log, wp_start, wp_end)
    wc = metric_window(accepted_log, wc_start, wc_end)
    blocks = natural_blocks(accepted_log, dataset.dates[OUTPUT_START_INDEX:OUTPUT_END_INDEX + 1])
    events_wp = emergency_events(wp)
    events_wc = emergency_events(wc)

    accepted_log.to_csv(OUTPUT_DIR / "q2_execution_log.csv", index=False)
    accepted_decisions.to_csv(OUTPUT_DIR / "q2_decision_audit.csv", index=False)
    blocks.to_csv(OUTPUT_DIR / "q2_natural_day_blocks.csv", index=False)
    events_wp.to_csv(OUTPUT_DIR / "q2_emergency_events_wp.csv", index=False)
    events_wc.to_csv(OUTPUT_DIR / "q2_emergency_events_wc.csv", index=False)
    focus_dates = pd.DatetimeIndex(pd.to_datetime(["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]))
    focus = accepted_log.loc[accepted_log["plan_day"].isin(focus_dates)].copy()
    if len(focus) != 4 * N_STEPS:
        raise AssertionError("四个指定日没有各自完整144段")
    focus.to_csv(OUTPUT_DIR / "q2_focus_days_144.csv", index=False)
    daily_parts = []
    monthly_parts = []
    for strategy_name, strategy_log in logs.items():
        strategy_wp = metric_window(strategy_log, wp_start, wp_end).copy()
        strategy_wp["cash_yuan"] = strategy_wp["plan_cost_yuan"] + strategy_wp["emergency_cost_yuan"]
        daily = strategy_wp.groupby("plan_day", as_index=False).agg(
            plan_purchase_kwh=("plan_grid_kwh", "sum"),
            emergency_purchase_kwh=("emergency_grid_kwh", "sum"),
            plan_cost_yuan=("plan_cost_yuan", "sum"),
            emergency_cost_yuan=("emergency_cost_yuan", "sum"),
            total_cash_yuan=("cash_yuan", "sum"),
            unused_supply_kwh=("unused_supply_kwh", "sum"),
            end_soc_kwh=("soc_after_kwh", "last"),
        )
        daily.insert(0, "strategy", strategy_name)
        daily_parts.append(daily)
        # W_P 的最后一段时间戳是 2026-01-01 00:00，但业务上属于
        # 2025-12-31 计划日的第 144 段；月度统计必须按 plan_day 归属，
        # 避免把这 10 分钟误列为一个孤立的“2026-01 月”。
        monthly = strategy_wp.assign(month=strategy_wp["plan_day"].dt.to_period("M").astype(str)).groupby("month", as_index=False).agg(
            plan_purchase_kwh=("plan_grid_kwh", "sum"),
            emergency_purchase_kwh=("emergency_grid_kwh", "sum"),
            plan_cost_yuan=("plan_cost_yuan", "sum"),
            emergency_cost_yuan=("emergency_cost_yuan", "sum"),
            total_cash_yuan=("cash_yuan", "sum"),
        )
        monthly.insert(0, "strategy", strategy_name)
        monthly_parts.append(monthly)
    pd.concat(daily_parts, ignore_index=True).to_csv(OUTPUT_DIR / "q2_daily_strategy_comparison.csv", index=False)
    pd.concat(monthly_parts, ignore_index=True).to_csv(OUTPUT_DIR / "q2_monthly_strategy_comparison.csv", index=False)
    workbook_path = OUTPUT_DIR / "result2.xlsx"
    workbook_readback = write_result2(workbook_path, dataset, accepted_log, accepted_plans, blocks, events_wc)
    audit = audit_full(dataset, accepted_log, accepted_decisions, accepted_plans, workbook_readback, blocks, events_wp, events_wc)

    accepted_wp_summary = summarize_window(wp)
    oracle = {"attempted": False}
    if not skip_oracle:
        oracle = {"attempted": True, **solve_oracle(dataset, float(wp.iloc[0]["soc_before_kwh"]), float(wp.iloc[-1]["soc_after_kwh"]))}
        oracle["strategy_gap_yuan"] = accepted_wp_summary["total_cash_yuan"] - oracle["cash_lower_bound_yuan"]
        oracle["strategy_gap_fraction"] = oracle["strategy_gap_yuan"] / oracle["cash_lower_bound_yuan"]

    actual_net = wp["load_actual_kwh"] - wp["pv_actual_kwh"]
    forecast_net = wp["load_forecast_kwh"] - wp["pv_forecast_kwh"]
    forecast_metrics = {
        "net_mae_kwh_per_10min": float(np.mean(np.abs(actual_net - forecast_net))),
        "net_rmse_kwh_per_10min": float(np.sqrt(np.mean((actual_net - forecast_net) ** 2))),
        "load_mae_kwh_per_10min": float(np.mean(np.abs(wp["load_actual_kwh"] - wp["load_forecast_kwh"]))),
        "pv_mae_kwh_per_10min": float(np.mean(np.abs(wp["pv_actual_kwh"] - wp["pv_forecast_kwh"]))),
    }
    result: Dict[str, object] = {
        "status": "PASS" if all(audit["checks"].values()) else "FAIL",
        "scope": "Q2 full 334-plan-day causal rolling backtest",
        "selected_r1": selected_r1.__dict__,
        "accepted_strategy": accepted.__dict__,
        "r2_gate": r2_result,
        "forecast_metrics": forecast_metrics,
        "strategy_comparison": strategy_table.to_dict(orient="records"),
        "accepted_wp": accepted_wp_summary,
        "accepted_wc": summarize_window(wc),
        "emergency_events_wp": int(len(events_wp)),
        "emergency_events_wc": int(len(events_wc)),
        "emergency_duration_wp_minutes": float(events_wp["duration_minutes"].sum()) if not events_wp.empty else 0.0,
        "oracle": oracle,
        "audit": audit,
        "template_readback": workbook_readback,
        "parameters": {
            "dt_hours": DT_HOURS,
            "eta_charge": ETA_C,
            "eta_discharge": ETA_D,
            "round_trip_efficiency": ETA_C * ETA_D,
            "soc_bounds_kwh": [SOC_MIN, SOC_MAX],
            "step_limit_kwh": STEP_LIMIT,
            "alphas": ALPHAS,
            "r2_grid_kwh": R2_GRID.tolist(),
            "r2_min_relative_saving": R2_MIN_RELATIVE_SAVING,
            "r2_anchor_sensitivity_kwh": list(R2_ANCHOR_SENSITIVITY),
            "r2_anchor_stability_rule": "both alternative anchors require positive asset-adjusted saving and positive 7-day block CI lower bound",
            "bootstrap_seed": BOOTSTRAP_SEED,
            "bootstrap_reps": BOOTSTRAP_REPS,
        },
        "manifest": {
            "inputs": {
                str((core.DATA_DIR / "附件1.xlsx").relative_to(ROOT)): sha256(core.DATA_DIR / "附件1.xlsx"),
                str((core.DATA_DIR / "附件2.xlsx").relative_to(ROOT)): sha256(core.DATA_DIR / "附件2.xlsx"),
                str(TEMPLATE_PATH.relative_to(ROOT)): sha256(TEMPLATE_PATH),
                core.ROUTE_PATH.name: sha256(core.ROUTE_PATH),
                core.ANALYSIS_PATH.name: sha256(core.ANALYSIS_PATH),
                core.TERMS_PATH.name: sha256(core.TERMS_PATH),
            },
            "code": {
                Path(__file__).name: sha256(Path(__file__)),
                Path(core.__file__).name: sha256(Path(core.__file__)),
            },
            "runtime": {
                "python": sys.version,
                "executable": sys.executable,
                "platform": platform.platform(),
                "numpy": np.__version__,
                "pandas": pd.__version__,
                "scipy": scipy.__version__,
                "openpyxl": openpyxl.__version__,
            },
            "command": "py -3.13 Q2/program/q2_reproduce_all.py",
        },
    }
    summary_path = OUTPUT_DIR / "q2_final_summary.json"
    summary_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    repro = {
        "random_seed": BOOTSTRAP_SEED,
        "input_sha256": result["manifest"]["inputs"],
        "code_sha256": result["manifest"]["code"],
        "runtime": result["manifest"]["runtime"],
        "parameters": result["parameters"],
        "unique_command": result["manifest"]["command"],
        "expected_status": result["status"],
        "expected_accepted_strategy": result["accepted_strategy"],
        "expected_wp_total_cash_yuan": result["accepted_wp"]["total_cash_yuan"],
    }
    (ROOT / "Q2" / "result" / "复现清单.json").write_text(json.dumps(repro, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("full",), default="full")
    parser.add_argument("--skip-r2", action="store_true")
    parser.add_argument("--skip-oracle", action="store_true")
    args = parser.parse_args()
    result = run_full(skip_r2=args.skip_r2, skip_oracle=args.skip_oracle)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["status"] != "PASS":
        raise SystemExit(2)


if __name__ == "__main__":
    main()
