"""Q3 P1 vertical slice under the frozen section-13 model contract.

This file deliberately runs one real plan day only.  It verifies the four
forecast releases, three non-overlapping adjustment blocks, B-settlement
ledger, physical feedback execution, the cross-day bridge, and causal-data
fault injections.  A full 334-day run is forbidden until an independent P1
review returns PASS.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import platform
import shutil
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
import pandas as pd
import scipy
from scipy import sparse
from scipy.optimize import linprog

PROGRAM_DIR = Path(__file__).resolve().parent
ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "Q2" / "program"))

import q2_pipeline_v1 as core
import q2_v2_compatible as q2full


DATA_DIR = ROOT / "附件"
TEMPLATE = DATA_DIR / "附件5" / "result3.xlsx"
OUTPUT_DIR = ROOT / "Q3" / "v2" / "result" / "p1"
TARGET_DATE = pd.Timestamp("2025-02-01")
TARGET_INDEX = 31
RELEASE_HOURS = (0, 6, 12, 18)
BLOCKS = {6: (35, 71), 12: (71, 107), 18: (107, 144)}
ALPHA = 0.65
# Q2 v2 upstream point forecasts.  Attachment 3 remains the within-day
# release channel; these models supply the causal baseline and terminal value.
LOAD_MODEL = "same_weekday_2w"
PV_MODEL_FOR_Q2_VALUE = "recent_5d"
RISK_DAYS = 60
GATE_DAYS = 14
MIN_POOL = 7
TOL = 1e-5


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(frozen=True)
class ForecastBook:
    dates: pd.DatetimeIndex
    issue_hours: np.ndarray
    values_kw: np.ndarray
    anchors_kw: np.ndarray

    def row(self, day_index: int, release_hour: int) -> np.ndarray:
        hit = np.flatnonzero(
            (self.dates == self.dates[4 * day_index])
            & (self.issue_hours == release_hour)
        )
        if len(hit) != 1:
            raise AssertionError(f"附件3行定位失败: day={day_index}, release={release_hour}")
        return self.values_kw[int(hit[0])].copy()

    def anchor(self, day_index: int, release_hour: int) -> float:
        release_index = RELEASE_HOURS.index(release_hour)
        return float(self.anchors_kw[day_index, release_index])


@dataclass
class ReleaseForecast:
    day_index: int
    release_hour: int
    issue_time: pd.Timestamp
    steps: np.ndarray
    load_kwh: np.ndarray
    pv_kwh: np.ndarray
    current_pv_anchor_kw: float
    anchor_source: str
    load_source_max: Optional[pd.Timestamp]
    load_correction_kwh: float
    load_correction_source_max: Optional[pd.Timestamp]

    @property
    def net_kwh(self) -> np.ndarray:
        return self.load_kwh - self.pv_kwh


def load_forecast_book(dataset: core.Dataset) -> ForecastBook:
    frame = pd.read_excel(DATA_DIR / "附件3.xlsx", sheet_name=0, header=0)
    if frame.shape != (1460, 26):
        raise AssertionError(f"附件3形状异常: {frame.shape}")
    date_series = pd.to_datetime(frame.iloc[:, 0].ffill())
    raw_issue = frame.iloc[:, 1].astype(str).str.strip()
    issue_hours = raw_issue.str.extract(r"(\d+)", expand=False).astype(int).to_numpy()
    values = frame.iloc[:, 2:].apply(pd.to_numeric, errors="raise").to_numpy(float)
    if not np.isfinite(values).all() or (values < 0).any():
        raise AssertionError("附件3预测值必须有限非负")
    expected = np.tile(np.asarray(RELEASE_HOURS), 365)
    if not np.array_equal(issue_hours, expected):
        raise AssertionError("附件3发布时间次序不是0/6/12/18")
    dates = pd.DatetimeIndex(date_series)
    if dates.nunique() != 365 or not (dates[::4] == dataset.dates).all():
        raise AssertionError("附件3日期与附件2不一致")
    # Freeze the point-anchor channel separately from the interval-energy
    # arrays.  Values at 06/12/18 are the best available timestamp-matched
    # proxies; midnight is a documented night-time zero.  Subsequent fault
    # injection into future interval actuals therefore cannot mutate a point
    # measurement that was already available at the release instant.
    anchors = np.zeros((365, 4), dtype=float)
    anchors[:, 1] = dataset.pv_kw[:, 35]
    anchors[:, 2] = dataset.pv_kw[:, 71]
    anchors[:, 3] = dataset.pv_kw[:, 107]
    return ForecastBook(
        dates=dates,
        issue_hours=issue_hours,
        values_kw=values,
        anchors_kw=anchors,
    )


def current_pv_anchor(
    book: ForecastBook,
    day_index: int,
    release_hour: int,
) -> Tuple[float, str]:
    """Return a point measurement proxy, never an interval-energy actual.

    Attachment 2 has no separate instantaneous meter channel.  At 06/12/18
    the value bearing that left-end timestamp is used only as a point anchor;
    the executor independently reads the interval energy later.  Midnight is
    a documented night-time 0 kW initialization, avoiding use of the not-yet
    completed 00:00-00:10 bridge interval.
    """
    if release_hour == 0:
        return book.anchor(day_index, release_hour), "nighttime_zero_point_anchor_not_bridge_interval"
    return book.anchor(day_index, release_hour), "frozen_current_point_measurement_proxy"


def release_steps(release_hour: int) -> np.ndarray:
    if release_hour == 0:
        return np.arange(core.N_STEPS, dtype=int)
    start = BLOCKS[release_hour][0]
    return np.arange(start, core.N_STEPS, dtype=int)


def build_release_forecast(
    dataset: core.Dataset,
    book: ForecastBook,
    q2_library: Dict[Tuple[str, str], Dict[int, core.Forecast]],
    day_index: int,
    release_hour: int,
    use_load_correction: bool = True,
    anchor_override_kw: Optional[float] = None,
) -> ReleaseForecast:
    issue = dataset.dates[day_index] + pd.Timedelta(hours=release_hour)
    steps = release_steps(release_hour)
    baseline = q2full.make_forecast(
        dataset,
        day_index,
        issue,
        (LOAD_MODEL, PV_MODEL_FOR_Q2_VALUE),
    )
    load = baseline.load_kwh[steps].copy()
    correction = 0.0
    correction_max: Optional[pd.Timestamp] = None
    if use_load_correction and release_hour > 0:
        completed = []
        for step in range(core.N_STEPS):
            end = dataset.interval_starts(day_index)[step] + pd.Timedelta(minutes=10)
            if end <= issue:
                completed.append(step)
        completed = completed[-12:]
        if completed:
            residual = dataset.load_kwh[day_index, completed] - baseline.load_kwh[completed]
            correction = float(np.median(residual))
            correction_max = max(
                dataset.interval_starts(day_index)[s] + pd.Timedelta(minutes=10)
                for s in completed
            )
            load = np.maximum(0.0, load + correction)

    anchor, anchor_source = current_pv_anchor(book, day_index, release_hour)
    if anchor_override_kw is not None:
        anchor = float(anchor_override_kw)
        anchor_source = "explicit_current_point_anchor_override"
    node_minutes = np.arange(0.0, 1440.0 + 60.0, 60.0)
    node_values = np.concatenate(([anchor], book.row(day_index, release_hour)))
    targets = dataset.interval_starts(day_index)[steps]
    lead_minutes = ((targets - issue) / pd.Timedelta(minutes=1)).to_numpy(float)
    if (lead_minutes < -TOL).any() or (lead_minutes > 1440.0 + TOL).any():
        raise AssertionError("附件3插值目标超出0—24小时")
    pv_kw = np.interp(lead_minutes, node_minutes, node_values)
    return ReleaseForecast(
        day_index=day_index,
        release_hour=release_hour,
        issue_time=issue,
        steps=steps,
        load_kwh=load,
        pv_kwh=pv_kw * core.DT_HOURS,
        current_pv_anchor_kw=anchor,
        anchor_source=anchor_source,
        load_source_max=baseline.source_max_timestamp,
        load_correction_kwh=correction,
        load_correction_source_max=correction_max,
    )


def residual_pools(
    dataset: core.Dataset,
    book: ForecastBook,
    q2_library: Dict[Tuple[str, str], Dict[int, core.Forecast]],
    day_index: int,
    release_hour: int,
    steps: np.ndarray,
    risk_days: int = RISK_DAYS,
    gate_days: int = GATE_DAYS,
) -> Tuple[np.ndarray, np.ndarray, List[int], List[int]]:
    issue = dataset.dates[day_index] + pd.Timedelta(hours=release_hour)
    historical = []
    for old_day in range(max(0, day_index - (risk_days + gate_days + 2)), day_index):
        latest_end = dataset.interval_starts(old_day)[int(np.max(steps))] + pd.Timedelta(minutes=10)
        if latest_end <= issue:
            historical.append(old_day)
    residuals: List[np.ndarray] = []
    valid_days: List[int] = []
    for old_day in historical:
        old = build_release_forecast(
            dataset, book, q2_library, old_day, release_hour, use_load_correction=True
        )
        positions = np.searchsorted(old.steps, steps)
        predicted = old.net_kwh[positions]
        actual = dataset.load_kwh[old_day, steps] - dataset.pv_kwh[old_day, steps]
        residuals.append(actual - predicted)
        valid_days.append(old_day)
    if not residuals:
        width = len(steps)
        return np.zeros((0, width)), np.zeros((0, width)), [], []
    matrix = np.vstack(residuals)
    risk_count = min(risk_days, len(matrix))
    risk = matrix[-risk_count:]
    gate_end = len(matrix) - risk_count
    gate_start = max(0, gate_end - gate_days)
    gate = matrix[gate_start:gate_end]
    return risk, gate, valid_days[-risk_count:], valid_days[gate_start:gate_end]


def solve_adjustment_lp(
    q0: np.ndarray,
    forecast: ReleaseForecast,
    margin_kwh: np.ndarray,
    start_soc: float,
    block: Tuple[int, int],
    price: np.ndarray,
    terminal_curve: Tuple[np.ndarray, np.ndarray],
) -> Dict[str, object]:
    steps = forecast.steps
    n = len(steps)
    ou, or_, oc, od, oh, ow = 0, n, 2 * n, 3 * n, 4 * n, 5 * n
    oe = 6 * n
    oz = oe + n + 1
    nvar = oz + 1
    objective = np.zeros(nvar)
    p = price[steps]
    objective[ou:ou + n] = 1.5 * p
    objective[or_:or_ + n] = -0.5 * p
    objective[oc:oc + n] = core.EPS_THROUGHPUT
    objective[od:od + n] = core.EPS_THROUGHPUT
    objective[oh:oh + n] = 5.0 * p
    objective[oz] = 1.0

    aeq = sparse.lil_matrix((2 * n, nvar))
    beq = np.zeros(2 * n)
    risk_net = forecast.net_kwh + margin_kwh
    for k, step in enumerate(steps):
        aeq[k, ou + k] = 1.0
        aeq[k, or_ + k] = -1.0
        aeq[k, oc + k] = -1.0
        aeq[k, od + k] = 1.0
        aeq[k, oh + k] = 1.0
        aeq[k, ow + k] = -1.0
        beq[k] = float(risk_net[k] - q0[step])
        row = n + k
        aeq[row, oe + k] = -1.0
        aeq[row, oe + k + 1] = 1.0
        aeq[row, oc + k] = -core.ETA_CHARGE
        aeq[row, od + k] = 1.0 / core.ETA_DISCHARGE

    grid, values = terminal_curve
    slopes = np.diff(values) / np.diff(grid)
    intercepts = values[:-1] - slopes * grid[:-1]
    aub = sparse.lil_matrix((len(slopes), nvar))
    bub = np.zeros(len(slopes))
    for k, (slope, intercept) in enumerate(zip(slopes, intercepts)):
        aub[k, oe + n] = float(slope)
        aub[k, oz] = -1.0
        bub[k] = -float(intercept)

    bounds: List[Tuple[Optional[float], Optional[float]]] = []
    for step in steps:
        bounds.append((0.0, None) if block[0] <= step < block[1] else (0.0, 0.0))
    for step in steps:
        bounds.append((0.0, float(q0[step])) if block[0] <= step < block[1] else (0.0, 0.0))
    bounds.extend([(0.0, core.STEP_LIMIT_KWH)] * n)
    bounds.extend([(0.0, core.STEP_LIMIT_KWH)] * n)
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(core.SOC_MIN_KWH, core.SOC_MAX_KWH)] * (n + 1))
    bounds.append((0.0, None))
    bounds[oe] = (start_soc, start_soc)

    started = time.perf_counter()
    result = linprog(
        objective,
        A_ub=aub.tocsr(),
        b_ub=bub,
        A_eq=aeq.tocsr(),
        b_eq=beq,
        bounds=bounds,
        method="highs",
    )
    elapsed = time.perf_counter() - started
    if not result.success:
        raise RuntimeError(f"Q3调整LP失败: {result.status} {result.message}")
    x = np.asarray(result.x)
    u, r = x[ou:ou + n], x[or_:or_ + n]
    candidate = q0.copy()
    candidate[steps] = q0[steps] + u - r
    residual = np.asarray(aeq.tocsr() @ x - beq)
    ub_residual = np.maximum(np.asarray(aub.tocsr() @ x - bub), 0.0)
    return {
        "candidate": candidate,
        "u": u,
        "r": r,
        "charge": x[oc:oc + n],
        "discharge": x[od:od + n],
        "planned_emergency": x[oh:oh + n],
        "soc": x[oe:oe + n + 1],
        "max_residual": float(max(np.abs(residual).max(), ub_residual.max())),
        "solve_seconds": float(elapsed),
        "objective_increment_plus_future_value": float(result.fun),
    }


def execute_net_step(soc: float, q: float, net_load: float) -> Dict[str, float]:
    return core.feedback_step(soc, q, max(net_load, 0.0), max(-net_load, 0.0))


def block_score(
    q: np.ndarray,
    q0: np.ndarray,
    block: Tuple[int, int],
    start_soc: float,
    forecast_net: np.ndarray,
    residual_scenarios: np.ndarray,
    price: np.ndarray,
) -> float:
    start, end = block
    u = np.maximum(q[start:end] - q0[start:end], 0.0)
    r = np.maximum(q0[start:end] - q[start:end], 0.0)
    ordinary = float(np.sum(price[start:end] * q0[start:end] + 1.5 * price[start:end] * u - 0.5 * price[start:end] * r))
    scenarios = residual_scenarios if len(residual_scenarios) else np.zeros((1, end - start))
    scores = []
    for residual in scenarios:
        soc = float(start_soc)
        emergency_cost = 0.0
        for local, step in enumerate(range(start, end)):
            action = execute_net_step(soc, float(q[step]), float(forecast_net[local] + residual[local]))
            soc = action["soc_after_kwh"]
            emergency_cost += 5.0 * price[step] * action["emergency_kwh"]
        lambda_asset = core.ETA_DISCHARGE * float(np.median(price))
        scores.append(ordinary + emergency_cost - lambda_asset * soc)
    return float(np.mean(scores))


def ledger(q0: np.ndarray, qfinal: np.ndarray, price: np.ndarray, emergency: np.ndarray) -> Dict[str, float]:
    u = np.maximum(qfinal - q0, 0.0)
    r = np.maximum(q0 - qfinal, 0.0)
    base = float(np.dot(price, q0))
    up = float(np.dot(1.5 * price, u))
    refund_b = float(np.dot(-0.5 * price, r))
    penalty_a = float(np.dot(0.5 * price, r))
    emergency_cost = float(np.dot(5.0 * price, emergency))
    return {
        "base_plan_cost_yuan": base,
        "up_adjustment_cost_yuan": up,
        "down_refund_B_yuan": refund_b,
        "down_penalty_A_yuan": penalty_a,
        "ordinary_cost_B_yuan": base + up + refund_b,
        "ordinary_cost_A_yuan": base + up + penalty_a,
        "emergency_cost_yuan": emergency_cost,
        "total_cash_B_yuan": base + up + refund_b + emergency_cost,
        "total_cash_A_revaluation_yuan": base + up + penalty_a + emergency_cost,
        "up_adjustment_kwh": float(u.sum()),
        "down_adjustment_kwh": float(r.sum()),
    }


def run_p1() -> Dict[str, object]:
    dataset = core.load_dataset()
    book = load_forecast_book(dataset)
    q2_config = q2full.StrategyConfig("R2", LOAD_MODEL, PV_MODEL_FOR_Q2_VALUE, ALPHA, "pwl")
    q2_library = q2full.build_forecast_library(dataset, [q2_config.combo])

    # Common January warm-up inherited from accepted Q2.  The row ending at
    # midnight is known; the 00:00-00:10 bridge is executed only after Q3 q0 is made.
    january_log, _, january_plans = q2full.simulate(dataset, 0, 30, q2_config, q2_library)
    midnight_row = january_log.loc[january_log["interval_start"] == TARGET_DATE - pd.Timedelta(minutes=10)].iloc[0]
    bridge_row = january_log.loc[january_log["interval_start"] == TARGET_DATE].iloc[0]
    soc_midnight = float(midnight_row["soc_after_kwh"])
    actual_start_soc = float(bridge_row["soc_after_kwh"])
    previous = january_plans[30]
    predicted_bridge = core.feedback_step(
        soc_midnight,
        float(previous.q[-1]),
        float(previous.forecast.load_kwh[-1]),
        float(previous.forecast.pv_kwh[-1]),
    )
    predicted_start_soc = float(predicted_bridge["soc_after_kwh"])

    day0 = build_release_forecast(dataset, book, q2_library, TARGET_INDEX, 0, use_load_correction=True)
    risk0, _, risk0_days, gate0_days = residual_pools(
        dataset, book, q2_library, TARGET_INDEX, 0, day0.steps,
        risk_days=core.MAX_RESIDUAL_DAYS, gate_days=0,
    )
    margin0 = np.quantile(risk0, ALPHA, axis=0) if len(risk0) >= MIN_POOL else np.zeros(core.N_STEPS)
    curve = q2full.next_day_value_curve(dataset, TARGET_INDEX, q2_config, q2_library)
    q0_solution = q2full.solve_plan(
        day0.net_kwh + margin0,
        dataset.price,
        predicted_start_soc,
        "pwl",
        pwl_curve=(curve["grid"], curve["values"]),
    )
    q0 = np.asarray(q0_solution["q"]).copy()
    qcurrent = q0.copy()
    qfinal = q0.copy()
    planned_emergency = np.zeros(core.N_STEPS)
    confirmation_count = np.zeros(core.N_STEPS, dtype=int)
    decisions: List[Dict[str, object]] = []
    records: List[Dict[str, object]] = []
    soc = actual_start_soc

    def execute_range(first: int, last: int) -> None:
        nonlocal soc
        for step in range(first, last):
            action = core.feedback_step(
                soc,
                float(qfinal[step]),
                float(dataset.load_kwh[TARGET_INDEX, step]),
                float(dataset.pv_kwh[TARGET_INDEX, step]),
            )
            start = dataset.interval_starts(TARGET_INDEX)[step]
            records.append({
                "interval_start": start,
                "interval_end": start + pd.Timedelta(minutes=10),
                "step": step,
                "template_label": dataset.template_labels[step],
                "q0_kwh": float(q0[step]),
                "qfinal_kwh": float(qfinal[step]),
                "load_actual_kwh": float(dataset.load_kwh[TARGET_INDEX, step]),
                "pv_actual_kwh": float(dataset.pv_kwh[TARGET_INDEX, step]),
                "charge_kwh": action["charge_kwh"],
                "discharge_kwh": action["discharge_kwh"],
                "actual_emergency_kwh": action["emergency_kwh"],
                "planned_emergency_kwh": float(planned_emergency[step]),
                "unused_supply_kwh": action["unused_supply_kwh"],
                "soc_before_kwh": soc,
                "soc_after_kwh": action["soc_after_kwh"],
                "price_yuan_per_kwh": float(dataset.price[step]),
            })
            soc = float(action["soc_after_kwh"])

    execute_range(0, 35)
    for release_hour in (6, 12, 18):
        block = BLOCKS[release_hour]
        forecast = build_release_forecast(
            dataset, book, q2_library, TARGET_INDEX, release_hour, use_load_correction=True
        )
        positions = np.flatnonzero((forecast.steps >= block[0]) & (forecast.steps < block[1]))
        block_steps = forecast.steps[positions]
        risk, gate, risk_days, gate_days = residual_pools(
            dataset, book, q2_library, TARGET_INDEX, release_hour, forecast.steps
        )
        margin = np.quantile(risk, ALPHA, axis=0) if len(risk) >= MIN_POOL else np.zeros(len(forecast.steps))
        solved = solve_adjustment_lp(
            q0,
            forecast,
            margin,
            soc,
            block,
            dataset.price,
            (curve["grid"], curve["values"]),
        )
        candidate = np.asarray(solved["candidate"]).copy()
        gate_block = gate[:, positions] if len(gate) else np.zeros((0, len(positions)))
        forecast_block = forecast.net_kwh[positions]
        keep_score = block_score(qcurrent, q0, block, soc, forecast_block, gate_block, dataset.price)
        new_score = block_score(candidate, q0, block, soc, forecast_block, gate_block, dataset.price)
        accepted = bool(len(gate_days) >= MIN_POOL and new_score < keep_score - 1e-9)
        before = qcurrent.copy()
        if accepted:
            qcurrent[block[0]:block[1]] = candidate[block[0]:block[1]]
            qfinal[block[0]:block[1]] = candidate[block[0]:block[1]]
            confirmation_count[block[0]:block[1]] += 1
            planned_emergency[forecast.steps] = np.asarray(solved["planned_emergency"])
        decisions.append({
            "release_hour": release_hour,
            "issue_time": forecast.issue_time,
            "block_start_step": block[0],
            "block_end_exclusive": block[1],
            "block_intervals": block[1] - block[0],
            "soc_at_release_kwh": soc,
            "anchor_kw": forecast.current_pv_anchor_kw,
            "anchor_source": forecast.anchor_source,
            "load_source_max": forecast.load_source_max,
            "load_correction_source_max": forecast.load_correction_source_max,
            "risk_days": len(risk_days),
            "risk_day_first": dataset.dates[risk_days[0]] if risk_days else pd.NaT,
            "risk_day_last": dataset.dates[risk_days[-1]] if risk_days else pd.NaT,
            "gate_days": len(gate_days),
            "gate_day_first": dataset.dates[gate_days[0]] if gate_days else pd.NaT,
            "gate_day_last": dataset.dates[gate_days[-1]] if gate_days else pd.NaT,
            "keep_score_yuan": keep_score,
            "new_score_yuan": new_score,
            "lambda_asset_yuan_per_kwh": core.ETA_DISCHARGE * float(np.median(dataset.price)),
            "accepted": accepted,
            "changed_intervals": int(np.count_nonzero(np.abs(qcurrent - before) > 1e-8)),
            "candidate_up_kwh": float(np.maximum(candidate - q0, 0).sum()),
            "candidate_down_kwh": float(np.maximum(q0 - candidate, 0).sum()),
            "lp_residual_kwh": solved["max_residual"],
            "planned_charge_discharge_overlap_count": int(np.count_nonzero(
                (np.asarray(solved["charge"]) > TOL)
                & (np.asarray(solved["discharge"]) > TOL)
            )),
            "solve_seconds": solved["solve_seconds"],
        })
        # The 18:00 block owns step 143=[next 00:00,next 00:10).  Stop at
        # midnight so that the next 00:00 plan is created before this bridge
        # interval's actual energy is revealed and executed.
        execute_range(block[0], block[1] if release_hour != 18 else block[1] - 1)

    bridge_contract_before_next_plan = float(qfinal[143])
    soc_next_midnight = float(soc)
    release18_forecast = forecast
    release18_margin = margin
    bridge_pos = int(np.flatnonzero(release18_forecast.steps == 143)[0])
    predicted_next_start = core.feedback_step(
        soc_next_midnight,
        bridge_contract_before_next_plan,
        float(release18_forecast.load_kwh[bridge_pos]),
        float(release18_forecast.pv_kwh[bridge_pos]),
    )["soc_after_kwh"]
    next_day_index = TARGET_INDEX + 1
    next_day_forecast = build_release_forecast(
        dataset, book, q2_library, next_day_index, 0, use_load_correction=True
    )
    next_risk, _, _, _ = residual_pools(
        dataset, book, q2_library, next_day_index, 0, next_day_forecast.steps,
        risk_days=core.MAX_RESIDUAL_DAYS, gate_days=0,
    )
    next_margin = np.quantile(next_risk, ALPHA, axis=0) if len(next_risk) >= MIN_POOL else np.zeros(core.N_STEPS)
    next_curve = q2full.next_day_value_curve(dataset, next_day_index, q2_config, q2_library)
    next_q0_solution = q2full.solve_plan(
        next_day_forecast.net_kwh + next_margin,
        dataset.price,
        float(predicted_next_start),
        "pwl",
        pwl_curve=(next_curve["grid"], next_curve["values"]),
    )
    next_plan_created_at = dataset.dates[next_day_index]
    bridge_contract_after_next_plan = float(qfinal[143])
    execute_range(143, 144)

    log = pd.DataFrame(records)
    decision_frame = pd.DataFrame(decisions)
    emergency = log["actual_emergency_kwh"].to_numpy(float)
    cash = ledger(q0, qfinal, dataset.price, emergency)

    # Hand-calculation ledger gate.
    hand_cases = []
    for original, final, expected_a, expected_b in [
        (100.0, 100.0, 100.0, 100.0),
        (100.0, 120.0, 130.0, 130.0),
        (100.0, 80.0, 110.0, 90.0),
        (0.0, 20.0, 30.0, 30.0),
    ]:
        calc = ledger(np.array([original]), np.array([final]), np.array([1.0]), np.array([0.0]))
        hand_cases.append(bool(abs(calc["ordinary_cost_A_yuan"] - expected_a) < TOL and abs(calc["ordinary_cost_B_yuan"] - expected_b) < TOL))

    # Three independent causality probes at 06:00.
    base6 = build_release_forecast(dataset, book, q2_library, TARGET_INDEX, 6, True)
    future_actual_data = copy.deepcopy(dataset)
    future_actual_data.load_kw[TARGET_INDEX, 35:] += 7777.0
    future_actual_data.load_kwh[TARGET_INDEX, 35:] += 7777.0 * core.DT_HOURS
    future_actual_data.pv_kw[TARGET_INDEX, 35:] += 5555.0
    future_actual_data.pv_kwh[TARGET_INDEX, 35:] += 5555.0 * core.DT_HOURS
    future_actual6 = build_release_forecast(future_actual_data, book, q2_library, TARGET_INDEX, 6, True)
    future_release_book = ForecastBook(
        book.dates.copy(), book.issue_hours.copy(), book.values_kw.copy(), book.anchors_kw.copy()
    )
    future_release_book.values_kw[4 * TARGET_INDEX + 2:, :] += 8888.0
    future_release6 = build_release_forecast(dataset, future_release_book, q2_library, TARGET_INDEX, 6, True)
    changed_anchor6 = build_release_forecast(
        dataset,
        book,
        q2_library,
        TARGET_INDEX,
        6,
        True,
        anchor_override_kw=base6.current_pv_anchor_kw + 100.0,
    )

    energy_residual = log["qfinal_kwh"] + log["pv_actual_kwh"] + log["discharge_kwh"] + log["actual_emergency_kwh"] - log["load_actual_kwh"] - log["charge_kwh"] - log["unused_supply_kwh"]
    soc_residual = log["soc_after_kwh"] - log["soc_before_kwh"] - core.ETA_CHARGE * log["charge_kwh"] + log["discharge_kwh"] / core.ETA_DISCHARGE
    audits = {
        "01_attachment3_365x4_finite": bool(len(book.dates) == 1460 and len(np.unique(book.dates)) == 365 and np.isfinite(book.values_kw).all()),
        "02_release_step_counts_144_109_73_37": bool([len(release_steps(h)) for h in RELEASE_HOURS] == [144, 109, 73, 37]),
        "03_commit_blocks_36_36_37": bool([b - a for a, b in BLOCKS.values()] == [36, 36, 37]),
        "04_q_arrays_independent": bool(not np.shares_memory(q0, qcurrent) and not np.shares_memory(q0, qfinal) and not np.shares_memory(qcurrent, qfinal)),
        "05_one_confirmation_per_interval": bool(confirmation_count.max() <= 1),
        "06_past_and_outside_blocks_frozen": bool(np.allclose(qfinal[:35], q0[:35])),
        "07_energy_balance": bool(float(np.abs(energy_residual).max()) <= TOL),
        "08_soc_balance": bool(float(np.abs(soc_residual).max()) <= TOL),
        "09_soc_bounds": bool(log["soc_after_kwh"].between(core.SOC_MIN_KWH - TOL, core.SOC_MAX_KWH + TOL).all()),
        "10_power_bounds": bool((log[["charge_kwh", "discharge_kwh"]].to_numpy() <= core.STEP_LIMIT_KWH + TOL).all()),
        "11_charge_discharge_mutex": bool(((log["charge_kwh"] <= TOL) | (log["discharge_kwh"] <= TOL)).all()),
        "12_all_adjustment_lps_feasible": bool((decision_frame["lp_residual_kwh"] <= TOL).all()),
        "13_history_sources_causal": bool(
            all(pd.isna(x) or pd.Timestamp(x) <= pd.Timestamp(t) for x, t in zip(decision_frame["load_source_max"], decision_frame["issue_time"]))
            and all(pd.isna(x) or pd.Timestamp(x) <= pd.Timestamp(t) for x, t in zip(decision_frame["load_correction_source_max"], decision_frame["issue_time"]))
        ),
        "14_future_actual_fault_invariant": bool(np.allclose(base6.net_kwh, future_actual6.net_kwh)),
        "15_future_release_fault_invariant": bool(np.allclose(base6.net_kwh, future_release6.net_kwh)),
        "16_current_anchor_is_separate_and_active": bool(base6.anchor_source == "frozen_current_point_measurement_proxy" and not np.allclose(base6.pv_kwh, changed_anchor6.pv_kwh)),
        "17_midnight_decision_precedes_bridge_actual": bool(day0.issue_time == TARGET_DATE and pd.Timestamp(bridge_row["interval_start"]) == TARGET_DATE),
        "18_bridge_contract_not_overwritten": bool(abs(float(previous.q[-1]) - float(bridge_row["plan_grid_kwh"])) <= TOL),
        "19_final_interval_is_next_day_bridge": bool(log.iloc[-1]["template_label"] == "0:00-0:10+1" and log.iloc[-1]["interval_start"] == TARGET_DATE + pd.Timedelta(days=1)),
        "20_ledger_hand_cases": bool(all(hand_cases)),
        "21_ledger_identity": bool(np.allclose(qfinal, q0 + np.maximum(qfinal - q0, 0.0) - np.maximum(q0 - qfinal, 0.0))),
        "22_planned_actual_emergency_separate": bool(not np.shares_memory(planned_emergency, emergency)),
        "23_144_execution_intervals": bool(len(log) == 144 and log["step"].tolist() == list(range(144))),
        "24_next_plan_before_bridge_and_no_overwrite": bool(
            next_plan_created_at == TARGET_DATE + pd.Timedelta(days=1)
            and abs(bridge_contract_after_next_plan - bridge_contract_before_next_plan) <= TOL
            and len(np.asarray(next_q0_solution["q"])) == 144
        ),
        "25_risk_gate_pools_nonoverlap": bool(
            all(
                pd.Timestamp(row.gate_day_last) < pd.Timestamp(row.risk_day_first)
                for row in decision_frame.itertuples()
                if pd.notna(row.gate_day_last) and pd.notna(row.risk_day_first)
            )
        ),
        "26_adjustment_plan_charge_discharge_mutex": bool(
            (decision_frame["planned_charge_discharge_overlap_count"] == 0).all()
        ),
        "28_midnight_residual_pool_excludes_unfinished_jan31_bridge": bool(
            len(risk0_days) >= MIN_POOL
            and max(risk0_days) <= 29
            and dataset.plan_end(max(risk0_days)) <= TARGET_DATE
        ),
    }
    status = "PASS" if all(audits.values()) else "FAIL"

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    log.to_csv(OUTPUT_DIR / "q3_p1_execution_log.csv", index=False, encoding="utf-8-sig")
    decision_frame.to_csv(OUTPUT_DIR / "q3_p1_decisions.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame({"step": np.arange(144), "q0_kwh": q0, "qfinal_kwh": qfinal}).to_csv(
        OUTPUT_DIR / "q3_p1_contracts.csv", index=False, encoding="utf-8-sig"
    )

    output_book = OUTPUT_DIR / "result3_p1_20250201.xlsx"
    shutil.copy2(TEMPLATE, output_book)
    wb = openpyxl.load_workbook(output_book)
    for sheet_name, values, cost in [
        ("计划购电量", q0, float(np.dot(dataset.price, q0))),
        ("调整购电量", qfinal, cash["ordinary_cost_B_yuan"]),
    ]:
        ws = wb[sheet_name]
        target_row = None
        for row in range(2, ws.max_row + 1):
            if pd.Timestamp(ws.cell(row, 1).value).normalize() == TARGET_DATE:
                target_row = row
                break
        if target_row is None:
            raise AssertionError(f"{sheet_name}找不到P1日期")
        for col, value in enumerate(values, start=2):
            ws.cell(target_row, col, float(value))
        ws.cell(target_row, 146, float(np.sum(values)))
        ws.cell(target_row, 147, float(cost))
    wb.save(output_book)
    check_wb = openpyxl.load_workbook(output_book, data_only=False, read_only=True)
    if check_wb.sheetnames != ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]:
        raise AssertionError("result3 P1工作表名称改变")
    readback = {}
    for sheet_name, expected, expected_cost in [
        ("计划购电量", q0, float(np.dot(dataset.price, q0))),
        ("调整购电量", qfinal, cash["ordinary_cost_B_yuan"]),
    ]:
        ws = check_wb[sheet_name]
        target_row = next(
            row for row in range(2, ws.max_row + 1)
            if pd.Timestamp(ws.cell(row, 1).value).normalize() == TARGET_DATE
        )
        actual = np.asarray([ws.cell(target_row, col).value for col in range(2, 146)], dtype=float)
        readback[sheet_name] = {
            "numeric_cells": int(np.isfinite(actual).sum()),
            "max_abs_contract_diff_kwh": float(np.max(np.abs(actual - expected))),
            "daily_amount_diff_kwh": float(abs(float(ws.cell(target_row, 146).value) - float(expected.sum()))),
            "daily_cost_diff_yuan": float(abs(float(ws.cell(target_row, 147).value) - expected_cost)),
        }
    workbook_readback_ok = bool(all(
        item["numeric_cells"] == 144
        and item["max_abs_contract_diff_kwh"] <= TOL
        and item["daily_amount_diff_kwh"] <= TOL
        and item["daily_cost_diff_yuan"] <= TOL
        for item in readback.values()
    ))
    audits["27_result3_p1_readback"] = workbook_readback_ok
    status = "PASS" if all(audits.values()) else "FAIL"

    result = {
        "status": status,
        "scope": "Q3 P1 one-real-day vertical slice; not annual result",
        "target_date": TARGET_DATE.date().isoformat(),
        "bridge": {
            "soc_midnight_kwh": soc_midnight,
            "predicted_start_soc_0010_kwh": predicted_start_soc,
            "actual_start_soc_0010_kwh": actual_start_soc,
            "previous_contract_bridge_kwh": float(previous.q[-1]),
        },
        "q0": {
            "purchase_kwh": float(q0.sum()),
            "cost_yuan": float(np.dot(dataset.price, q0)),
            "lp_residual_kwh": float(q0_solution["max_residual"]),
            "risk_days": len(risk0_days),
            "gate_days_reserved": len(gate0_days),
        },
        "decisions": json.loads(decision_frame.to_json(orient="records", date_format="iso")),
        "execution": {
            "final_purchase_kwh": float(qfinal.sum()),
            "actual_emergency_kwh": float(emergency.sum()),
            "unused_supply_kwh": float(log["unused_supply_kwh"].sum()),
            "end_soc_kwh": float(log.iloc[-1]["soc_after_kwh"]),
            "soc_min_kwh": float(log["soc_after_kwh"].min()),
            "soc_max_kwh": float(log["soc_after_kwh"].max()),
        },
        "ledger": cash,
        "audits": audits,
        "max_energy_residual_kwh": float(np.abs(energy_residual).max()),
        "max_soc_residual_kwh": float(np.abs(soc_residual).max()),
        "template_readback": readback,
        "manifest": {
            "inputs": {
                "attachment1": sha256(DATA_DIR / "附件1.xlsx"),
                "attachment2": sha256(DATA_DIR / "附件2.xlsx"),
                "attachment3": sha256(DATA_DIR / "附件3.xlsx"),
                "result3_template": sha256(TEMPLATE),
                "analysis": sha256(ROOT / "题目分析报告.md"),
                "terms": sha256(ROOT / "术语表格.md"),
            },
            "code": sha256(Path(__file__)),
            "python": platform.python_version(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
            "parameters": {
                "alpha": ALPHA,
                "load_model": LOAD_MODEL,
                "q2_value_pv_model": PV_MODEL_FOR_Q2_VALUE,
                "risk_days": RISK_DAYS,
                "gate_days": GATE_DAYS,
                "settlement_main": "B",
            },
            "command": f'"{sys.executable}" Q3/v2/program/q3_pipeline_v2.py --mode p1',
        },
    }
    (OUTPUT_DIR / "q3_p1_summary.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if status != "PASS":
        failed = [name for name, ok in audits.items() if not ok]
        raise AssertionError(f"Q3 P1审计失败: {failed}")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["p1"], default="p1")
    args = parser.parse_args()
    result = run_p1()
    print(json.dumps({
        "status": result["status"],
        "scope": result["scope"],
        "ledger": result["ledger"],
        "execution": result["execution"],
        "audits_passed": sum(result["audits"].values()),
        "audits_total": len(result["audits"]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
