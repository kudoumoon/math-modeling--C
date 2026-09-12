"""Full-year strict causal solution for Q4-2 and Q4-3.

The official evaluation window is 2025-02-01..2025-12-31.  Planning always
uses causal price forecasts; cash settlement always uses Attachment 4 actual
prices.  Invalid sources, solver failures, physical violations, and template
read-back mismatches are fatal.  There is no runtime model fallback.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import platform
import shutil
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
import pandas as pd
import scipy

PROGRAM_DIR = Path(__file__).resolve().parent
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROGRAM_DIR))
sys.path.insert(0, str(ROOT / "Q2" / "program"))
sys.path.insert(0, str(ROOT / "Q3" / "v2" / "program"))

import q2_pipeline_v1 as core
import q2_v2_compatible as q2full
import q3_pipeline_v2 as q3p1
import q3_full_pipeline_v2 as q3full
import q4_core as q4


OUT = ROOT / "Q4" / "result"
AUDIT = ROOT / "Q4" / "audit" / "p2"
TEMPLATE42 = ROOT / "附件" / "附件5" / "result4-2.xlsx"
TEMPLATE43 = ROOT / "附件" / "附件5" / "result4-3.xlsx"
FIRST_DAY = 31
LAST_DAY = 364
EVAL_DAYS = LAST_DAY - FIRST_DAY + 1
TOL = 1e-5
COMBO = ("same_weekday_2w", "recent_5d")
RELEASE_SETS = {"M0": (), "M1": (6,), "M2": (6, 12), "M3": (6, 12, 18)}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass
class Q43Plan:
    day: int
    q0: np.ndarray
    qfinal: np.ndarray
    predicted_start_soc: float
    actual_start_soc: float
    confirmation_count: np.ndarray
    q0_lp_residual: float
    release_forecasts: Dict[int, q3p1.ReleaseForecast]
    release_prices: Dict[int, q4.PriceForecast]


class ReleaseForecastCache:
    def __init__(self, dataset, book, library) -> None:
        self.dataset = dataset
        self.book = book
        self.library = library
        self.cache: Dict[Tuple[int, int], q3p1.ReleaseForecast] = {}

    def get(self, day: int, release: int) -> q3p1.ReleaseForecast:
        key = (int(day), int(release))
        if key not in self.cache:
            self.cache[key] = q3p1.build_release_forecast(
                self.dataset, self.book, self.library, day, release,
                use_load_correction=True,
            )
        return self.cache[key]


class DynamicCurveCache:
    """Cache an identical daily curve shared by Q4-2 and Q4-3 schedules."""

    def __init__(self, dataset, actual_price, library) -> None:
        self.dataset = dataset
        self.actual_price = actual_price
        self.library = library
        self.cache: Dict[int, Tuple[np.ndarray, np.ndarray]] = {}

    def get(self, day: int) -> Tuple[np.ndarray, np.ndarray]:
        if day >= LAST_DAY:
            raise ValueError("最后一天必须使用硬终端SOC，不能请求次日价值曲线")
        if day not in self.cache:
            curve = q4.dynamic_terminal_curve(
                self.dataset, self.actual_price, self.library, COMBO,
                day, self.dataset.dates[day],
            )
            self.cache[day] = (
                np.asarray(curve["grid"], dtype=float),
                np.asarray(curve["values"], dtype=float),
            )
        grid, values = self.cache[day]
        return grid.copy(), values.copy()


def residual_pools(
    dataset: core.Dataset,
    cache: ReleaseForecastCache,
    day: int,
    release: int,
    steps: np.ndarray,
    risk_days: int,
    gate_days: int,
) -> Tuple[np.ndarray, np.ndarray, List[int], List[int]]:
    issue = dataset.dates[day] + pd.Timedelta(hours=release)
    candidates: List[int] = []
    lookback = risk_days + gate_days + 2
    for old in range(max(0, day - lookback), day):
        latest_end = dataset.interval_starts(old)[int(np.max(steps))] + pd.Timedelta(minutes=10)
        if latest_end <= issue:
            candidates.append(old)
    residuals = []
    for old in candidates:
        old_forecast = cache.get(old, release)
        positions = np.searchsorted(old_forecast.steps, steps)
        if not np.array_equal(old_forecast.steps[positions], steps):
            raise AssertionError("历史预测步索引不能映射到当前提交块")
        predicted = old_forecast.net_kwh[positions]
        actual = dataset.load_kwh[old, steps] - dataset.pv_kwh[old, steps]
        residuals.append(actual - predicted)
    if not residuals:
        width = len(steps)
        return np.zeros((0, width)), np.zeros((0, width)), [], []
    matrix = np.vstack(residuals)
    nr = min(risk_days, len(matrix))
    risk = matrix[-nr:]
    gate_end = len(matrix) - nr
    gate_start = max(0, gate_end - gate_days)
    return risk, matrix[gate_start:gate_end], candidates[-nr:], candidates[gate_start:gate_end]


def margin_from_risk(risk: np.ndarray, width: int) -> np.ndarray:
    if len(risk) < q3p1.MIN_POOL:
        return np.zeros(width)
    return np.maximum(np.quantile(risk, q4.ALPHA, axis=0), 0.0)


def solve_q42_plan(
    dataset: core.Dataset,
    actual_price: np.ndarray,
    library,
    curves: DynamicCurveCache,
    day: int,
    predicted_start: float,
) -> q4.DayPlan:
    issue = dataset.dates[day]
    forecast = library[COMBO][day]
    margin, _, _ = q4.causal_margin_60(dataset, library, COMBO, issue)
    price = q4.build_price_forecast(dataset, actual_price, day, issue)
    if day == LAST_DAY:
        solved = q2full.solve_plan(
            forecast.net_kwh + margin, price.values, predicted_start,
            "hard", hard_anchor=core.INITIAL_SOC_KWH,
        )
    else:
        solved = q2full.solve_plan(
            forecast.net_kwh + margin, price.values, predicted_start,
            "pwl", pwl_curve=curves.get(day),
        )
    if float(solved["max_residual"]) > TOL:
        raise AssertionError(f"Q4-2计划LP残差超限 day={day}")
    return q4.DayPlan(
        day=day,
        decision_time=issue,
        predicted_start_soc=float(predicted_start),
        forecast=forecast,
        margin=margin,
        price_forecast=price,
        q=np.asarray(solved["q"], dtype=float),
        allow_storage=True,
        max_lp_residual=float(solved["max_residual"]),
    )


def simulate_q42(dataset, actual_price, library, curves, warm):
    records: List[Dict[str, object]] = []
    decisions: List[Dict[str, object]] = []
    plans: Dict[int, q4.DayPlan] = {}
    soc_midnight = float(warm.soc_midnight)
    previous = warm.previous
    for day in range(FIRST_DAY, LAST_DAY + 1):
        predicted_start = q4.predict_plan_start(soc_midnight, previous)
        plan = solve_q42_plan(dataset, actual_price, library, curves, day, predicted_start)

        bridge = q4.execute_step(dataset, actual_price, previous, core.N_STEPS - 1, soc_midnight)
        records.append(bridge)
        actual_start = float(bridge["soc_after_kwh"])
        soc = actual_start
        for step in range(core.N_STEPS - 1):
            row = q4.execute_step(dataset, actual_price, plan, step, soc)
            records.append(row)
            soc = float(row["soc_after_kwh"])
        plans[day] = plan
        decisions.append({
            "day_index": day,
            "date": dataset.dates[day],
            "decision_time": plan.decision_time,
            "predicted_start_soc_kwh": predicted_start,
            "actual_start_soc_kwh": actual_start,
            "price_max_source_time": plan.price_forecast.max_source_time,
            "load_source_max": plan.forecast.source_max_timestamp,
            "risk_history_days": len(q4.causal_margin_60(dataset, library, COMBO, plan.decision_time)[1]),
            "q0_lp_residual_kwh": plan.max_lp_residual,
        })
        previous = plan
        soc_midnight = soc
        if (day - FIRST_DAY + 1) % 30 == 0 or day == LAST_DAY:
            print(f"PROGRESS Q4-2 {dataset.dates[day].date()} plans={day-FIRST_DAY+1}", flush=True)
    final_bridge = q4.execute_step(dataset, actual_price, previous, core.N_STEPS - 1, soc_midnight)
    records.append(final_bridge)
    return (
        pd.DataFrame(records).sort_values("interval_start", kind="stable").reset_index(drop=True),
        pd.DataFrame(decisions),
        plans,
    )


def simulate_q43(dataset, actual_price, library, curves, cache, warm, schedule: str):
    allowed = RELEASE_SETS[schedule]
    records: List[Dict[str, object]] = []
    decisions: List[Dict[str, object]] = []
    plans: Dict[int, Q43Plan] = {}
    soc_midnight = float(warm.soc_midnight)
    previous: Dict[str, object] = {
        "day": warm.previous.day,
        "q0_bridge": float(warm.previous.q[-1]),
        "qfinal_bridge": float(warm.previous.q[-1]),
        "load_forecast_bridge": float(warm.previous.forecast.load_kwh[-1]),
        "pv_forecast_bridge": float(warm.previous.forecast.pv_kwh[-1]),
        "price_forecast_bridge": float(warm.previous.price_forecast.values[-1]),
        "price_max_source_time": warm.previous.price_forecast.max_source_time,
        "release_used": "Q4_January_warmup",
    }

    def append_execution(
        plan_day: int,
        step: int,
        q0_value: float,
        qfinal_value: float,
        soc_before: float,
        release_used: object,
        price_forecast_value: float,
        price_max_source_time: pd.Timestamp,
    ) -> float:
        action = core.feedback_step(
            soc_before, qfinal_value,
            float(dataset.load_kwh[plan_day, step]),
            float(dataset.pv_kwh[plan_day, step]),
        )
        start = dataset.interval_starts(plan_day)[step]
        price_actual = float(actual_price[plan_day, step])
        base = price_actual * q0_value
        up = 1.5 * price_actual * max(qfinal_value - q0_value, 0.0)
        refund = -0.5 * price_actual * max(q0_value - qfinal_value, 0.0)
        emergency_cost = 5.0 * price_actual * float(action["emergency_kwh"])
        records.append({
            "interval_start": start,
            "interval_end": start + pd.Timedelta(minutes=10),
            "plan_day": dataset.dates[plan_day],
            "natural_day": start.normalize(),
            "step": step,
            "template_label": dataset.template_labels[step],
            "schedule": schedule,
            "release_used": release_used,
            "price_forecast_yuan_per_kwh": price_forecast_value,
            "price_actual_yuan_per_kwh": price_actual,
            "price_max_source_time": price_max_source_time,
            "load_actual_kwh": float(dataset.load_kwh[plan_day, step]),
            "pv_actual_kwh": float(dataset.pv_kwh[plan_day, step]),
            "q0_kwh": q0_value,
            "qfinal_kwh": qfinal_value,
            "base_plan_cost_yuan": base,
            "up_adjustment_cost_yuan": up,
            "down_refund_B_yuan": refund,
            "ordinary_cost_B_yuan": base + up + refund,
            "emergency_kwh": float(action["emergency_kwh"]),
            "emergency_cost_yuan": emergency_cost,
            "charge_kwh": float(action["charge_kwh"]),
            "discharge_kwh": float(action["discharge_kwh"]),
            "unused_supply_kwh": float(action["unused_supply_kwh"]),
            "soc_before_kwh": float(soc_before),
            "soc_after_kwh": float(action["soc_after_kwh"]),
        })
        return float(action["soc_after_kwh"])

    for day in range(FIRST_DAY, LAST_DAY + 1):
        issue0 = dataset.dates[day]
        forecast0 = cache.get(day, 0)
        risk0, _, risk0_days, _ = residual_pools(
            dataset, cache, day, 0, forecast0.steps,
            risk_days=core.MAX_RESIDUAL_DAYS, gate_days=0,
        )
        margin0 = margin_from_risk(risk0, core.N_STEPS)
        price0 = q4.build_price_forecast(dataset, actual_price, day, issue0)
        predicted_start = float(core.feedback_step(
            soc_midnight,
            float(previous["qfinal_bridge"]),
            float(previous["load_forecast_bridge"]),
            float(previous["pv_forecast_bridge"]),
        )["soc_after_kwh"])
        if day == LAST_DAY:
            q0_solved = q2full.solve_plan(
                forecast0.net_kwh + margin0, price0.values, predicted_start,
                "hard", hard_anchor=core.INITIAL_SOC_KWH,
            )
            terminal_curve = None
        else:
            terminal_curve = curves.get(day)
            q0_solved = q2full.solve_plan(
                forecast0.net_kwh + margin0, price0.values, predicted_start,
                "pwl", pwl_curve=terminal_curve,
            )
        if float(q0_solved["max_residual"]) > TOL:
            raise AssertionError(f"Q4-3 q0 LP残差超限 day={day} schedule={schedule}")
        q0 = np.asarray(q0_solved["q"], dtype=float)
        qfinal = q0.copy()
        confirmations = np.zeros(core.N_STEPS, dtype=int)
        release_forecasts = {0: forecast0}
        release_prices = {0: price0}

        actual_start = append_execution(
            int(previous["day"]), 143,
            float(previous["q0_bridge"]), float(previous["qfinal_bridge"]),
            soc_midnight, previous["release_used"],
            float(previous["price_forecast_bridge"]),
            pd.Timestamp(previous["price_max_source_time"]),
        )
        soc = actual_start
        latest_forecast = forecast0
        latest_price = price0
        latest_release: object = 0

        def execute(start: int, end: int) -> None:
            nonlocal soc
            for step in range(start, end):
                pos = int(np.flatnonzero(latest_forecast.steps == step)[0])
                soc = append_execution(
                    day, step, float(q0[step]), float(qfinal[step]), soc,
                    latest_release, float(latest_price.values[step]),
                    latest_price.max_source_time,
                )

        execute(0, 35)
        for release in (6, 12, 18):
            block = q3p1.BLOCKS[release]
            if release in allowed:
                issue = dataset.dates[day] + pd.Timedelta(hours=release)
                forecast = cache.get(day, release)
                price = q4.build_price_forecast(dataset, actual_price, day, issue)
                latest_forecast = forecast
                latest_price = price
                latest_release = release
                release_forecasts[release] = forecast
                release_prices[release] = price
                risk, gate, risk_days, gate_days = residual_pools(
                    dataset, cache, day, release, forecast.steps,
                    risk_days=60, gate_days=14,
                )
                margin = margin_from_risk(risk, len(forecast.steps))
                if day == LAST_DAY:
                    solved = q3full.solve_adjustment(
                        q0, forecast, margin, soc, block, price.values,
                        "B", None, hard_anchor=core.INITIAL_SOC_KWH,
                    )
                    if bool(solved.get("fallback_keep_contract", False)):
                        raise AssertionError("B主模型不得触发任何solver fallback")
                else:
                    solved = q4.solve_adjustment_strict(
                        q0, forecast, margin, soc, block, price.values,
                        terminal_curve,
                    )
                if float(solved["max_residual"]) > TOL:
                    raise AssertionError(f"Q4-3调整LP残差超限 day={day} release={release}")
                candidate = np.asarray(solved["candidate"], dtype=float)
                positions = np.flatnonzero(
                    (forecast.steps >= block[0]) & (forecast.steps < block[1])
                )
                accepted = False
                keep_score = np.nan
                new_score = np.nan
                if len(gate_days) >= q3p1.MIN_POOL:
                    gate_block = gate[:, positions]
                    keep_score = q4.block_score_b(
                        qfinal, q0, block, soc, forecast.net_kwh[positions],
                        gate_block, price.values,
                    )
                    new_score = q4.block_score_b(
                        candidate, q0, block, soc, forecast.net_kwh[positions],
                        gate_block, price.values,
                    )
                    accepted = bool(new_score < keep_score - 1e-9)
                if accepted:
                    qfinal[block[0]:block[1]] = candidate[block[0]:block[1]]
                    confirmations[block[0]:block[1]] += 1
                decisions.append({
                    "day": dataset.dates[day],
                    "issue_time": issue,
                    "schedule": schedule,
                    "release_hour": release,
                    "block_start": block[0],
                    "block_end_exclusive": block[1],
                    "soc_kwh": soc,
                    "accepted": accepted,
                    "keep_score_yuan": keep_score,
                    "new_score_yuan": new_score,
                    "changed_intervals": int(np.count_nonzero(
                        np.abs(candidate[block[0]:block[1]] - q0[block[0]:block[1]]) > 1e-8
                    )) if accepted else 0,
                    "risk_days": len(risk_days),
                    "risk_day_first": dataset.dates[risk_days[0]] if risk_days else pd.NaT,
                    "risk_day_last": dataset.dates[risk_days[-1]] if risk_days else pd.NaT,
                    "gate_days": len(gate_days),
                    "gate_day_first": dataset.dates[gate_days[0]] if gate_days else pd.NaT,
                    "gate_day_last": dataset.dates[gate_days[-1]] if gate_days else pd.NaT,
                    "load_source_max": forecast.load_source_max,
                    "load_correction_source_max": forecast.load_correction_source_max,
                    "price_max_source_time": price.max_source_time,
                    "price_current_actual_step": price.current_actual_step,
                    "lp_residual_kwh": float(solved["max_residual"]),
                    "planned_mutex_violations": int(np.count_nonzero(
                        (np.asarray(solved["charge"]) > TOL)
                        & (np.asarray(solved["discharge"]) > TOL)
                    )),
                })
            execute(block[0], block[1] if release != 18 else 143)

        plans[day] = Q43Plan(
            day=day,
            q0=q0,
            qfinal=qfinal.copy(),
            predicted_start_soc=predicted_start,
            actual_start_soc=actual_start,
            confirmation_count=confirmations,
            q0_lp_residual=float(q0_solved["max_residual"]),
            release_forecasts=release_forecasts,
            release_prices=release_prices,
        )
        bridge_pos = int(np.flatnonzero(latest_forecast.steps == 143)[0])
        previous = {
            "day": day,
            "q0_bridge": float(q0[143]),
            "qfinal_bridge": float(qfinal[143]),
            "load_forecast_bridge": float(latest_forecast.load_kwh[bridge_pos]),
            "pv_forecast_bridge": float(latest_forecast.pv_kwh[bridge_pos]),
            "price_forecast_bridge": float(latest_price.values[143]),
            "price_max_source_time": latest_price.max_source_time,
            "release_used": latest_release,
        }
        soc_midnight = soc
        if (day - FIRST_DAY + 1) % 30 == 0 or day == LAST_DAY:
            print(f"PROGRESS Q4-3 {schedule} {dataset.dates[day].date()} plans={day-FIRST_DAY+1}", flush=True)

    append_execution(
        int(previous["day"]), 143,
        float(previous["q0_bridge"]), float(previous["qfinal_bridge"]),
        soc_midnight, previous["release_used"],
        float(previous["price_forecast_bridge"]),
        pd.Timestamp(previous["price_max_source_time"]),
    )
    return (
        pd.DataFrame(records).sort_values("interval_start", kind="stable").reset_index(drop=True),
        pd.DataFrame(decisions),
        plans,
    )


def plan_window(log: pd.DataFrame, dataset: core.Dataset) -> pd.DataFrame:
    start = dataset.dates[FIRST_DAY] + pd.Timedelta(minutes=10)
    end = dataset.dates[LAST_DAY] + pd.Timedelta(days=1, minutes=10)
    return log.loc[(log["interval_start"] >= start) & (log["interval_start"] < end)].copy()


def natural_window(log: pd.DataFrame, dataset: core.Dataset) -> pd.DataFrame:
    start = dataset.dates[FIRST_DAY]
    end = dataset.dates[LAST_DAY] + pd.Timedelta(days=1)
    return log.loc[(log["interval_start"] >= start) & (log["interval_start"] < end)].copy()


def summarize_q42(log, dataset) -> Dict[str, float]:
    frame = plan_window(log, dataset)
    return {
        "intervals": int(len(frame)),
        "plan_grid_kwh": float(frame["plan_grid_kwh"].sum()),
        "plan_cost_actual_yuan": float(frame["plan_cost_actual_yuan"].sum()),
        "emergency_kwh": float(frame["emergency_kwh"].sum()),
        "emergency_cost_actual_yuan": float(frame["emergency_cost_actual_yuan"].sum()),
        "total_cash_yuan": float((frame["plan_cost_actual_yuan"] + frame["emergency_cost_actual_yuan"]).sum()),
        "charge_kwh": float(frame["charge_kwh"].sum()),
        "discharge_kwh": float(frame["discharge_kwh"].sum()),
        "unused_supply_kwh": float(frame["unused_supply_kwh"].sum()),
        "start_soc_kwh": float(frame.iloc[0]["soc_before_kwh"]),
        "end_soc_kwh": float(frame.iloc[-1]["soc_after_kwh"]),
        "soc_min_kwh": float(frame["soc_after_kwh"].min()),
        "soc_max_kwh": float(frame["soc_after_kwh"].max()),
    }


def summarize_q43(log, dataset) -> Dict[str, float]:
    frame = plan_window(log, dataset)
    return {
        "intervals": int(len(frame)),
        "q0_kwh": float(frame["q0_kwh"].sum()),
        "qfinal_kwh": float(frame["qfinal_kwh"].sum()),
        "base_plan_cost_yuan": float(frame["base_plan_cost_yuan"].sum()),
        "up_adjustment_cost_yuan": float(frame["up_adjustment_cost_yuan"].sum()),
        "down_refund_B_yuan": float(frame["down_refund_B_yuan"].sum()),
        "ordinary_cost_B_yuan": float(frame["ordinary_cost_B_yuan"].sum()),
        "emergency_kwh": float(frame["emergency_kwh"].sum()),
        "emergency_cost_yuan": float(frame["emergency_cost_yuan"].sum()),
        "total_cash_B_yuan": float((frame["ordinary_cost_B_yuan"] + frame["emergency_cost_yuan"]).sum()),
        "charge_kwh": float(frame["charge_kwh"].sum()),
        "discharge_kwh": float(frame["discharge_kwh"].sum()),
        "unused_supply_kwh": float(frame["unused_supply_kwh"].sum()),
        "start_soc_kwh": float(frame.iloc[0]["soc_before_kwh"]),
        "end_soc_kwh": float(frame.iloc[-1]["soc_after_kwh"]),
        "soc_min_kwh": float(frame["soc_after_kwh"].min()),
        "soc_max_kwh": float(frame["soc_after_kwh"].max()),
    }


def physical_audit(log, dataset, q43: bool) -> Dict[str, object]:
    frame = plan_window(log, dataset).sort_values("interval_start", kind="stable")
    qcol = "qfinal_kwh" if q43 else "plan_grid_kwh"
    ecol = "emergency_kwh"
    balance = (
        frame[qcol] + frame["pv_actual_kwh"] + frame["discharge_kwh"] + frame[ecol]
        - frame["load_actual_kwh"] - frame["charge_kwh"] - frame["unused_supply_kwh"]
    ).to_numpy(float)
    state = (
        frame["soc_after_kwh"] - frame["soc_before_kwh"]
        - core.ETA_CHARGE * frame["charge_kwh"]
        + frame["discharge_kwh"] / core.ETA_DISCHARGE
    ).to_numpy(float)
    starts = pd.DatetimeIndex(frame["interval_start"])
    ends = pd.DatetimeIndex(frame["interval_end"])
    continuity = np.abs(
        frame["soc_before_kwh"].to_numpy(float)[1:]
        - frame["soc_after_kwh"].to_numpy(float)[:-1]
    )
    checks = {
        "01_plan_window_334x144": len(frame) == EVAL_DAYS * core.N_STEPS,
        "02_ten_minute_intervals": bool(np.all((ends - starts).total_seconds() == 600)),
        "03_no_gap_or_duplicate": bool(not starts.duplicated().any() and np.all(starts[1:].asi8 == ends[:-1].asi8)),
        "04_energy_balance": float(np.max(np.abs(balance))) <= TOL,
        "05_soc_balance": float(np.max(np.abs(state))) <= TOL,
        "06_soc_continuity": float(np.max(continuity)) <= TOL,
        "07_soc_bounds": bool(frame[["soc_before_kwh", "soc_after_kwh"]].to_numpy().min() >= core.SOC_MIN_KWH - TOL and frame[["soc_before_kwh", "soc_after_kwh"]].to_numpy().max() <= core.SOC_MAX_KWH + TOL),
        "08_power_bounds": bool(frame[["charge_kwh", "discharge_kwh"]].to_numpy().max() <= core.STEP_LIMIT_KWH + TOL),
        "09_charge_discharge_mutex": bool(not ((frame["charge_kwh"] > TOL) & (frame["discharge_kwh"] > TOL)).any()),
        "10_emergency_not_charging": bool(not ((frame[ecol] > TOL) & (frame["charge_kwh"] > TOL)).any()),
        "11_nonnegative_contracts": bool(frame[qcol].min() >= -TOL and frame[ecol].min() >= -TOL),
    }
    return {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "checks": checks,
        "metrics": {
            "max_energy_residual_kwh": float(np.max(np.abs(balance))),
            "max_soc_residual_kwh": float(np.max(np.abs(state))),
            "max_soc_continuity_error_kwh": float(np.max(continuity)),
        },
    }


def natural_blocks(log: pd.DataFrame, dataset: core.Dataset) -> pd.DataFrame:
    frame = natural_window(log, dataset)
    labels = ("0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00")
    rows = []
    for day in dataset.dates[FIRST_DAY:LAST_DAY + 1]:
        one = frame.loc[(frame["interval_start"] >= day) & (frame["interval_start"] < day + pd.Timedelta(days=1))].sort_values("interval_start")
        if len(one) != core.N_STEPS:
            raise AssertionError(f"自然日储能块不完整 {day.date()} rows={len(one)}")
        for block, label in enumerate(labels):
            part = one.iloc[block * 24:(block + 1) * 24]
            rows.append({
                "natural_day": day,
                "block": label,
                "charge_kwh": float(part["charge_kwh"].sum()),
                "discharge_kwh": float(part["discharge_kwh"].sum()),
            })
    return pd.DataFrame(rows)


def emergency_events(log: pd.DataFrame, dataset: core.Dataset) -> pd.DataFrame:
    frame = natural_window(log, dataset).rename(columns={"emergency_kwh": "emergency_grid_kwh"})
    events = core.merge_emergency_events(frame)
    if events.empty:
        return pd.DataFrame(columns=["date", "time_period", "emergency_kwh"])
    events["date"] = pd.to_datetime(events["start"]).dt.normalize()
    events["time_period"] = [
        core.format_event_time(pd.Timestamp(start), pd.Timestamp(end))
        for start, end in zip(events["start"], events["end"])
    ]
    return events[["date", "time_period", "emergency_kwh"]]


def copy_style(source, target) -> None:
    target._style = copy.copy(source._style)
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)
    target.number_format = source.number_format


def write_storage_and_emergency(wb, log, blocks, events, dataset) -> None:
    days = dataset.dates[FIRST_DAY:LAST_DAY + 1]
    storage = wb["充放电量"]
    styles = [[copy.copy(storage.cell(r, c)) for c in range(1, 7)] for r in range(2, 8)]
    storage.delete_rows(2, storage.max_row - 1)
    natural = natural_window(log, dataset)
    for day_offset, day in enumerate(days):
        part = blocks.loc[blocks["natural_day"] == day].sort_values("block", kind="stable")
        one = natural.loc[(natural["interval_start"] >= day) & (natural["interval_start"] < day + pd.Timedelta(days=1))].sort_values("interval_start")
        if len(part) != 6 or len(one) != 144:
            raise AssertionError(f"写模板时自然日不完整 {day.date()}")
        # Preserve the prescribed block order instead of lexical sorting.
        part = blocks.loc[blocks["natural_day"] == day]
        for block_index, (_, item) in enumerate(part.iterrows()):
            row = 2 + day_offset * 6 + block_index
            for col in range(1, 7):
                copy_style(styles[block_index][col - 1], storage.cell(row, col))
            storage.cell(row, 1, day.to_pydatetime() if block_index == 0 else None)
            storage.cell(row, 2, item["block"])
            storage.cell(row, 3, float(item["charge_kwh"]))
            storage.cell(row, 4, float(item["discharge_kwh"]))
            storage.cell(row, 5, "0:00" if block_index == 0 else ("24:00" if block_index == 1 else None))
            storage.cell(row, 6, float(one.iloc[0]["soc_before_kwh"]) if block_index == 0 else (float(one.iloc[-1]["soc_after_kwh"]) if block_index == 1 else None))

    emergency = wb["紧急购电量"]
    emergency_styles = [[copy.copy(emergency.cell(r, c)) for c in range(1, 4)] for r in range(2, 5)]
    emergency.delete_rows(2, emergency.max_row - 1)
    rows = events.copy().reset_index(drop=True)
    if rows.empty:
        rows = pd.DataFrame([{"date": days[0], "time_period": None, "emergency_kwh": 0.0}])
    previous_date = None
    for index, item in rows.iterrows():
        row = 2 + index
        for col in range(1, 4):
            copy_style(emergency_styles[index % 3][col - 1], emergency.cell(row, col))
        date = pd.Timestamp(item["date"])
        emergency.cell(row, 1, date.to_pydatetime() if date != previous_date else None)
        emergency.cell(row, 2, item["time_period"])
        emergency.cell(row, 3, float(item["emergency_kwh"]))
        previous_date = date


def write_result42(dataset, actual_price, plans, log) -> Tuple[Path, Dict[str, object]]:
    output = OUT / "result4-2.xlsx"
    shutil.copy2(TEMPLATE42, output)
    wb = openpyxl.load_workbook(output)
    ws = wb["计划购电量"]
    for row, day in enumerate(range(FIRST_DAY, LAST_DAY + 1), start=2):
        q = plans[day].q
        ws.cell(row, 1, dataset.dates[day].to_pydatetime())
        for col, value in enumerate(q, start=2):
            ws.cell(row, col, float(value))
        ws.cell(row, 146, float(q.sum()))
        ws.cell(row, 147, float(np.dot(actual_price[day], q)))
    blocks = natural_blocks(log, dataset)
    events = emergency_events(log, dataset)
    write_storage_and_emergency(wb, log, blocks, events, dataset)
    wb.save(output)
    check = openpyxl.load_workbook(output, data_only=False)
    values = np.asarray([[check["计划购电量"].cell(r, c).value for c in range(2, 146)] for r in range(2, 336)], dtype=float)
    expected = np.vstack([plans[d].q for d in range(FIRST_DAY, LAST_DAY + 1)])
    costs = np.asarray([check["计划购电量"].cell(r, 147).value for r in range(2, 336)], dtype=float)
    expected_costs = np.asarray([np.dot(actual_price[d], plans[d].q) for d in range(FIRST_DAY, LAST_DAY + 1)])
    readback = {
        "sheet_names": check.sheetnames,
        "plan_rows": check["计划购电量"].max_row - 1,
        "storage_rows": check["充放电量"].max_row - 1,
        "placeholder_count": sum(1 for sheet in check.worksheets for row in sheet.iter_rows() for cell in row if cell.value == "⁝"),
        "plan_max_diff_kwh": float(np.max(np.abs(values - expected))),
        "plan_cost_max_diff_yuan": float(np.max(np.abs(costs - expected_costs))),
    }
    return output, readback


def write_result43(dataset, actual_price, plans, log) -> Tuple[Path, Dict[str, object]]:
    output = OUT / "result4-3.xlsx"
    shutil.copy2(TEMPLATE43, output)
    wb = openpyxl.load_workbook(output)
    plan_ws = wb["计划购电量"]
    final_ws = wb["调整购电量"]
    for row, day in enumerate(range(FIRST_DAY, LAST_DAY + 1), start=2):
        plan = plans[day]
        ledger = q4.ledger_b(plan.q0, plan.qfinal, actual_price[day], np.zeros(core.N_STEPS))
        for ws, values, cost in (
            (plan_ws, plan.q0, ledger["base_plan_cost_yuan"]),
            (final_ws, plan.qfinal, ledger["ordinary_cost_B_yuan"]),
        ):
            ws.cell(row, 1, dataset.dates[day].to_pydatetime())
            for col, value in enumerate(values, start=2):
                ws.cell(row, col, float(value))
            ws.cell(row, 146, float(values.sum()))
            ws.cell(row, 147, float(cost))
    blocks = natural_blocks(log, dataset)
    events = emergency_events(log, dataset)
    write_storage_and_emergency(wb, log, blocks, events, dataset)
    wb.save(output)
    check = openpyxl.load_workbook(output, data_only=False)
    read0 = np.asarray([[check["计划购电量"].cell(r, c).value for c in range(2, 146)] for r in range(2, 336)], dtype=float)
    readf = np.asarray([[check["调整购电量"].cell(r, c).value for c in range(2, 146)] for r in range(2, 336)], dtype=float)
    expected0 = np.vstack([plans[d].q0 for d in range(FIRST_DAY, LAST_DAY + 1)])
    expectedf = np.vstack([plans[d].qfinal for d in range(FIRST_DAY, LAST_DAY + 1)])
    readback = {
        "sheet_names": check.sheetnames,
        "plan_rows": check["计划购电量"].max_row - 1,
        "adjust_rows": check["调整购电量"].max_row - 1,
        "storage_rows": check["充放电量"].max_row - 1,
        "placeholder_count": sum(1 for sheet in check.worksheets for row in sheet.iter_rows() for cell in row if cell.value == "⁝"),
        "plan_max_diff_kwh": float(np.max(np.abs(read0 - expected0))),
        "adjust_max_diff_kwh": float(np.max(np.abs(readf - expectedf))),
    }
    return output, readback


def release_price_sensitivity(plans, actual_price, emergency_cost: float) -> Dict[str, float]:
    base = up = refund = 0.0
    for day, plan in plans.items():
        base += float(np.dot(actual_price[day], plan.q0))
        for release, (start, end) in q3p1.BLOCKS.items():
            issue_step = q4.RELEASE_STEP[release]
            issue_price = float(actual_price[day, issue_step])
            u = np.maximum(plan.qfinal[start:end] - plan.q0[start:end], 0.0)
            r = np.maximum(plan.q0[start:end] - plan.qfinal[start:end], 0.0)
            up += 1.5 * issue_price * float(u.sum())
            refund += -0.5 * issue_price * float(r.sum())
    return {
        "base_plan_cost_yuan": base,
        "up_adjustment_cost_yuan": up,
        "down_refund_B_yuan": refund,
        "emergency_cost_yuan": float(emergency_cost),
        "total_cash_release_price_yuan": float(base + up + refund + emergency_cost),
    }


def audit_q42(dataset, log, decisions, plans, readback) -> Dict[str, object]:
    physical = physical_audit(log, dataset, q43=False)
    source_ok = bool((pd.to_datetime(decisions["price_max_source_time"]) <= pd.to_datetime(decisions["decision_time"])).all())
    load_ok = bool(all(
        pd.isna(row.load_source_max) or pd.Timestamp(row.load_source_max) <= pd.Timestamp(row.decision_time)
        for row in decisions.itertuples()
    ))
    checks = {
        **physical["checks"],
        "12_price_sources_causal": source_ok,
        "13_load_pv_sources_causal": load_ok,
        "14_plan_lp_residual": max(plan.max_lp_residual for plan in plans.values()) <= TOL,
        "15_template_sheets": readback["sheet_names"] == ["计划购电量", "充放电量", "紧急购电量"],
        "16_template_shape": readback["plan_rows"] == EVAL_DAYS and readback["storage_rows"] == EVAL_DAYS * 6,
        "17_template_no_placeholder": readback["placeholder_count"] == 0,
        "18_template_values": readback["plan_max_diff_kwh"] <= TOL and readback["plan_cost_max_diff_yuan"] <= TOL,
    }
    return {"status": "PASS" if all(checks.values()) else "FAIL", "checks": checks, "physical_metrics": physical["metrics"], "template_readback": readback}


def audit_q43(dataset, actual_price, log, decisions, plans, schedule: str, readback=None) -> Dict[str, object]:
    physical = physical_audit(log, dataset, q43=True)
    expected_decisions = EVAL_DAYS * len(RELEASE_SETS[schedule])
    source_ok = True
    load_ok = True
    pool_ok = True
    if len(decisions):
        source_ok = bool((pd.to_datetime(decisions["price_max_source_time"]) <= pd.to_datetime(decisions["issue_time"])).all())
        load_ok = bool(all(
            (pd.isna(row.load_source_max) or pd.Timestamp(row.load_source_max) <= pd.Timestamp(row.issue_time))
            and (pd.isna(row.load_correction_source_max) or pd.Timestamp(row.load_correction_source_max) <= pd.Timestamp(row.issue_time))
            for row in decisions.itertuples()
        ))
        pool_ok = bool(all(
            pd.isna(row.gate_day_last) or pd.isna(row.risk_day_first)
            or pd.Timestamp(row.gate_day_last) < pd.Timestamp(row.risk_day_first)
            for row in decisions.itertuples()
        ))
    frame = plan_window(log, dataset)
    u = np.maximum(frame["qfinal_kwh"] - frame["q0_kwh"], 0.0)
    r = np.maximum(frame["q0_kwh"] - frame["qfinal_kwh"], 0.0)
    p = frame["price_actual_yuan_per_kwh"]
    ledger_error = max(
        float(np.max(np.abs(frame["base_plan_cost_yuan"] - p * frame["q0_kwh"]))),
        float(np.max(np.abs(frame["up_adjustment_cost_yuan"] - 1.5 * p * u))),
        float(np.max(np.abs(frame["down_refund_B_yuan"] + 0.5 * p * r))),
    )
    checks = {
        **physical["checks"],
        "12_expected_decision_count": len(decisions) == expected_decisions,
        "13_price_sources_causal": source_ok,
        "14_load_pv_sources_causal": load_ok,
        "15_risk_gate_disjoint": pool_ok,
        "16_adjustment_lp_residual": len(decisions) == 0 or float(decisions["lp_residual_kwh"].max()) <= TOL,
        "17_planned_mutex": len(decisions) == 0 or int(decisions["planned_mutex_violations"].sum()) == 0,
        "18_q0_lp_residual": max(plan.q0_lp_residual for plan in plans.values()) <= TOL,
        "19_contract_arrays_separate": all(not np.shares_memory(plan.q0, plan.qfinal) for plan in plans.values()),
        "20_one_confirmation_per_interval": all(plan.confirmation_count.max() <= 1 for plan in plans.values()),
        "21_early_block_frozen": all(np.allclose(plan.q0[:35], plan.qfinal[:35]) for plan in plans.values()),
        "22_actual_price_ledger": ledger_error <= TOL,
        "23_price_channels_separate": "price_forecast_yuan_per_kwh" in frame and "price_actual_yuan_per_kwh" in frame,
    }
    if readback is not None:
        checks.update({
            "24_template_sheets": readback["sheet_names"] == ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
            "25_template_shape": readback["plan_rows"] == EVAL_DAYS and readback["adjust_rows"] == EVAL_DAYS and readback["storage_rows"] == EVAL_DAYS * 6,
            "26_template_no_placeholder": readback["placeholder_count"] == 0,
            "27_template_values": readback["plan_max_diff_kwh"] <= TOL and readback["adjust_max_diff_kwh"] <= TOL,
        })
    return {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "checks": checks,
        "physical_metrics": physical["metrics"],
        "ledger_max_error_yuan": ledger_error,
        "template_readback": readback,
    }


def run_full() -> Dict[str, object]:
    OUT.mkdir(parents=True, exist_ok=True)
    AUDIT.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    dataset = core.load_dataset()
    actual_price = q4.load_actual_prices(dataset)
    library = q2full.build_forecast_library(dataset, [COMBO])
    warm = q4.build_warm_context(30)
    curves = DynamicCurveCache(dataset, actual_price, library)

    log42, decisions42, plans42 = simulate_q42(dataset, actual_price, library, curves, warm)
    summary42 = summarize_q42(log42, dataset)
    log42.to_csv(OUT / "q4_2_execution_log.csv", index=False, encoding="utf-8-sig")
    decisions42.to_csv(OUT / "q4_2_daily_decisions.csv", index=False, encoding="utf-8-sig")
    workbook42, readback42 = write_result42(dataset, actual_price, plans42, log42)
    audit42 = audit_q42(dataset, log42, decisions42, plans42, readback42)
    if audit42["status"] != "PASS":
        raise AssertionError([key for key, ok in audit42["checks"].items() if not ok])
    print(f"PASS Q4-2 total={summary42['total_cash_yuan']:.6f}", flush=True)

    book = q3p1.load_forecast_book(dataset)
    forecast_cache = ReleaseForecastCache(dataset, book, library)
    schedule_summaries: Dict[str, Dict[str, float]] = {}
    schedule_audits: Dict[str, Dict[str, object]] = {}
    schedule_logs: Dict[str, pd.DataFrame] = {}
    schedule_decisions: Dict[str, pd.DataFrame] = {}
    schedule_plans: Dict[str, Dict[int, Q43Plan]] = {}
    for schedule in ("M0", "M1", "M2", "M3"):
        log, decisions, plans = simulate_q43(
            dataset, actual_price, library, curves, forecast_cache, warm, schedule
        )
        summary = summarize_q43(log, dataset)
        audit = audit_q43(dataset, actual_price, log, decisions, plans, schedule)
        if audit["status"] != "PASS":
            raise AssertionError(f"{schedule} audit failed: {[key for key, ok in audit['checks'].items() if not ok]}")
        log.to_csv(OUT / f"q4_3_{schedule}_execution_log.csv", index=False, encoding="utf-8-sig")
        decisions.to_csv(OUT / f"q4_3_{schedule}_decisions.csv", index=False, encoding="utf-8-sig")
        schedule_summaries[schedule] = summary
        schedule_audits[schedule] = audit
        schedule_logs[schedule] = log
        schedule_decisions[schedule] = decisions
        schedule_plans[schedule] = plans
        print(f"PASS Q4-3 {schedule} total={summary['total_cash_B_yuan']:.6f}", flush=True)

    # Q3-v2 M3 is the frozen official baseline; no post-hoc strategy switching.
    log43 = schedule_logs["M3"]
    decisions43 = schedule_decisions["M3"]
    plans43 = schedule_plans["M3"]
    workbook43, readback43 = write_result43(dataset, actual_price, plans43, log43)
    audit43 = audit_q43(dataset, actual_price, log43, decisions43, plans43, "M3", readback43)
    if audit43["status"] != "PASS":
        raise AssertionError([key for key, ok in audit43["checks"].items() if not ok])

    sensitivity = release_price_sensitivity(
        plans43, actual_price, schedule_summaries["M3"]["emergency_cost_yuan"]
    )
    comparison = pd.DataFrame([
        {"strategy": name, **metrics,
         "saving_vs_M0_yuan": schedule_summaries["M0"]["total_cash_B_yuan"] - metrics["total_cash_B_yuan"]}
        for name, metrics in schedule_summaries.items()
    ])
    comparison.to_csv(OUT / "q4_3_M0_M3_comparison.csv", index=False, encoding="utf-8-sig")

    final_checks = {
        "q4_2_audit": audit42["status"] == "PASS",
        "q4_3_all_schedules_audit": all(value["status"] == "PASS" for value in schedule_audits.values()),
        "q4_3_official_M3_template_audit": audit43["status"] == "PASS",
        "same_warm_start": abs(float(plans42[FIRST_DAY].predicted_start_soc) - float(plans43[FIRST_DAY].predicted_start_soc)) <= TOL,
        "official_files_exist": workbook42.exists() and workbook43.exists(),
    }
    result = {
        "status": "PASS" if all(final_checks.values()) else "FAIL",
        "evaluation_window": {"first_day": "2025-02-01", "last_day": "2025-12-31", "days": EVAL_DAYS, "intervals": EVAL_DAYS * core.N_STEPS},
        "price_model": "causal lag-7 same-slot; actual current slot revealed at 06/12/18 only",
        "q4_2": summary42,
        "q4_3": schedule_summaries,
        "q4_3_official": "M3",
        "release_time_price_sensitivity": sensitivity,
        "checks": final_checks,
        "audits": {"q4_2": audit42, "q4_3": {**schedule_audits, "M3_template": audit43}},
        "runtime_seconds": time.perf_counter() - started,
        "manifest": {
            "inputs": {
                "attachment1": sha256(ROOT / "附件" / "附件1.xlsx"),
                "attachment2": sha256(ROOT / "附件" / "附件2.xlsx"),
                "attachment3": sha256(ROOT / "附件" / "附件3.xlsx"),
                "attachment4": sha256(ROOT / "附件" / "附件4.xlsx"),
                "template4_2": sha256(TEMPLATE42),
                "template4_3": sha256(TEMPLATE43),
            },
            "upstream_q3_v2": {
                "pipeline": sha256(ROOT / "Q3" / "v2" / "program" / "q3_pipeline_v2.py"),
                "M3_decisions": sha256(ROOT / "Q3" / "v2" / "result" / "q3_M3_decisions.csv"),
                "M3_execution": sha256(ROOT / "Q3" / "v2" / "result" / "q3_M3_execution_log.csv"),
                "summary": sha256(ROOT / "Q3" / "v2" / "result" / "q3_final_summary.json"),
            },
            "code": {"q4_core": sha256(PROGRAM_DIR / "q4_core.py"), "q4_p1": sha256(PROGRAM_DIR / "q4_p1.py"), "q4_full": sha256(Path(__file__))},
            "python": platform.python_version(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
            "command": "python Q4/program/q4_full.py --mode full",
        },
    }
    (OUT / "q4_final_summary.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    (AUDIT / "q4_p2_audit.json").write_text(json.dumps({"status": result["status"], "checks": final_checks, "audits": result["audits"]}, ensure_ascii=False, indent=2), encoding="utf-8")
    if result["status"] != "PASS":
        raise AssertionError([key for key, ok in final_checks.items() if not ok])
    print(json.dumps({
        "status": result["status"],
        "q4_2_total_yuan": summary42["total_cash_yuan"],
        "q4_3_totals_yuan": {key: value["total_cash_B_yuan"] for key, value in schedule_summaries.items()},
        "official": "M3",
        "runtime_seconds": result["runtime_seconds"],
    }, ensure_ascii=False, indent=2), flush=True)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("full",), default="full")
    parser.parse_args()
    run_full()


if __name__ == "__main__":
    main()
