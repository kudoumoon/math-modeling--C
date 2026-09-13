"""Output-only tables, template readback and figures; no policy reruns."""
from __future__ import annotations
import json
from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd
from q3_api import ROOT, TERMINAL_VALUE

FOCUS = ('2025-03-20','2025-06-21','2025-09-23','2025-12-21')
CASH = ['base_cost_yuan','up_cost_yuan','down_cost_yuan','ordinary_cost_yuan','emergency_cost_yuan','total_cost_yuan']
ENERGY = ['q0_kwh','qfinal_kwh','u_kwh','r_kwh','charge_kwh','discharge_kwh','emergency_kwh','unused_supply_kwh']


def write_json(path, value):
    Path(path).write_text(json.dumps(value,indent=2,ensure_ascii=False,allow_nan=False)+'\n',encoding='utf-8')


def daily_metrics(f):
    g=f.groupby('plan_day',sort=True)
    d=g[CASH+ENERGY].sum()
    d['initial_soc_kwh'],d['end_soc_kwh']=g.soc_before_kwh.first(),g.soc_after_kwh.last()
    d['asset_adjusted_cost_yuan']=d.total_cost_yuan+TERMINAL_VALUE*(d.initial_soc_kwh-d.end_soc_kwh)
    d['interval_count']=g.size()
    return d.reset_index()


def emergency_events(f):
    rows,active=[],None
    for x in f.itertuples():
        if x.emergency_kwh<=1e-7:
            active=None
            continue
        date=x.interval_start.normalize()
        if active is not None and active['interval_end']==x.interval_start and active['date']==date:
            active['interval_end']=x.interval_end
            active['emergency_kwh']+=x.emergency_kwh
        else:
            active=dict(date=date,interval_start=x.interval_start,interval_end=x.interval_end,emergency_kwh=x.emergency_kwh)
            rows.append(active)
    out=pd.DataFrame(rows,columns=['date','interval_start','interval_end','emergency_kwh'])
    out['time_period']=[a.strftime('%H:%M')+'-'+('24:00' if b.normalize()>a.normalize() else b.strftime('%H:%M')) for a,b in zip(out.interval_start,out.interval_end)]
    return out


def write_workbook(path,f,blocks,events):
    wb=openpyxl.load_workbook(ROOT/'data/附件5/result3.xlsx')
    for ws in wb:
        for m in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(m))
        ws.delete_rows(2,ws.max_row)
    expected={ws.title:[] for ws in wb}
    def append(name,values):
        wb[name].append(values)
        expected[name].append(values)
    for date,part in f.groupby('plan_day',sort=True):
        part=part.sort_values('time_index')
        if part.time_index.tolist()!=list(range(144)):
            raise ValueError('workbook needs complete planning days')
        for name,key,cost in [('计划购电量','q0_kwh','base_cost_yuan'),('调整购电量','qfinal_kwh','ordinary_cost_yuan')]:
            append(name,[date.to_pydatetime(),*part[key].astype(float),float(part[key].sum()),float(part[cost].sum())])
        for k,x in enumerate(blocks[blocks.plan_day==date].itertuples()):
            append('充放电量',[date.to_pydatetime() if k==0 else None,f'{k*4}:00-{(k+1)*4}:00',
                float(x.charge_kwh),float(x.discharge_kwh),'0:00' if k==0 else ('24:00' if k==1 else None),
                float(part.iloc[0].soc_before_kwh) if k==0 else (float(part.iloc[-1].soc_after_kwh) if k==1 else None)])
    for x in events.itertuples():
        append('紧急购电量',[x.date.to_pydatetime(),x.time_period,float(x.emergency_kwh)])
    from openpyxl.comments import Comment
    wb['充放电量'].cell(1,6).comment=Comment('Time A: nominal 0:00/24:00 are 00:10/next 00:10. Blocks group 24 template rows. Exact natural-day bridge totals are in natural_day_blocks.csv.','Q3 V3')
    wb['计划购电量'].cell(1,1).comment=Comment('PROVISIONAL restricted baseline: q0 is Q2 G; midnight official forecast value not evaluated.','Q3 V3')
    wb.save(path)
    check=openpyxl.load_workbook(path,data_only=True)
    error,count=0.,0
    for name,rows in expected.items():
        got=list(check[name].iter_rows(min_row=2,values_only=True)) if rows else []
        if len(got)!=len(rows):
            raise ValueError('workbook row mismatch')
        for actual,want in zip(got,rows):
            for a,b in zip(actual,want):
                count+=1
                if isinstance(b,(int,float)):
                    error=max(error,abs(float(a)-b))
                elif a!=b:
                    raise ValueError(f'workbook cell mismatch {name}: {a!r} != {b!r}')
    if error>1e-6:
        raise ValueError('workbook numeric mismatch')
    return dict(status='PASS',cells_checked=count,max_numeric_diff=error)


def write_variant(directory,ledger,decisions,plans,report_start,report_end):
    directory.mkdir(parents=True,exist_ok=False)
    ledger.to_csv(directory/'interval_ledger.csv',index=False)
    decisions.to_csv(directory/'decisions.csv',index=False)
    plans.to_csv(directory/'daily_plans.csv',index=False)
    report=ledger[(ledger.day_index>=report_start)&(ledger.day_index<=report_end)].copy()
    daily=daily_metrics(report)
    daily.to_csv(directory/'daily_ledger.csv',index=False)
    grouped=daily.assign(month=daily.plan_day.dt.strftime('%Y-%m')).groupby('month')
    monthly=grouped[CASH+ENERGY+['asset_adjusted_cost_yuan']].sum()
    monthly['daily_cost_p95_yuan']=grouped.total_cost_yuan.quantile(.95)
    monthly['daily_cost_cvar95_yuan']=grouped.total_cost_yuan.apply(lambda x: x.nlargest(max(1,int(np.ceil(.05*len(x))))).mean())
    monthly.to_csv(directory/'monthly_metrics.csv')
    monthly.loc[monthly.index.str[-2:].isin(['03','06','09','12'])].to_csv(directory/'sample_months.csv')
    blocks=report.assign(block_index=report.time_index//24).groupby(['plan_day','block_index'])[['charge_kwh','discharge_kwh']].sum().reset_index()
    blocks.to_csv(directory/'template_day_blocks.csv',index=False)
    start=report.plan_day.min()
    end=report.plan_day.max()+pd.Timedelta(days=1)
    natural=ledger[(ledger.interval_start>=start)&(ledger.interval_start<end)]
    ng=natural.assign(block_index=natural.interval_start.dt.hour//4).groupby(['natural_day','block_index'])
    nb=ng.agg(charge_kwh=('charge_kwh','sum'),discharge_kwh=('discharge_kwh','sum'),initial_soc_kwh=('soc_before_kwh','first'),end_soc_kwh=('soc_after_kwh','last'),interval_count=('time_index','size')).reset_index()
    nb['complete_block']=nb.interval_count==24
    nb.to_csv(directory/'natural_day_blocks.csv',index=False)
    events=emergency_events(natural)
    events.to_csv(directory/'emergency_events.csv',index=False)
    focus=report[report.plan_day.dt.strftime('%Y-%m-%d').isin(FOCUS)]
    focus[focus.time_index.isin([59,71,83,95,107,119])].to_csv(directory/'designated_table1_slots.csv',index=False)
    daily[daily.plan_day.dt.strftime('%Y-%m-%d').isin(FOCUS)].to_csv(directory/'designated_table1_daily.csv',index=False)
    blocks[blocks.plan_day.dt.strftime('%Y-%m-%d').isin(FOCUS)].to_csv(directory/'designated_table2.csv',index=False)
    events[events.date.isin(pd.to_datetime(FOCUS))].to_csv(directory/'designated_table3.csv',index=False)
    write_json(directory/'template_readback.json',write_workbook(directory/'result3.xlsx',report,blocks,events))
    costs=np.sort(daily.total_cost_yuan.to_numpy())
    result=dict(days=len(daily),intervals=len(report),**{c:float(report[c].sum()) for c in CASH+ENERGY},
                initial_soc_kwh=float(report.iloc[0].soc_before_kwh),end_soc_kwh=float(report.iloc[-1].soc_after_kwh),
                asset_adjusted_cost_yuan=float(report.total_cost_yuan.sum()+TERMINAL_VALUE*(6000-report.iloc[-1].soc_after_kwh)),
                boundary_adjusted_cost_yuan=float(daily.asset_adjusted_cost_yuan.sum()),
                daily_cost_p95_yuan=float(np.quantile(costs,.95)),daily_cost_cvar95_yuan=float(costs[-max(1,int(np.ceil(.05*len(costs)))):].mean()),
                accepted_releases=int(decisions.accepted.sum()) if len(decisions) else 0,settlement=str(report.iloc[0].settlement),formal_use=False)
    write_json(directory/'summary.json',result)
    return result,daily


def paired_comparison(dailies,path):
    base=dailies['baseline'].set_index('plan_day')
    rows=[]
    for name,d in dailies.items():
        if name=='baseline':
            continue
        pair=d.set_index('plan_day').join(base,lsuffix='_policy',rsuffix='_base',how='inner')
        if len(pair)!=len(base):
            raise ValueError('comparison windows differ')
        for period,part in [('all',pair),('development',pair[pair.index.month<=6]),('audit',pair[pair.index.month>=7])]:
            if part.empty:
                continue
            diff=(part.asset_adjusted_cost_yuan_policy-part.asset_adjusted_cost_yuan_base).to_numpy()
            for block in [3,7,14]:
                rng=np.random.default_rng(20260912)
                b=min(block,len(diff))
                starts=rng.integers(0,len(diff)-b+1,size=(2000,int(np.ceil(len(diff)/b))))
                ix=(starts[:,:,None]+np.arange(b)).reshape(2000,-1)[:,:len(diff)]
                low,high=np.quantile(diff[ix].mean(axis=1),[.025,.975])
                rows.append(dict(variant=name,period=period,days=len(diff),block_days=block,
                    cash_difference_yuan=float((part.total_cost_yuan_policy-part.total_cost_yuan_base).sum()),
                    boundary_adjusted_difference_yuan=float(diff.sum()),mean_difference_yuan=float(diff.mean()),bootstrap_low=float(low),bootstrap_high=float(high)))
    pd.DataFrame(rows).to_csv(path,index=False)


def plot_comparison(dailies,path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    fig,axes=plt.subplots(2,1,figsize=(10,7),sharex=True,layout='constrained')
    for name,d in dailies.items():
        if name not in ('baseline','A','B'):
            continue
        axes[0].plot(d.plan_day,d.total_cost_yuan.cumsum()/1e4,label=name)
        axes[1].plot(d.plan_day,d.emergency_kwh.cumsum(),label=name)
    axes[0].set_ylabel('Cumulative cash / 10,000 yuan')
    axes[1].set_ylabel('Cumulative emergency / kWh')
    axes[0].set_title('Q3 provisional: fixed Q2 midnight contracts')
    for ax in axes:
        ax.grid(alpha=.2)
        ax.legend()
    fig.savefig(path,dpi=180)
    plt.close(fig)
