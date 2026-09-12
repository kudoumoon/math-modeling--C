"""Q4 template delivery, cell readback and ledger-derived tables."""
from copy import copy
import numpy as np
import openpyxl
import pandas as pd
from q4_core import ROOT, validate_ledger


def write_template(frame, path, part, annual=False, full_ledger=None):
    validate_ledger(frame)
    days = sorted(frame.plan_day.unique())
    if annual and days != list(pd.date_range("2025-02-01", "2025-12-31").strftime("%Y-%m-%d")):
        raise ValueError("annual template requires all 334 report days")
    book = openpyxl.load_workbook(ROOT / "data/附件5" / f"result4-{part}.xlsx")
    book.worksheets[-2].cell(1, 6).comment = openpyxl.comments.Comment(
        "PROVISIONAL. Time A nominal 00:00/24:00 represent actual 00:10/next 00:10. Storage blocks group 24 template rows.", "Q4 V1")
    if part == 3:
        book.worksheets[1].cell(1, 147).comment = openpyxl.comments.Comment(
            "Final adjusted ordinary purchase cost: base contract plus up/down fees. Emergency cash is separate; deltas are in interval ledger.", "Q4 V1")
    names = book.sheetnames.copy()
    headers = {s.title: [c.value for c in s[1]] for s in book}
    styles = {s.title: [copy(c._style) for c in s[2]] for s in book}
    expected = {name: [] for name in names}
    for sheet in book:
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        sheet.delete_rows(2, sheet.max_row)
    for day in days:
        data = frame.loc[frame.plan_day == day].sort_values("time_index")
        date = pd.Timestamp(day).to_pydatetime()
        expected[names[0]].append([date, *data.planned_grid_kwh.tolist(), float(data.planned_grid_kwh.sum()), float(data.base_cost_yuan.sum())])
        if part == 3:
            expected[names[1]].append([date, *data.final_grid_kwh.tolist(), float(data.final_grid_kwh.sum()), float(data.ordinary_cost_yuan.sum())])
        for block in range(6):
            chunk = data.iloc[24*block:24*(block+1)]
            expected[names[-2]].append([date if block == 0 else None,
                f"{block*4}:00-{(block+1)*4}:00", float(chunk.charge_kwh.sum()), float(chunk.discharge_kwh.sum()),
                "0:00" if block == 0 else ("24:00" if block == 1 else None),
                float(data.iloc[0].actual_soc_before_kwh) if block == 0 else (float(data.iloc[-1].actual_soc_after_kwh) if block == 1 else None)])
    natural = frame if full_ledger is None else full_ledger
    starts = pd.to_datetime(natural.interval_start)
    first, last = pd.Timestamp(days[0]), pd.Timestamp(days[-1])+pd.Timedelta(days=1)
    natural = natural[(starts >= first) & (starts < last)].sort_values("interval_start")
    if annual and (len(natural) != 334*144 or pd.Timestamp(natural.iloc[0].interval_start) != first):
        raise ValueError("annual natural-clock emergency window needs Jan31 bridge")
    previous_date = None
    for _, row in natural[natural.emergency_kwh > 1e-8].iterrows():
        start, end = pd.Timestamp(row.interval_start), pd.Timestamp(row.interval_end)
        date = start.normalize().to_pydatetime()
        end_label = "24:00" if end.normalize() > start.normalize() else f"{end.hour}:{end.minute:02d}"
        expected[names[-1]].append([date if date != previous_date else None,
            f"{start.hour}:{start.minute:02d}-{end_label}", float(row.emergency_kwh)])
        previous_date = date
    if not expected[names[-1]]:
        expected[names[-1]] = [[pd.Timestamp(days[0]).to_pydatetime(), None, 0.]]
    for name, rows in expected.items():
        for ri, row in enumerate(rows, 2):
            for ci, value in enumerate(row, 1):
                cell = book[name].cell(ri, ci, value)
                cell._style = copy(styles[name][ci-1])
    book.save(path)
    return readback_template(path, expected, headers)


def readback_template(path, expected, headers):
    book = openpyxl.load_workbook(path, data_only=False)
    if book.sheetnames != list(expected):
        raise AssertionError("sheet names mismatch")
    count = 0
    for name, rows in expected.items():
        sheet = book[name]
        if [c.value for c in sheet[1]] != headers[name] or sheet.max_row != len(rows)+1:
            raise AssertionError("template header/shape mismatch")
        for ri, row in enumerate(rows, 2):
            for ci, value in enumerate(row, 1):
                actual = sheet.cell(ri, ci).value
                if isinstance(value, (float, int)):
                    if not isinstance(actual, (float, int)) or not np.isclose(actual, value, atol=1e-6, rtol=0):
                        raise AssertionError(f"numeric mismatch {name}:{ri},{ci}")
                elif actual != value:
                    raise AssertionError(f"label mismatch {name}:{ri},{ci}")
                count += 1
    return {"passed": True, "checked_cells": count, "sheet_names": book.sheetnames,
            "scope": "PROVISIONAL", "soc_labels": "nominal 00:00/24:00 denote actual 00:10/next 00:10"}


def tables_and_figures(frame, result_dir, figure_dir, prefix, full_ledger=None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    daily = frame.groupby("plan_day").agg(cash_cost_yuan=("cash_cost_yuan", "sum"),
        ordinary_cost_yuan=("ordinary_cost_yuan", "sum"), emergency_cost_yuan=("emergency_cost_yuan", "sum"),
        emergency_kwh=("emergency_kwh", "sum"), unused_supply_kwh=("unused_supply_kwh", "sum"),
        end_soc_kwh=("actual_soc_after_kwh", "last"))
    daily.to_csv(result_dir / f"{prefix}_daily.csv")
    months = pd.to_datetime(daily.index).strftime("%Y-%m")
    monthly = daily.groupby(months).sum()
    monthly["end_soc_kwh"] = daily.groupby(months).end_soc_kwh.last()
    monthly["daily_cvar95_yuan"] = daily.groupby(months).cash_cost_yuan.apply(lambda s: s[s >= s.quantile(.95)].mean())
    monthly.to_csv(result_dir / f"{prefix}_monthly.csv")
    dates = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
    selected = frame[frame.plan_day.isin(dates)]
    selected.to_csv(result_dir / f"{prefix}_designated_intervals.csv", index=False)
    daily.loc[daily.index.isin(dates)].to_csv(result_dir / f"{prefix}_designated_daily.csv")
    natural = frame.copy() if full_ledger is None else full_ledger.copy()
    starts = pd.to_datetime(natural.interval_start)
    first, last = pd.Timestamp(frame.plan_day.min()), pd.Timestamp(frame.plan_day.max())+pd.Timedelta(days=1)
    natural = natural[(starts >= first) & (starts < last)].copy()
    natural["natural_day"] = pd.to_datetime(natural.interval_start).dt.strftime("%Y-%m-%d")
    natural["block_index"] = pd.to_datetime(natural.interval_start).dt.hour // 4
    blocks = natural.groupby(["natural_day", "block_index"]).agg(
        grid_kwh=("final_grid_kwh", "sum"), charge_kwh=("charge_kwh", "sum"),
        discharge_kwh=("discharge_kwh", "sum"), emergency_kwh=("emergency_kwh", "sum"),
        cash_cost_yuan=("cash_cost_yuan", "sum"), initial_soc_kwh=("actual_soc_before_kwh", "first"),
        end_soc_kwh=("actual_soc_after_kwh", "last"))
    blocks.to_csv(result_dir / f"{prefix}_natural_four_hour_blocks.csv")
    blocks.loc[blocks.index.get_level_values(0).isin(dates)].to_csv(result_dir / f"{prefix}_designated_natural_blocks.csv")
    natural[natural.emergency_kwh > 1e-8].to_csv(result_dir / f"{prefix}_natural_emergency_events.csv", index=False)
    fig, axes = plt.subplots(2, 1, figsize=(10, 6), sharex=True)
    axes[0].plot(pd.to_datetime(daily.index), daily.cash_cost_yuan.cumsum(), color="#13795b")
    axes[0].set_ylabel("Cumulative cash / yuan")
    axes[1].plot(pd.to_datetime(daily.index), daily.emergency_kwh, color="#b3455a")
    axes[1].set_ylabel("Emergency / kWh")
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(figure_dir / f"{prefix}_cost_emergency.png", dpi=160)
    plt.close(fig)
