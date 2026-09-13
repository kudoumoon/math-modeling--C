from dataclasses import replace
from pathlib import Path
import sys

import numpy as np
import openpyxl
import pandas as pd
import pytest

from q4_prices import forecast_prices
from q4_core import ROOT, MidnightPolicy, model, online, interval_ledger, validate_ledger
from q4_delivery import write_template
from run_q4 import q42


def fixture_data():
    dates = pd.date_range("2025-01-01", periods=12)
    load = np.tile(3000 + 500*np.sin(np.arange(144)/13), (12, 1))
    pv = np.tile(np.maximum(0, 2200*np.sin((np.arange(144)-36)/72*np.pi)), (12, 1))
    prior = np.r_[np.full(72, .3), np.full(72, .8)]
    data = model.Q2Data(dates, prior, load, pv, [])
    prices = np.tile(prior, (12, 1))
    return data, prices


@pytest.mark.parametrize("day", [7, 8])
def test_price_future_and_uncompleted_invariance(day):
    data, prices = fixture_data()
    for hour in (0, 6, 12, 18):
        issue = data.dates[day]+pd.Timedelta(hours=hour)
        expected = forecast_prices(data.dates, prices, data.price, day, issue)
        corrupted = prices.copy()
        completion = data.dates.to_numpy()[:, None] + pd.to_timedelta((np.arange(144)+2)*10, unit="min").to_numpy()[None, :]
        corrupted[completion > issue.to_datetime64()] = np.nan
        actual = forecast_prices(data.dates, corrupted, data.price, day, issue)
        np.testing.assert_array_equal(actual.values, expected.values)
        assert (actual.source_completed_at <= issue).all()
    cold = forecast_prices(data.dates, prices*np.nan, data.price, 0, data.dates[0])
    np.testing.assert_array_equal(cold.values, data.price)


def test_real_risk_core_matches_q2_and_blocks_future():
    data, prices = fixture_data()
    day = 2
    forecast = forecast_prices(data.dates, prices, data.price, day, data.dates[day])
    p = MidnightPolicy(data, data.load.copy(), data.pv.copy())
    chosen, grids, diagnostics = p.plan(day, 6000., forecast)
    assert len(grids) == 15
    assert chosen == online.DEFAULT_CANDIDATE
    online._worker_init(data, data.load, data.pv, "A", 0)
    key, q2 = online._solve_candidate((day, 6000., 6000., online.candidate_grid()[1], 10))
    np.testing.assert_allclose(grids[key], q2["grid"], atol=1e-7)
    changed = replace(data, load=data.load.copy(), pv=data.pv.copy())
    changed.load[1:] = np.nan
    changed.pv[1:] = np.nan
    p2 = MidnightPolicy(changed, data.load.copy(), data.pv.copy())
    _, mutated, _ = p2.plan(day, 6000., forecast)
    for key in grids:
        np.testing.assert_array_equal(grids[key], mutated[key])
    p.settle_candidate_scores(day, grids, 6000., prices[day])
    assert not any(p.histories.values())
    p.pending = online.release_mature_scores(p.pending, p.histories, day+1)
    assert not any(p.histories.values())
    p.pending = online.release_mature_scores(p.pending, p.histories, day+2)
    assert all(len(v) == 1 for v in p.histories.values())


def test_actual_price_settlement_physics_and_tamper(tmp_path):
    data, prices = fixture_data()
    f = forecast_prices(data.dates, prices, data.price, 0, data.dates[0])
    q0 = np.full(144, 200.)
    q = q0.copy()
    q[35:71] += 100
    q[71:107] -= 50
    actual = np.linspace(.2, 1.2, 144)
    settled = model.settle_causally(q, data.load[0], data.pv[0], 6000.)
    a = interval_ledger(data, 0, q0, q, settled, f, 6000., 6000., actual, "test", "A")
    b = interval_ledger(data, 0, q0, q, settled, f, 6000., 6000., actual, "test", "B")
    assert np.isclose(a.cash_cost_yuan.sum()-b.cash_cost_yuan.sum(), actual @ np.maximum(q0-q, 0))
    assert not np.isclose(a.base_cost_yuan.sum(), f.values @ q0)
    for part in (2, 3):
        path = tmp_path / f"part{part}.xlsx"
        assert write_template(a if part == 3 else interval_ledger(data, 0, q0, q0,
            model.settle_causally(q0, data.load[0], data.pv[0], 6000.), f, 6000., 6000., actual, "test"), path, part)["passed"]
        book = openpyxl.load_workbook(path)
        assert book.worksheets[0].cell(2, 2).value == 200.
    corrupt = a.copy()
    corrupt.loc[5, "cash_cost_yuan"] += 1
    with pytest.raises(AssertionError, match="mismatch"):
        validate_ledger(corrupt)
    corrupt = a.copy()
    corrupt.loc[0, "actual_soc_after_kwh"] += 2
    with pytest.raises(AssertionError):
        validate_ledger(corrupt)


def test_q42_next_midnight_does_not_use_uncompleted_truth():
    data, prices = fixture_data()
    load_pred, pv_pred = data.load.copy(), data.pv.copy()
    baseline, _ = q42(data, load_pred, pv_pred, prices, 1)
    changed = replace(data, pv=data.pv.copy())
    changed.pv[0, -1] += 50000
    perturbed, _ = q42(changed, load_pred, pv_pred, prices, 1)
    a, b = baseline[baseline.plan_day == "2025-01-02"], perturbed[perturbed.plan_day == "2025-01-02"]
    np.testing.assert_array_equal(a.planned_grid_kwh, b.planned_grid_kwh)
    np.testing.assert_array_equal(a.planned_initial_energy_kwh, b.planned_initial_energy_kwh)
    assert a.executed_initial_energy_kwh.iloc[0] != b.executed_initial_energy_kwh.iloc[0]


def test_q3_release_future_invariance_and_adjustment_prices():
    sys.path.insert(0, str(ROOT / "solutions/Q3_V3_ALIGNED/program"))
    import q3_api
    from types import SimpleNamespace
    data, prices = fixture_data()
    releases = np.full((12, 4, 24), 1000.)
    inputs = q3_api.Inputs(data.dates, data.load, data.pv, data.price, releases)
    upstream = SimpleNamespace(load_forecast_kw=data.load.copy(), pv_forecast_kw=data.pv.copy())
    f = q3_api.build_release_forecast(inputs, upstream, 8, 6)
    dirty = replace(inputs, load_kw=inputs.load_kw.copy(), pv_kw=inputs.pv_kw.copy(), releases_kw=releases.copy())
    dirty.load_kw[8, 35:] = np.nan
    dirty.pv_kw[8, 35:] = np.nan
    dirty.releases_kw[8, 2:] = np.nan
    after = q3_api.build_release_forecast(dirty, upstream, 8, 6)
    np.testing.assert_array_equal(after.net_kwh, f.net_kwh)
    q0 = np.full(12, 800.)
    result = q3_api.solve_adjustment(q0, np.zeros(12), np.ones(12), 6000., np.ones(12, bool), "B")
    assert result["r"].sum() > 0
    assert result["max_residual"] < 1e-6


def test_mature_real_residual_pools_and_accepted_heldout_gate():
    sys.path.insert(0, str(ROOT / "solutions/Q3_V3_ALIGNED/program"))
    import q3_api
    from q4_provenance import upstream_path
    inputs = q3_api.load_inputs()
    upstream = q3_api.load_upstream(run_path=upstream_path())
    risk, gate, risk_ids, gate_ids = q3_api.residual_pools(inputs, upstream, 78, 6)
    assert len(risk) >= 7 and len(gate) >= 7
    assert not set(risk_ids) & set(gate_ids)
    assert max(risk_ids+gate_ids) < 78
    forecast = q3_api.build_release_forecast(inputs, upstream, 78, 6)
    # A deliberately undercontracted real March-20 release tests the accepted
    # branch with disjoint real historical residuals, without choosing a policy.
    q0 = np.zeros(len(forecast.steps))
    p = np.full(len(q0), .5)
    candidate = q3_api.solve_adjustment(q0, forecast.net_kwh, p, 1200., np.ones(len(q0), bool), "A")
    keep = q3_api.gate_score(q0, q0, forecast.net_kwh, p, 1200., gate, "A")
    new = q3_api.gate_score(q0, candidate["q"], forecast.net_kwh, p, 1200., gate, "A")
    assert new < keep-1e-6


def test_natural_midnight_emergency_and_final_adjusted_template(tmp_path):
    data, prices = fixture_data()
    frames = []
    soc = 1200.
    for day in range(2):
        f = forecast_prices(data.dates, prices, data.price, day, data.dates[day])
        q0, q = np.zeros(144), np.zeros(144)
        q[35:71] = 100.
        settled = model.settle_causally(q, data.load[day], data.pv[day], soc)
        frames.append(interval_ledger(data, day, q0, q, settled, f, soc, soc, prices[day], "test"))
        soc = settled["soc"][-1]
    full = pd.concat(frames, ignore_index=True)
    path = tmp_path / "midnight.xlsx"
    write_template(frames[1], path, 3, full_ledger=full)
    book = openpyxl.load_workbook(path)
    emergency = book.worksheets[-1]
    assert emergency.cell(2, 2).value == "0:00-0:10"
    assert emergency.cell(2, 3).value == pytest.approx(frames[0].iloc[-1].emergency_kwh)
    assert book.worksheets[1].cell(2, 37).value == 100.
    assert book.worksheets[1].cell(2, 147).value == pytest.approx(frames[1].ordinary_cost_yuan.sum())
