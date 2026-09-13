"""Freeze verified Q2 V3 runs into a versioned, reviewable release."""

from __future__ import annotations

import argparse
import copy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

from q2_forecasts import shift_cross_day
from q2_model import (
    DT,
    E_MAX,
    E_MIN,
    ETA_C,
    ETA_D,
    K_EMERGENCY,
    Q_MAX,
    load_data,
)


REPO_ROOT = Path(__file__).resolve().parents[3]
VERSION_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = REPO_ROOT / "data/附件5/result2.xlsx"
DEFAULT_RUNS = {
    "main": REPO_ROOT / "runs/q2v3-r2-full-b0-a-20260913-01",
    "battery_b": REPO_ROOT / "runs/q2v3-r2-full-b0-batb-20260913-01",
    "time_b": REPO_ROOT / "runs/q2v3-r2-full-b0-timeb-20260913-01",
    "delay_1": REPO_ROOT / "runs/q2v3-r2-full-b0-delay1-20260913-01",
    "reset_feb": REPO_ROOT / "runs/q2v3-r2-full-b0-resetfeb-20260913-01",
}
DEFAULT_RUNS = {
    name: VERSION_ROOT / "results/annual_runs" / path.name
    if (VERSION_ROOT / "results/annual_runs" / path.name).is_dir() else path
    for name, path in DEFAULT_RUNS.items()
}
EXPECTED = {
    "main": {"time_mapping": "A", "battery_interpretation": "A", "end_day": 364, "observation_delay_slots": 0, "compatibility_reset": False},
    "battery_b": {"time_mapping": "A", "battery_interpretation": "B", "end_day": 364, "observation_delay_slots": 0, "compatibility_reset": False},
    "time_b": {"time_mapping": "B", "battery_interpretation": "A", "end_day": 363, "observation_delay_slots": 0, "compatibility_reset": False},
    "delay_1": {"time_mapping": "A", "battery_interpretation": "A", "end_day": 364, "observation_delay_slots": 1, "compatibility_reset": False},
    "reset_feb": {"time_mapping": "A", "battery_interpretation": "A", "end_day": 364, "observation_delay_slots": 0, "compatibility_reset": True},
}
TOL = 1e-7


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )


def _git_blob_sha256(commit: str, relative_path: str) -> str:
    content = subprocess.check_output(
        ["git", "show", f"{commit}:{relative_path}"], cwd=REPO_ROOT
    )
    return hashlib.sha256(content).hexdigest()


def audit_soc_boundaries(
    soc: np.ndarray, run_days: np.ndarray, *, compatibility_reset: bool,
    report_start_day: int = 31, initial_soc_kwh: float = 6000.0,
) -> dict[str, Any]:
    """Validate every transition, isolating only the declared RESET1 boundary."""
    destinations = run_days[1:]
    jumps = soc[destinations, 0] - soc[run_days[:-1], -1]
    exempt = destinations == report_start_day if compatibility_reset else np.zeros(len(jumps), dtype=bool)
    ordinary = jumps[~exempt]
    reset_present = not compatibility_reset or int(exempt.sum()) == 1
    reset_target_matches = not compatibility_reset or (
        reset_present and abs(float(soc[report_start_day, 0]) - initial_soc_kwh) <= TOL
    )
    residual = float(np.max(np.abs(ordinary), initial=0.0))
    return {
        "passed": bool(np.all(np.isfinite(jumps)) and residual <= 1e-8 and reset_target_matches),
        "continuous_boundaries_checked": int((~exempt).sum()),
        "max_continuous_boundary_residual_kwh": residual,
        "declared_reset_day_index": report_start_day if compatibility_reset else None,
        "declared_reset_jump_kwh": float(jumps[exempt][0]) if compatibility_reset and reset_present else None,
        "declared_reset_target_kwh": initial_soc_kwh if compatibility_reset else None,
        "declared_reset_target_matches": bool(reset_target_matches),
        "reset_is_sensitivity_only": compatibility_reset,
    }


def verify_run(name: str, run_root: Path) -> dict[str, Any]:
    manifest = json.loads((run_root / "manifest.json").read_text(encoding="utf-8"))
    config = json.loads((run_root / "config.json").read_text(encoding="utf-8"))
    summary = json.loads((run_root / "summary.json").read_text(encoding="utf-8"))
    expected = EXPECTED[name]
    checks: dict[str, bool] = {
        "manifest_complete": manifest.get("status") == "complete",
        "arm_is_frozen_b0": config.get("forecast_arm") == "B0",
        "forecast_model_is_frozen": config.get("forecast_model_id") == "Q2V2-SW2-REC5",
        "starts_january_1": config.get("start_day") == 0,
        "reports_from_february_1": config.get("report_start_day") == 31,
        "time_mapping_matches": config.get("time_mapping") == expected["time_mapping"],
        "battery_interpretation_matches": config.get("battery_interpretation") == expected["battery_interpretation"],
        "end_day_matches": config.get("end_day") == expected["end_day"],
        "observation_delay_matches": config.get("observation_delay_slots") == expected["observation_delay_slots"],
        "compatibility_reset_matches": config.get("compatibility_reset_on_report_start") is expected["compatibility_reset"],
        "initial_soc_is_6000": abs(float(config.get("initial_soc_kwh", np.nan)) - 6000.0) <= TOL,
        "candidate_score_delay_is_two_days": config.get("candidate_full_score_delay_days") == 2,
        "scenario_cutoff_is_day_minus_two": config.get("scenario_latest_complete_plan_day") == "day_index - 2",
    }

    output_hashes = manifest.get("outputs", {})
    checks["all_manifest_outputs_present"] = bool(output_hashes) and all(
        (run_root / relative).is_file() for relative in output_hashes
    )
    checks["all_manifest_output_hashes_match"] = bool(output_hashes) and all(
        sha256(run_root / relative) == digest
        for relative, digest in output_hashes.items()
    )
    source_hashes = manifest.get("source_hashes", {})
    commit = manifest.get("git_commit", "")
    checks["all_live_source_hashes_match"] = bool(source_hashes) and all(
        (REPO_ROOT / relative).is_file()
        and sha256(REPO_ROOT / relative) == digest
        for relative, digest in source_hashes.items()
    )
    checks["all_git_blob_hashes_match"] = bool(source_hashes) and all(
        _git_blob_sha256(commit, relative) == digest
        for relative, digest in source_hashes.items()
    )
    input_hashes = manifest.get("input_hashes", {})
    checks["all_live_input_hashes_match"] = bool(input_hashes) and all(
        (REPO_ROOT / relative).is_file()
        and sha256(REPO_ROOT / relative) == digest
        for relative, digest in input_hashes.items()
    )

    with np.load(run_root / "policy_arrays.npz") as stored:
        arrays = {key: stored[key] for key in stored.files}
    shapes = {key: value.shape for key, value in arrays.items()}
    checks["policy_array_shapes"] = (
        shapes == {
            "G": (365, 144),
            "C": (365, 144),
            "D": (365, 144),
            "CommandedC": (365, 144),
            "CommandedD": (365, 144),
            "Emergency": (365, 144),
            "Spill": (365, 144),
            "SOC": (365, 145),
        }
    )

    data = load_data()
    if expected["time_mapping"] == "B":
        price = shift_cross_day(np.tile(data.price, (365, 1)))[0]
        load = shift_cross_day(data.load)
        pv = shift_cross_day(data.pv)
    else:
        price, load, pv = data.price, data.load, data.pv
    run_days = np.arange(0, expected["end_day"] + 1)
    report_days = np.arange(31, expected["end_day"] + 1)
    g = arrays["G"][report_days]
    c = arrays["C"][report_days]
    discharge = arrays["D"][report_days]
    emergency = arrays["Emergency"][report_days]
    spill = arrays["Spill"][report_days]
    soc = arrays["SOC"][report_days]
    balance = g + pv[report_days] * DT + discharge + emergency - load[report_days] * DT - c - spill
    state = soc[:, 1:] - soc[:, :-1] - ETA_C * c + discharge / ETA_D
    continuity = arrays["SOC"][run_days[:-1], -1] - arrays["SOC"][run_days[1:], 0]
    boundary_audit = audit_soc_boundaries(
        arrays["SOC"], run_days,
        compatibility_reset=expected["compatibility_reset"],
    )
    calculated = {
        "days": int(len(report_days)),
        "plan_cost_yuan": float(np.sum(g * price)),
        "emergency_cost_yuan": float(np.sum(emergency * price * K_EMERGENCY)),
        "grid_purchase_kwh": float(g.sum()),
        "emergency_kwh": float(emergency.sum()),
        "emergency_intervals": int(np.sum(emergency > TOL)),
        "emergency_days": int(np.sum(emergency.sum(axis=1) > TOL)),
        "spill_kwh": float(spill.sum()),
        "charge_kwh": float(c.sum()),
        "discharge_kwh": float(discharge.sum()),
        "soc_min_kwh": float(np.nanmin(soc)),
        "soc_max_kwh": float(np.nanmax(soc)),
        "end_soc_kwh": float(soc[-1, -1]),
    }
    calculated["total_cost_yuan"] = calculated["plan_cost_yuan"] + calculated["emergency_cost_yuan"]
    numeric_differences = {
        key: abs(float(summary[key]) - value) for key, value in calculated.items()
    }
    checks.update({
        "summary_recomputed": max(numeric_differences.values(), default=0.0) <= 1e-5,
        "energy_balance_closes": float(np.max(np.abs(balance))) <= 1e-8,
        "soc_balance_closes": float(np.max(np.abs(state))) <= 1e-8,
        "cross_day_soc_contract": boundary_audit["passed"],
        "soc_bounds_hold": float(np.nanmin(arrays["SOC"][run_days])) >= E_MIN - TOL
        and float(np.nanmax(arrays["SOC"][run_days])) <= E_MAX + TOL,
        "power_bounds_hold": min(float(c.min()), float(discharge.min()), float(g.min()), float(emergency.min()), float(spill.min())) >= -TOL
        and float(max(c.max(), discharge.max())) <= Q_MAX + TOL,
        "no_simultaneous_charge_discharge": not bool(np.any((c > TOL) & (discharge > TOL))),
    })

    ledger = pd.read_csv(run_root / "interval_ledger.csv")
    expected_rows = len(report_days) * 144
    flat = {
        "planned_grid_kwh": g.ravel(),
        "charge_kwh": c.ravel(),
        "discharge_kwh": discharge.ravel(),
        "commanded_charge_kwh": arrays["CommandedC"][report_days].ravel(),
        "commanded_discharge_kwh": arrays["CommandedD"][report_days].ravel(),
        "emergency_kwh": emergency.ravel(),
        "unused_supply_kwh": spill.ravel(),
        "actual_soc_before_kwh": soc[:, :-1].ravel(),
        "actual_soc_after_kwh": soc[:, 1:].ravel(),
    }
    ledger_differences = {
        key: float(np.max(np.abs(ledger[key].to_numpy(float) - value)))
        for key, value in flat.items()
    } if len(ledger) == expected_rows else {key: float("inf") for key in flat}
    checks["interval_ledger_rows"] = len(ledger) == expected_rows
    checks["interval_ledger_matches_policy"] = max(ledger_differences.values()) <= 1e-8

    daily = pd.read_csv(run_root / "daily_ledger.csv")
    checks["daily_ledger_days"] = daily["day_index"].tolist() == run_days.tolist()
    if expected["compatibility_reset"]:
        reset_row = daily.loc[daily["day_index"] == 31]
        checks["declared_reset_daily_ledger_matches"] = len(reset_row) == 1 and bool(
            np.all(np.abs(reset_row[["planned_initial_energy_kwh", "executed_initial_energy_kwh"]]
                          .to_numpy(float) - 6000.0) <= TOL)
        )
    checks["daily_score_release_contract"] = bool(
        np.array_equal(
            daily["latest_full_score_day_released"].to_numpy(int),
            daily["day_index"].to_numpy(int) - 2,
        )
    )
    audit = {
        "run_name": name,
        "run_id": manifest.get("run_id"),
        "run_root": str(run_root.relative_to(REPO_ROOT)),
        "git_commit": commit,
        "checks": checks,
        "passed": all(checks.values()),
        "calculated_summary": calculated,
        "summary_max_abs_difference": max(numeric_differences.values(), default=0.0),
        "interval_ledger_max_abs_differences": ledger_differences,
        "max_energy_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_soc_balance_residual_kwh": float(np.max(np.abs(state))),
        "max_cross_day_soc_residual_kwh": float(np.max(np.abs(continuity))),
        "soc_boundary_audit": boundary_audit,
    }
    if not audit["passed"]:
        failed = [key for key, passed in checks.items() if not passed]
        raise AssertionError(f"{name} run verification failed: {failed}")
    return {"manifest": manifest, "config": config, "summary": summary, "arrays": arrays, "daily": daily, "audit": audit}


def emergency_events(ledger: pd.DataFrame) -> pd.DataFrame:
    active = ledger.loc[ledger["emergency_kwh"] > TOL].copy()
    columns = ["date", "start", "end", "time_period", "emergency_kwh", "duration_minutes", "interval_count"]
    if active.empty:
        return pd.DataFrame(columns=columns)
    active["interval_start"] = pd.to_datetime(active["interval_start"])
    active["interval_end"] = pd.to_datetime(active["interval_end"])
    active = active.sort_values("interval_start").reset_index(drop=True)
    groups = (active["interval_start"] != active["interval_end"].shift()).cumsum()
    rows = []
    for _, event in active.groupby(groups, sort=False):
        start = pd.Timestamp(event.iloc[0]["interval_start"])
        end = pd.Timestamp(event.iloc[-1]["interval_end"])
        day_offset = (end.normalize() - start.normalize()).days
        suffix = f"+{day_offset}d" if day_offset else ""
        rows.append({
            "date": start.normalize(),
            "start": start,
            "end": end,
            "time_period": f"{start:%H:%M}-{end:%H:%M}{suffix}",
            "emergency_kwh": float(event["emergency_kwh"].sum()),
            "duration_minutes": float((end - start).total_seconds() / 60.0),
            "interval_count": int(len(event)),
        })
    return pd.DataFrame(rows, columns=columns)


def natural_day_ledger(arrays: dict[str, np.ndarray], dates: pd.DatetimeIndex) -> pd.DataFrame:
    """Map plan-day slots to the official natural-day reporting window."""
    rows = []
    for day in range(31, 365):
        source_days = np.concatenate(([day - 1], np.full(143, day, dtype=int)))
        source_slots = np.concatenate(([143], np.arange(143, dtype=int)))
        natural_date = pd.Timestamp(dates[day])
        for natural_slot, (source_day, source_slot) in enumerate(zip(source_days, source_slots)):
            start = natural_date + pd.Timedelta(minutes=10 * natural_slot)
            rows.append({
                "natural_day": natural_date.date().isoformat(),
                "natural_time_index": natural_slot,
                "source_plan_day": pd.Timestamp(dates[source_day]).date().isoformat(),
                "source_time_index": int(source_slot),
                "interval_start": start,
                "interval_end": start + pd.Timedelta(minutes=10),
                "planned_grid_kwh": float(arrays["G"][source_day, source_slot]),
                "charge_kwh": float(arrays["C"][source_day, source_slot]),
                "discharge_kwh": float(arrays["D"][source_day, source_slot]),
                "commanded_charge_kwh": float(arrays["CommandedC"][source_day, source_slot]),
                "commanded_discharge_kwh": float(arrays["CommandedD"][source_day, source_slot]),
                "emergency_kwh": float(arrays["Emergency"][source_day, source_slot]),
                "unused_supply_kwh": float(arrays["Spill"][source_day, source_slot]),
                "actual_soc_before_kwh": float(arrays["SOC"][source_day, source_slot]),
                "actual_soc_after_kwh": float(arrays["SOC"][source_day, source_slot + 1]),
            })
    result = pd.DataFrame(rows)
    if len(result) != 334 * 144:
        raise AssertionError("natural-day bridge did not produce 334 complete days")
    starts = pd.DatetimeIndex(result["interval_start"])
    ends = pd.DatetimeIndex(result["interval_end"])
    if not bool(np.all(starts[1:].asi8 == ends[:-1].asi8)):
        raise AssertionError("natural-day bridge has a timestamp gap")
    if not np.allclose(result["actual_soc_after_kwh"].to_numpy()[:-1],
                       result["actual_soc_before_kwh"].to_numpy()[1:], atol=1e-8, rtol=0):
        raise AssertionError("natural-day bridge has a SOC discontinuity")
    return result


def storage_blocks(ledger: pd.DataFrame) -> pd.DataFrame:
    work = ledger.copy()
    time_column = "natural_time_index" if "natural_time_index" in work else "time_index"
    day_column = "natural_day" if "natural_day" in work else "plan_day"
    work["block_index"] = work[time_column].astype(int) // 24
    work["date"] = pd.to_datetime(work[day_column])
    work = work.sort_values(["date", time_column])
    if work.duplicated(["date", time_column]).any():
        raise AssertionError("storage ledger has duplicate intervals")
    if not all(np.array_equal(group[time_column].to_numpy(), np.arange(144))
               for _, group in work.groupby("date")):
        raise AssertionError("storage ledger must contain 144 ordered intervals per day")
    grouped = work.groupby(["date", "block_index"], as_index=False).agg(
        charge_kwh=("charge_kwh", "sum"),
        discharge_kwh=("discharge_kwh", "sum"),
    )
    grouped["time_period"] = grouped["block_index"].map(
        {i: f"{4 * i}:00-{4 * (i + 1)}:00" for i in range(6)}
    )
    starts = work.groupby("date")["actual_soc_before_kwh"].first()
    ends = work.groupby("date")["actual_soc_after_kwh"].last()
    grouped["soc_0000_kwh"] = grouped["date"].map(starts)
    grouped["soc_2400_kwh"] = grouped["date"].map(ends)
    return grouped[["date", "block_index", "time_period", "charge_kwh", "discharge_kwh", "soc_0000_kwh", "soc_2400_kwh"]]


def _copy_style(source: openpyxl.cell.Cell, target: openpyxl.cell.Cell) -> None:
    target._style = copy.copy(source._style)
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)
    target.number_format = source.number_format


def write_result2(
    output_path: Path,
    ledger: pd.DataFrame,
    blocks: pd.DataFrame,
    events: pd.DataFrame,
) -> dict[str, Any]:
    shutil.copy2(TEMPLATE, output_path)
    book = openpyxl.load_workbook(output_path)
    book.properties.created = datetime(2026, 9, 13)
    book.properties.modified = datetime(2026, 9, 13)
    plan_sheet = book["计划购电量"]
    storage_sheet = book["充放电量"]
    emergency_sheet = book["紧急购电量"]
    plan_sheet.cell(1, 2).comment = openpyxl.comments.Comment(
        "Archived Time A purchase contract: 00:10 to next-day 00:10. "
        "Storage and emergency sheets use actual natural-clock windows reconstructed "
        "from these same stored trajectories. This export does not implement the new interval-end model.",
        "Q2 delivery",
    )
    storage_sheet.cell(1, 2).comment = openpyxl.comments.Comment(
        "Actual natural-clock four-hour blocks; 00:00/24:00 SOC uses the prior-day "
        "slot143 bridge under archived Time A. See q2_natural_day_ledger.csv.", "Q2 delivery",
    )
    days = pd.DatetimeIndex(pd.to_datetime(ledger["plan_day"].drop_duplicates()))
    if len(days) != 334:
        raise AssertionError(f"official result requires 334 plan days, got {len(days)}")
    plans = ledger.pivot(index="plan_day", columns="time_index", values="planned_grid_kwh").loc[[d.date().isoformat() for d in days]].to_numpy(float)
    prices = ledger.pivot(index="plan_day", columns="time_index", values="price_yuan_per_kwh").loc[[d.date().isoformat() for d in days]].to_numpy(float)
    for offset, day in enumerate(days, start=2):
        plan_sheet.cell(offset, 1, day.to_pydatetime())
        for slot, value in enumerate(plans[offset - 2], start=2):
            plan_sheet.cell(offset, slot, float(value))
        plan_sheet.cell(offset, 146, float(plans[offset - 2].sum()))
        plan_sheet.cell(offset, 147, float(np.dot(plans[offset - 2], prices[offset - 2])))

    storage_styles = [[copy.copy(storage_sheet.cell(row, col)) for col in range(1, 7)] for row in range(2, 8)]
    storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
    for day_offset, day in enumerate(days):
        day_blocks = blocks.loc[blocks["date"] == day].sort_values("block_index")
        if len(day_blocks) != 6:
            raise AssertionError(f"storage blocks incomplete for {day.date()}")
        for block_offset, (_, item) in enumerate(day_blocks.iterrows()):
            row = 2 + day_offset * 6 + block_offset
            for col in range(1, 7):
                _copy_style(storage_styles[block_offset][col - 1], storage_sheet.cell(row, col))
            storage_sheet.cell(row, 1, day.to_pydatetime() if block_offset == 0 else None)
            storage_sheet.cell(row, 2, item["time_period"])
            storage_sheet.cell(row, 3, float(item["charge_kwh"]))
            storage_sheet.cell(row, 4, float(item["discharge_kwh"]))
            storage_sheet.cell(row, 5, "0:00" if block_offset == 0 else ("24:00" if block_offset == 1 else None))
            storage_sheet.cell(row, 6, float(item["soc_0000_kwh"]) if block_offset == 0 else (float(item["soc_2400_kwh"]) if block_offset == 1 else None))

    emergency_styles = [[copy.copy(emergency_sheet.cell(row, col)) for col in range(1, 4)] for row in range(2, 5)]
    emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
    event_rows = events.copy()
    if event_rows.empty:
        event_rows = pd.DataFrame([{"date": days[0], "time_period": None, "emergency_kwh": 0.0}])
    previous_date = None
    for offset, (_, item) in enumerate(event_rows.iterrows(), start=2):
        for col in range(1, 4):
            _copy_style(emergency_styles[(offset - 2) % 3][col - 1], emergency_sheet.cell(offset, col))
        date = pd.Timestamp(item["date"])
        emergency_sheet.cell(offset, 1, date.to_pydatetime() if date != previous_date else None)
        emergency_sheet.cell(offset, 2, item["time_period"])
        emergency_sheet.cell(offset, 3, float(item["emergency_kwh"]))
        previous_date = date
    book.save(output_path)
    return readback_result2(output_path, days, plans, prices, blocks, events)


def readback_result2(
    path: Path,
    days: pd.DatetimeIndex,
    plans: np.ndarray,
    prices: np.ndarray,
    blocks: pd.DataFrame,
    events: pd.DataFrame,
) -> dict[str, Any]:
    book = openpyxl.load_workbook(path, data_only=False)
    plan = book["计划购电量"]
    storage = book["充放电量"]
    emergency = book["紧急购电量"]
    read_plans = np.asarray([[plan.cell(row, col).value for col in range(2, 146)] for row in range(2, 336)], dtype=float)
    read_dates = pd.DatetimeIndex(pd.to_datetime([plan.cell(row, 1).value for row in range(2, 336)]))
    read_totals = np.asarray([plan.cell(row, 146).value for row in range(2, 336)], dtype=float)
    read_costs = np.asarray([plan.cell(row, 147).value for row in range(2, 336)], dtype=float)
    expected_costs = np.sum(plans * prices, axis=1)
    read_charge = np.asarray([storage.cell(row, 3).value for row in range(2, 2 + 334 * 6)], dtype=float)
    read_discharge = np.asarray([storage.cell(row, 4).value for row in range(2, 2 + 334 * 6)], dtype=float)
    read_soc_start = np.asarray([storage.cell(2 + day * 6, 6).value for day in range(334)], dtype=float)
    read_soc_end = np.asarray([storage.cell(3 + day * 6, 6).value for day in range(334)], dtype=float)
    expected_starts = blocks.groupby("date")["soc_0000_kwh"].first().to_numpy(float)
    expected_ends = blocks.groupby("date")["soc_2400_kwh"].first().to_numpy(float)
    expected_event_rows = max(1, len(events))
    carried = None
    event_dates, event_periods, event_values = [], [], []
    for row in range(2, emergency.max_row + 1):
        if emergency.cell(row, 1).value is not None:
            carried = pd.Timestamp(emergency.cell(row, 1).value)
        event_dates.append(carried)
        event_periods.append(emergency.cell(row, 2).value)
        event_values.append(float(emergency.cell(row, 3).value))
    event_match = len(event_values) == expected_event_rows
    if len(events):
        event_match = event_match and pd.DatetimeIndex(event_dates).equals(pd.DatetimeIndex(events["date"]))
        event_match = event_match and event_periods == events["time_period"].tolist()
        event_match = event_match and float(np.max(np.abs(np.asarray(event_values) - events["emergency_kwh"].to_numpy(float)))) <= TOL
    checks = {
        "sheet_names_preserved": book.sheetnames == ["计划购电量", "充放电量", "紧急购电量"],
        "plan_shape": (plan.max_row, plan.max_column) == (335, 147),
        "plan_dates_match": read_dates.equals(days),
        "plan_values_match": float(np.max(np.abs(read_plans - plans))) <= TOL,
        "plan_totals_match": float(np.max(np.abs(read_totals - plans.sum(axis=1)))) <= TOL,
        "plan_costs_match": float(np.max(np.abs(read_costs - expected_costs))) <= TOL,
        "storage_rows": storage.max_row - 1 == 334 * 6,
        "storage_charge_match": float(np.max(np.abs(read_charge - blocks["charge_kwh"].to_numpy(float)))) <= TOL,
        "storage_discharge_match": float(np.max(np.abs(read_discharge - blocks["discharge_kwh"].to_numpy(float)))) <= TOL,
        "storage_soc_start_match": float(np.max(np.abs(read_soc_start - expected_starts))) <= TOL,
        "storage_soc_end_match": float(np.max(np.abs(read_soc_end - expected_ends))) <= TOL,
        "emergency_events_match": bool(event_match),
    }
    result = {
        "checks": checks,
        "passed": all(checks.values()),
        "plan_rows": plan.max_row - 1,
        "storage_rows": storage.max_row - 1,
        "emergency_rows": emergency.max_row - 1,
        "max_plan_value_difference_kwh": float(np.max(np.abs(read_plans - plans))),
    }
    if not result["passed"]:
        raise AssertionError(f"result2 readback failed: {[key for key, value in checks.items() if not value]}")
    return result


def monthly_summary(ledger: pd.DataFrame) -> pd.DataFrame:
    work = ledger.copy()
    work["month"] = pd.to_datetime(work["plan_day"]).dt.to_period("M").astype(str)
    work["total_cost_yuan"] = work["plan_cost_yuan"] + work["emergency_cost_yuan"]
    return work.groupby("month", as_index=False).agg(
        plan_cost_yuan=("plan_cost_yuan", "sum"),
        emergency_cost_yuan=("emergency_cost_yuan", "sum"),
        total_cost_yuan=("total_cost_yuan", "sum"),
        grid_purchase_kwh=("planned_grid_kwh", "sum"),
        emergency_kwh=("emergency_kwh", "sum"),
        charge_kwh=("charge_kwh", "sum"),
        discharge_kwh=("discharge_kwh", "sum"),
    )


def make_figures(output: Path, ledger: pd.DataFrame, monthly: pd.DataFrame, sensitivity: pd.DataFrame) -> list[Path]:
    output.mkdir(parents=True, exist_ok=False)
    plt.rcParams.update({"font.size": 9, "axes.grid": True, "grid.alpha": 0.25})
    paths = []

    fig, ax = plt.subplots(figsize=(8.2, 4.4), constrained_layout=True)
    x = np.arange(len(monthly))
    ax.bar(x, monthly["plan_cost_yuan"] / 1e4, color="#287271", label="Scheduled purchase")
    ax.bar(x, monthly["emergency_cost_yuan"] / 1e4, bottom=monthly["plan_cost_yuan"] / 1e4, color="#D1495B", label="Emergency purchase")
    ax.set(xlabel="Month", ylabel="Cost (10,000 yuan)", title="Q2 V3 monthly operating cost")
    ax.set_xticks(x, monthly["month"], rotation=35, ha="right")
    ax.legend(frameon=False)
    path = output / "q2_v3_monthly_cost.png"
    fig.savefig(path, dpi=180, metadata={"Software": "Q2 V3 finalize_q2.py"})
    plt.close(fig)
    paths.append(path)

    daily = ledger.groupby("plan_day", as_index=False).agg(
        emergency_kwh=("emergency_kwh", "sum"),
        end_soc_kwh=("actual_soc_after_kwh", "last"),
    )
    dates = pd.to_datetime(daily["plan_day"])
    fig, left = plt.subplots(figsize=(9.0, 4.4), constrained_layout=True)
    right = left.twinx()
    left.fill_between(dates, daily["emergency_kwh"], color="#D1495B", alpha=0.55, linewidth=0, label="Emergency energy")
    right.plot(dates, daily["end_soc_kwh"], color="#287271", linewidth=1.0, label="End-of-day SOC")
    left.set(xlabel="Plan day", ylabel="Emergency energy (kWh)", title="Q2 V3 daily risk and battery state")
    right.set_ylabel("End-of-day SOC (kWh)")
    left.grid(True, alpha=0.2)
    right.grid(False)
    lines = left.get_legend_handles_labels()[0] + right.get_legend_handles_labels()[0]
    labels = left.get_legend_handles_labels()[1] + right.get_legend_handles_labels()[1]
    left.legend(lines, labels, frameon=False, loc="upper right")
    path = output / "q2_v3_daily_risk_soc.png"
    fig.savefig(path, dpi=180, metadata={"Software": "Q2 V3 finalize_q2.py"})
    plt.close(fig)
    paths.append(path)

    fig, axes = plt.subplots(1, 2, figsize=(8.6, 4.1), constrained_layout=True)
    labels = sensitivity["interpretation"].tolist()
    colors = ["#287271", "#D1495B", "#E9C46A", "#4C78A8", "#7A5195"]
    axes[0].bar(labels, sensitivity["common_333d_total_cost_yuan"] / 1e4, color=colors)
    axes[0].set(ylabel="Cost (10,000 yuan)", title="Common 333-day operating cost")
    axes[1].bar(labels, sensitivity["common_333d_emergency_kwh"] / 1e3, color=colors)
    axes[1].set(ylabel="Emergency energy (1,000 kWh)", title="Common 333-day emergency purchase")
    for axis in axes:
        axis.tick_params(axis="x", rotation=20)
    path = output / "q2_v3_interpretation_sensitivity.png"
    fig.savefig(path, dpi=180, metadata={"Software": "Q2 V3 finalize_q2.py"})
    plt.close(fig)
    paths.append(path)
    return paths


def _copy(source: Path, destination: Path) -> None:
    shutil.copy2(source, destination)
    if sha256(source) != sha256(destination):
        raise AssertionError(f"copy hash mismatch: {source} -> {destination}")


def _manifest_key(path: Path) -> str:
    try:
        return str(path.relative_to(VERSION_ROOT))
    except ValueError:
        return path.name


def finalize(args: argparse.Namespace) -> Path:
    output = Path(args.output_dir).resolve()
    figures = Path(args.figures_dir).resolve()
    if output.exists() or (not args.skip_figures and figures.exists()):
        raise FileExistsError("release output paths must not already exist")
    run_paths = {
        "main": Path(args.main_run).resolve(),
        "battery_b": Path(args.battery_b_run).resolve(),
        "time_b": Path(args.time_b_run).resolve(),
        "delay_1": Path(args.delay_1_run).resolve(),
        "reset_feb": Path(args.reset_feb_run).resolve(),
    }
    verified = {name: verify_run(name, path) for name, path in run_paths.items()}
    commits = {item["audit"]["git_commit"] for item in verified.values()}
    if len(commits) != 1:
        raise AssertionError(f"full runs use different source commits: {sorted(commits)}")

    output.mkdir(parents=True, exist_ok=False)
    main_root = run_paths["main"]
    main_ledger = pd.read_csv(main_root / "interval_ledger.csv")
    main_ledger["plan_day"] = main_ledger["plan_day"].astype(str)
    natural_ledger = natural_day_ledger(
        verified["main"]["arrays"], load_data().dates
    )
    blocks = storage_blocks(natural_ledger)
    plan_blocks = storage_blocks(main_ledger)
    events = emergency_events(natural_ledger)
    plan_day_events = emergency_events(main_ledger)
    monthly = monthly_summary(main_ledger)
    official_daily = verified["main"]["daily"].loc[lambda frame: frame["day_index"] >= 31].copy()

    sensitivity_rows = []
    labels = {
        "main": "Main A/A",
        "battery_b": "Battery B",
        "time_b": "Time B",
        "delay_1": "1-slot delay",
        "reset_feb": "Feb SOC reset",
    }
    for name, item in verified.items():
        common = item["daily"].loc[
            item["daily"]["day_index"].between(31, 363)
        ]
        sensitivity_rows.append({
            "case": name,
            "interpretation": labels[name],
            "report_start_day": item["config"]["report_start_day"],
            "report_end_day": item["config"]["report_end_day"],
            "common_333d_total_cost_yuan": float(common["realized_total_cost_yuan"].sum()),
            "common_333d_emergency_kwh": float(common["emergency_kwh"].sum()),
            "common_333d_end_soc_kwh": float(common.iloc[-1]["actual_end_energy_kwh"]),
            **{key: item["summary"][key] for key in (
                "days", "plan_cost_yuan", "emergency_cost_yuan", "total_cost_yuan",
                "grid_purchase_kwh", "emergency_kwh", "end_soc_kwh",
                "balance_residual_kwh", "soc_residual_kwh",
            )},
            "scope_note": {
                "main": "334-day official plan window; common columns stop at 2025-12-30",
                "battery_b": "frozen battery action; common columns stop at 2025-12-30",
                "time_b": "historical cross-day-shift sensitivity; not the newly confirmed natural-day mapping; 333 days through 2025-12-30",
                "delay_1": "Battery A with one 10-minute measurement delay; common columns stop at 2025-12-30",
                "reset_feb": "January history retained; planned and executed SOC reset to 6000 kWh on 2025-02-01",
            }[name],
        })
    sensitivity = pd.DataFrame(sensitivity_rows)

    _copy(main_root / "interval_ledger.csv", output / "q2_interval_ledger.csv")
    _copy(main_root / "policy_arrays.npz", output / "q2_policy_arrays.npz")
    _copy(main_root / "forecast_bundle.npz", output / "q2_forecast_bundle.npz")
    _copy(main_root / "forecast_fit_log.csv", output / "q2_forecast_fit_log.csv")
    _copy(main_root / "january_warmup_summary.json", output / "q2_january_warmup_summary.json")
    official_daily.to_csv(output / "q2_daily_ledger.csv", index=False)
    natural_ledger.to_csv(output / "q2_natural_day_ledger.csv", index=False)
    blocks.to_csv(output / "q2_natural_day_storage_blocks.csv", index=False)
    plan_blocks.to_csv(output / "q2_plan_day_storage_blocks.csv", index=False)
    events.to_csv(output / "q2_emergency_events_natural_day.csv", index=False)
    plan_day_events.to_csv(output / "q2_emergency_events_plan_day.csv", index=False)
    monthly.to_csv(output / "q2_monthly_summary.csv", index=False)
    sensitivity.to_csv(output / "q2_sensitivity_summary.csv", index=False)
    write_json(output / "q2_final_summary.json", verified["main"]["summary"])
    write_json(output / "q2_full_run_audits.json", {name: item["audit"] for name, item in verified.items()})
    for name, item in verified.items():
        write_json(output / f"q2_{name}_config.json", item["config"])
        write_json(output / f"q2_{name}_source_manifest.json", item["manifest"])
        if name != "main":
            write_json(output / f"q2_{name}_summary.json", item["summary"])
            _copy(run_paths[name] / "policy_arrays.npz", output / f"q2_{name}_policy_arrays.npz")
            _copy(run_paths[name] / "daily_ledger.csv", output / f"q2_{name}_daily_ledger_all.csv")

    readback = write_result2(output / "result2.xlsx", main_ledger, blocks, events)
    write_json(output / "q2_result2_readback.json", readback)
    figure_paths = [] if args.skip_figures else make_figures(figures, main_ledger, monthly, sensitivity)
    release_files = sorted(path for path in output.iterdir() if path.is_file() and path.name != "q2_release_manifest.json")
    release_commit = subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=REPO_ROOT, text=True
    ).strip()
    release_sources = [
        Path(__file__),
        Path(__file__).with_name("run_q2_v3.py"),
        Path(__file__).with_name("q2_forecasts.py"),
        Path(__file__).with_name("q2_model.py"),
        Path(__file__).with_name("q2_deep_core.py"),
        Path(__file__).with_name("evaluate_development.py"),
    ]
    release_source_hashes = {
        str(path.relative_to(REPO_ROOT)): sha256(path) for path in release_sources
    }
    if not all(
        _git_blob_sha256(release_commit, relative) == digest
        for relative, digest in release_source_hashes.items()
    ):
        raise AssertionError("release source does not match the release Git commit")
    release_manifest = {
        "schema_version": 1,
        "release_id": "Q2-V3-CAUSAL-FORECAST-FINAL-V1-R2",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "status": "DELIVERY_EXPORT_CHECKS_PASS_P2_PENDING",
        "model_rerun": False,
        "new_conventions_implemented": False,
        "formal_use": False,
        "selected_arm": "B0",
        "selected_model_id": "Q2V2-SW2-REC5",
        "predictor_upgrade_adopted": False,
        "source_run_commit": next(iter(commits)),
        "release_git_commit": release_commit,
        "release_source_hashes": release_source_hashes,
        "template_sha256": sha256(TEMPLATE),
        "source_runs": {name: str(path.relative_to(REPO_ROOT)) for name, path in run_paths.items()},
        "run_audits_passed": all(item["audit"]["passed"] for item in verified.values()),
        "result2_readback_passed": readback["passed"],
        "reporting_windows": {
            "plan_and_cost": "334 plan days from 2025-02-01 00:10 through 2026-01-01 00:10",
            "storage_template": "334 actual natural days, six four-hour blocks each, with true 00:00 and 24:00 SOC projected from archived Time A trajectories",
            "plan_day_storage_reference": "q2_plan_day_storage_blocks.csv retains six nominal groups of 24 plan-row slots; physical boundaries are shifted 10 minutes under mapping A",
            "emergency_template": "334 natural days from 2025-02-01 00:00 through 2026-01-01 00:00",
            "natural_day_bridge": "natural slot 0 uses prior plan-day slot 143; slots 1-143 use current plan-day slots 0-142",
        },
        "output_hashes": {_manifest_key(path): sha256(path) for path in release_files},
        "figure_hashes": {_manifest_key(path): sha256(path) for path in figure_paths},
    }
    write_json(output / "q2_release_manifest.json", release_manifest)
    print(json.dumps({
        "release": str(output),
        "main_total_cost_yuan": verified["main"]["summary"]["total_cost_yuan"],
        "emergency_events": len(events),
        "files": len(release_files) + 1,
        "figures": len(figure_paths),
        "status": "DELIVERY_EXPORT_CHECKS_PASS_P2_PENDING",
    }, ensure_ascii=False, indent=2))
    return output


def parser() -> argparse.ArgumentParser:
    value = argparse.ArgumentParser(description=__doc__)
    value.add_argument("--main-run", default=str(DEFAULT_RUNS["main"]))
    value.add_argument("--battery-b-run", default=str(DEFAULT_RUNS["battery_b"]))
    value.add_argument("--time-b-run", default=str(DEFAULT_RUNS["time_b"]))
    value.add_argument("--delay-1-run", default=str(DEFAULT_RUNS["delay_1"]))
    value.add_argument("--reset-feb-run", default=str(DEFAULT_RUNS["reset_feb"]))
    value.add_argument("--output-dir", default=str(VERSION_ROOT / "results/final_v1"))
    value.add_argument("--figures-dir", default=str(VERSION_ROOT / "figures/final_v1"))
    value.add_argument("--skip-figures", action="store_true")
    return value


if __name__ == "__main__":
    finalize(parser().parse_args())
